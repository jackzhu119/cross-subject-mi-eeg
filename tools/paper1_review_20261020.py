#!/usr/bin/env python3
"""Render only source-reviewed paragraphs into new Paper1 submission files."""
import difflib, hashlib, json, re, subprocess, zipfile
from pathlib import Path
from docx import Document
import fitz

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/"research_runs/PAPER1_SUBMISSION_CORRECTION_20261018"
NEW=ROOT/"research_runs/PAPER1_REVIEW_20261020"

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def blocks(text):
    return [p.strip() for p in re.split(r"\n\s*\n",text.strip())]

def plain(s):
    return re.sub(r"^#{1,6}\s+", "", s).replace("**","").replace(chr(96),"").strip()

def normal(s):
    return re.sub(r"\s+", " ",s.replace("\xa0"," ").replace("–","-")).strip()

def check_tables(a,b):
    assert len(a.tables)==len(b.tables), "table count changed"
    for i,(t,u) in enumerate(zip(a.tables,b.tables)):
        assert [[c.text for c in r.cells] for r in t.rows] == [[c.text for c in r.cells] for r in u.rows], ("scientific table changed",i)

def rewrite_docx(docpath, oldmd, newmd, out):
    before=blocks(oldmd.read_text(encoding="utf-8"))
    after=blocks(newmd.read_text(encoding="utf-8"))
    assert len(before)==len(after), "source markdown structure changed"
    revisions=[(i,a,b) for i,(a,b) in enumerate(zip(before,after)) if a!=b]
    assert 0<len(revisions)<70, ("unexpected revision count",len(revisions))
    base=Document(docpath)
    doc=Document(docpath)
    seen=set()
    mapping=[]
    for blockidx,orig,revised in revisions:
        assert not orig.startswith("|"),"numeric table edit prohibited"
        op=plain(orig); np=plain(revised)
        choices=[]
        for k,p in enumerate(doc.paragraphs):
            if k not in seen and p.text.strip():
                ratio=difflib.SequenceMatcher(None,normal(op),normal(p.text),autojunk=False).ratio()
                choices.append((ratio,k))
        ratio,k=max(choices)
        if ratio<0.83:
            raise AssertionError(f"Unmatched source block {blockidx}: {ratio:.3f} source={op[:100]!r}, closest={doc.paragraphs[k].text[:100]!r}")
        p=doc.paragraphs[k]
        cap=re.match(r"^\*\*(Fig\s+\d+\.[^*]+)\*\*(.*)$",revised,flags=re.S)
        if cap:
            p.clear()
            p.add_run(cap.group(1)).bold=True
            p.add_run(cap.group(2).replace(chr(96),""))
        elif len(p.runs)==1:
            p.runs[0].text=np
        else:
            p.text=np
        seen.add(k)
        mapping.append({"block":blockidx,"docx_paragraph":k,"match":round(ratio,5),
                        "from":op[:70],"to":np[:70]})
    check_tables(base,doc)
    assert len(base.paragraphs)==len(doc.paragraphs)
    for k,(a,b) in enumerate(zip(base.paragraphs,doc.paragraphs)):
        if k not in seen: assert a.text==b.text,("unreviewed text modified",k)
    out.parent.mkdir(parents=True,exist_ok=True)
    doc.save(out)
    saved=Document(out)
    check_tables(base,saved)
    for k,(a,b) in enumerate(zip(base.paragraphs,saved.paragraphs)):
        if k not in seen: assert a.text==b.text,("unreviewed saved text modified",k)
    with zipfile.ZipFile(docpath) as z0,zipfile.ZipFile(out) as z1:
        for name in (x for x in z0.namelist() if x.startswith("word/media/")):
            assert z0.read(name)==z1.read(name),("figure/media changed",name)
    return {"changed_source_blocks":len(revisions),"changed_word_paragraphs":len(seen),
            "tables_preserved":len(base.tables),"original_sha256":sha(docpath),
            "new_sha256":sha(out),"change_map":mapping}

def convert(name):
    doc=NEW/(name+".docx")
    p=subprocess.run(["libreoffice","-env:UserInstallation=file:///tmp/p1review_20261020_lo",
          "--headless","--convert-to","pdf","--outdir",str(NEW),str(doc)],
          text=True,capture_output=True,timeout=180)
    assert p.returncode==0,(name,p.stdout,p.stderr)
    path=NEW/(name+".pdf")
    assert path.exists() and path.stat().st_size>3000
    with fitz.open(path) as pdf:
        pages=len(pdf)
        if name=="Cover_Letter": assert pages in (1,2,3)
        else: assert 10<=pages<=40
        assert all(p.get_text().strip() for p in pdf)
        if name=="main_manuscript":
            text="\n".join(p.get_text() for p in pdf)
            for s in ("21.1667","65.3775","S29","S60","Article 32(1)"):
                assert s in text,("rendered PDF lacks essential text",s)
            assert "01:53:31" not in text
    return {"pages":pages,"sha256":sha(path)}

def main():
    NEW.mkdir(parents=True,exist_ok=True)
    result={}
    for key,docfile,oldmd,newmd,out in [
     ("main","Main_Manuscript_REVISED.docx","manuscript_en_REVISED.md","manuscript_en.md","main_manuscript.docx"),
     ("supplement","S1_Appendix_REVISED.docx","supplementary_en_REVISED.md","supplementary_en.md","S1_Appendix.docx"),
     ("cover","Cover_Letter_REVISED.docx","Cover_Letter_REVISED.md","Cover_Letter.md","Cover_Letter.docx")]:
        result[key]=rewrite_docx(BASE/docfile,BASE/oldmd,NEW/newmd,NEW/out)
    result["pdfs"]={p:convert(p) for p in ("main_manuscript","S1_Appendix","Cover_Letter")}
    m=(NEW/"manuscript_en.md").read_text(encoding="utf-8")
    abstract=m.split("## Abstract\n\n",1)[1].split("\n\n## Introduction",1)[0]
    assert abstract.strip()==(NEW/"Abstract_for_Portal.txt").read_text(encoding="utf-8").strip()
    assert "66/85." in abstract
    receipt={"status":"PASS","source_commit":"e53318846e4a3ce8604de2ffc8f12c7b58ff3953",
        "scope":"reviewed prose only","scientific_results_reused_without_training":True,
        "source_tables_and_media_preserved":True,"abstract_text_matches_main":True,
        "files":result}
    (NEW/"RELEASE_VALIDATION.json").write_text(json.dumps(receipt,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    print(json.dumps({k:{"changed_blocks":v["changed_source_blocks"],"tables":v["tables_preserved"]}
                      for k,v in result.items() if k!="pdfs"}))
if __name__=="__main__":main()
