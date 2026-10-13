#!/usr/bin/env python3
"""Render the redesign work package from saved results and authored Markdown.

No raw EEG is read and no model, prediction, selection, or statistical test is run.
External intervals are copied from the frozen audit; figure source CSVs are saved.
Use --figures-only before the manuscript and supporting text have been authored.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import subprocess
import tempfile

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
import fitz


ROOT = Path(__file__).resolve().parents[1]
SOFFICE = Path("/opt/codex/runtimes/codex-primary-runtime/dependencies/bin/override/soffice")
CSP_RUN = "cpu-csp-20261013T004218Z-2828deb4"
BLUE, AMBER, DARK, GREY = "#235789", "#D27A22", "#263238", "#727C82"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def clean_inline(s: str) -> str:
    s = re.sub(r"!\[([^]]*)\]\([^)]*\)", r"\1", s)
    s = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", s)
    return s.replace("**", "").replace("`", "")


def save_figure(fig: plt.Figure, stem: str, out: Path) -> None:
    for suffix in ("png", "pdf", "svg"):
        kwargs = {"dpi": 300} if suffix == "png" else {}
        fig.savefig(out / f"{stem}.{suffix}", bbox_inches="tight", pad_inches=0.08, **kwargs)
    plt.close(fig)


def figures() -> list[Path]:
    figdir, sourcedir = ROOT / "figures", ROOT / "figures" / "source_data"
    sourcedir.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9,
                         "axes.spines.top": False, "axes.spines.right": False,
                         "axes.labelcolor": DARK, "text.color": DARK,
                         "svg.fonttype": "none", "pdf.fonttype": 42})
    paired = pd.read_csv(ROOT / "results" / "historical_internal_paired.csv")
    audit = json.loads((ROOT / "evidence" / "historical_reaggregation.json").read_text())
    csp = pd.read_csv(ROOT / "results" / CSP_RUN / "participant_confusion_and_metrics.csv")
    loso = csp[csp.evaluation == "new_source_loso"].copy().sort_values("subject")
    if len(paired) != 9 or len(loso) != 9:
        raise ValueError("Expected nine known development participants in each plot")

    psource = sourcedir / "figure1_development_contrasts.csv"
    cols = ["subject", "rank_minus_loss_pp", "fixed_minus_matched_loss_pp", "fixed_minus_historical_rank_pp"]
    paired[cols].to_csv(psource, index=False, float_format="%.12g")
    titles = [("rank_minus_loss_pp", "A  Mean-rank − mean-loss", "Exploratory development\ncomparison"),
              ("fixed_minus_matched_loss_pp", "B  Fixed20 − replayed\nmean-loss", "Matched-runtime replay"),
              ("fixed_minus_historical_rank_pp", "C  Fixed20 − earlier\nmean-rank", "Cross-run comparison")]
    # Author at approximately the document's final width so labels do not
    # become tiny when the three-panel figure is inserted on an A4 page.
    fig, axes = plt.subplots(1, 3, figsize=(7.4, 5.4), sharey=True)
    y = np.arange(len(paired))
    for ax, (col, title, subtitle) in zip(axes, titles):
        values = paired[col].to_numpy()
        ax.barh(y, values, color=[AMBER if s in (3, 8) else BLUE for s in paired.subject], height=0.64)
        ax.axvline(0, color=DARK, lw=0.8)
        ax.axvline(values.mean(), color=GREY, lw=1, ls="--")
        ax.set_yticks(y, [f"S{s}" for s in paired.subject])
        ax.set_xlim(-20, 59)
        ax.set_xticks([0, 20, 40])
        for yi, val in zip(y, values):
            ax.text(val + (0.7 if val >= 0 else -0.7), yi, f"{val:+.2f}", va="center",
                    ha="left" if val >= 0 else "right", fontsize=8.7,
                    bbox={"facecolor": "white", "edgecolor": "none", "pad": 0.5})
        ax.set_title(title + "\n" + subtitle, fontsize=8.8, loc="left", pad=12)
        ax.set_xlabel("BA difference (pp)")
        ax.text(0.02, -0.36, f"Mean {values.mean():+.2f} pp\nMedian {np.median(values):+.2f} pp",
                transform=ax.transAxes, fontsize=8.7)
        ax.grid(axis="x", alpha=0.15)
        ax.set_axisbelow(True)
        if ax is not axes[0]:
            ax.tick_params(axis="y", left=False)
    axes[0].invert_yaxis()
    fig.text(0.08, 0.015, "Amber: S3 and S8; 87.57% of the signed aggregate gain in panel A.", fontsize=8.7)
    fig.subplots_adjust(left=0.07, right=0.99, wspace=0.22, bottom=0.34, top=0.80)
    save_figure(fig, "figure1_development_heterogeneity", figdir)

    external = []
    for key, label in (("q14_physionet", "PhysioNet (109)"), ("q15_cho", "Cho2017 (52)"), ("q15_lee", "Lee2019_MI (54)")):
        v = audit[key]
        freeze = v.get("frozen_primary_resampling_replayed", v.get("frozen_primary_exact_zero_policy_replayed"))
        lo, hi = freeze["participant_bootstrap_percentile_95_ci_pp"]
        external.append({"dataset": label, "participant_count": v["shared_minus_broad"]["n_participants"],
                         "shared_minus_broad_balanced_accuracy_pp": v["shared_minus_broad"]["mean_difference_pp"],
                         "frozen_bootstrap_95_ci_lower_pp": lo, "frozen_bootstrap_95_ci_upper_pp": hi,
                         "scope": "historical_complete_frozen_pipeline_contrast_not_pooled"})
    ext = pd.DataFrame(external)
    ext.to_csv(sourcedir / "figure2_external_complete_pipeline_contrasts.csv", index=False, float_format="%.12g")
    fig, ax = plt.subplots(figsize=(7.4, 3.3))
    for yi, row in ext.iterrows():
        val, lo, hi = row.shared_minus_broad_balanced_accuracy_pp, row.frozen_bootstrap_95_ci_lower_pp, row.frozen_bootstrap_95_ci_upper_pp
        ax.errorbar(val, yi, xerr=[[val-lo], [hi-val]], fmt="o", color=AMBER if yi == 1 else BLUE,
                    markersize=7, capsize=4, elinewidth=1.7)
        ax.text(1.7, yi, f"{val:+.3f} [{lo:+.3f}, {hi:+.3f}]", va="center", fontsize=8.5)
    ax.axvline(0, color=DARK, lw=1, ls="--")
    ax.set_yticks(range(3), ext.dataset)
    ax.set_xlim(-2.6, 4.6)
    ax.set_xticks([-2, -1, 0, 1])
    ax.set_ylim(2.55, -0.70)
    ax.set_xlabel("Shared-input − broad balanced accuracy (pp)")
    fig.suptitle("Historical external evaluations: complete frozen pipeline contrasts", x=0.23, y=0.98,
                 ha="left", fontsize=10)
    ax.text(1.7, -0.43, "Difference [frozen 95% CI]", fontsize=8.5)
    ax.grid(axis="x", alpha=0.15)
    fig.text(0.13, 0.015, "Participant bootstrap intervals; cohorts, tasks and pipelines are heterogeneous. No pooled estimate.", fontsize=8)
    fig.subplots_adjust(left=0.23, right=0.99, bottom=0.25, top=0.88)
    save_figure(fig, "figure2_historical_external_contrasts", figdir)

    csource = loso[["subject", "n_trials", "recall_left", "recall_right", "balanced_accuracy_percent"]].copy()
    csource.to_csv(sourcedir / "figure3_new_cpu_csp_source_loso.csv", index=False, float_format="%.12g")
    fig, ax = plt.subplots(figsize=(7.4, 4.4))
    y = np.arange(9)
    ax.plot(loso.recall_left.to_numpy()*100, y-0.16, "o", color=BLUE, label="Left-hand recall", markersize=5.5)
    ax.plot(loso.recall_right.to_numpy()*100, y+0.16, "s", color=AMBER, label="Right-hand recall", markersize=5.5)
    ax.plot(loso.balanced_accuracy_percent.to_numpy(), y, "D", color=DARK, label="Balanced accuracy", markersize=5.5)
    for yi, row in enumerate(loso.itertuples()):
        ax.text(102, yi, f"{row.balanced_accuracy_percent:.2f}", va="center", fontsize=8.7)
    ax.set_xlim(0, 117)
    ax.set_xticks([0, 25, 50, 75, 100])
    ax.set_yticks(y, [f"S{s}" for s in loso.subject])
    ax.set_ylim(8.7, -0.7)
    ax.axvline(50, lw=1, ls="--", color=GREY)
    ax.set_xlabel("Recall or balanced accuracy (%)")
    ax.set_title("New CPU source-only diagnostic: nine-participant CSP–LDA LOSO", loc="left", fontsize=10)
    ax.text(102, -0.95, "BA (%)", fontsize=8.7)
    ax.legend(loc="upper center", bbox_to_anchor=(0.48, -0.14), ncol=3, frameon=False, fontsize=8.5)
    ax.grid(axis="x", alpha=0.15)
    fig.text(0.11, 0.014, "New source fits only; 288 trials/person. Equal-person BA: 61.65%; sample SD: 14.12 pp.", fontsize=8.2)
    fig.subplots_adjust(left=0.12, right=0.99, bottom=0.25, top=0.86)
    save_figure(fig, "figure3_new_cpu_csp_source_diagnostic", figdir)

    captions = """# Figure captions

**Figure 1. Participant heterogeneity in historical source-duration comparisons.** Each bar is the saved, seed-averaged balanced-accuracy difference for one of nine BNCI participants. Panel A compares the exploratory mean-rank follow-up with the earlier mean-loss development condition. S3 and S8 contribute 87.57% of its signed aggregate gain. Panel B compares fixed20 with the matched-runtime replay of the historical mean-loss schedule. Panel C is a retrospective cross-run comparison of fixed20 with the earlier mean-rank run; the earlier effective runtime is not authenticated. Dashed lines show equal-participant means. These four-class development comparisons do not confirm a generally superior duration selector. pp, percentage points.

**Figure 2. Historical external complete-pipeline contrasts.** Points show the equal-participant shared-input minus broad balanced-accuracy difference, with the original frozen participant-bootstrap 95% confidence interval. The three cohorts are shown separately; their task definitions, eligible trials, and complete pipelines differ. The negative Cho2017 result and intervals spanning zero for PhysioNet and Lee2019_MI are retained. The comparison does not isolate spectral representation, parameter sharing, or training duration. No pooled cross-cohort estimate is shown.

**Figure 3. New CPU source-only CSP–LDA diagnostic.** Left-hand recall, right-hand recall, and balanced accuracy are calculated from the stored held-out predictions of nine newly executed leave-one-participant-out fits. All eight remaining participants contribute both sessions to each training fold; each held-out participant contributes 288 trials, balanced between classes. Equal-participant mean balanced accuracy is 61.65% with sample SD 14.12 pp. The separate all-source resubstitution fit is not plotted. These are new retrospective source fits in a recorded current runtime, not a replay of the archived external checkpoint, a prospective target-cohort result, or evidence for a superior neural method. The dashed line marks 50% and does not represent a participant-level statistical test.
"""
    (ROOT / "FIGURE_CAPTIONS.md").write_text(captions, encoding="utf-8")
    return [ROOT / "results" / "historical_internal_paired.csv", ROOT / "evidence" / "historical_reaggregation.json",
            ROOT / "results" / CSP_RUN / "participant_confusion_and_metrics.csv"]


def field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText"); text.text = instruction
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, end])


def inline(paragraph, text: str) -> None:
    # Preserve inline links as legible labels and URL text; soft-break long URLs.
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", lambda m: m[1] + " (" + m[2] + ")", text)
    tokens = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text)
    for token in tokens:
        bold = token.startswith("**") and token.endswith("**")
        italic = token.startswith("*") and token.endswith("*") and not bold
        code = token.startswith("`") and token.endswith("`")
        token = token[2:-2] if bold else token[1:-1] if italic or code else token
        # Prevent long repository paths and links from protruding into margins.
        if "https://" in token or len(token) > 80:
            token = re.sub(r"([/_-])", lambda m: m[1] + "\u200b", token)
        run = paragraph.add_run(token)
        run.bold, run.italic = bold, italic
        if code:
            run.font.name = "DejaVu Sans Mono"; run.font.size = Pt(8)


def table_rows(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        row = [v.strip() for v in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", v.replace(" ", "")) for v in row):
            continue
        rows.append(row)
    width = max(map(len, rows), default=0)
    if width == 0:
        raise ValueError("Empty Markdown table")
    return [row + [""] * (width-len(row)) for row in rows]


def markdown_docx(source: Path, output: Path, include_figures: bool) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.72)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Liberation Serif", Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.12
    for name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        doc.styles[name].font.name = "Liberation Sans"
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor.from_string("263238")
        doc.styles[name].paragraph_format.keep_with_next = True
    doc.core_properties.title = clean_inline(source.read_text().splitlines()[0].lstrip("# "))
    doc.core_properties.author = "Ziyuan Zhu"
    doc.core_properties.subject = "Research redesign work version; independent confirmation and neural fits pending"
    header = section.header.paragraphs[0]
    header.text = "RESEARCH WORK VERSION — INDEPENDENT CONFIRMATION PENDING"
    for run in header.runs:
        run.font.name = "Liberation Sans"; run.font.size = Pt(7.3); run.font.color.rgb = RGBColor.from_string("727C82")
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("Page "); field(footer, "PAGE")
    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0
    images_inserted = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line or line == "---":
            i += 1; continue
        if line.startswith("```"):
            block = []; i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i]); i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(block)); run.font.name = "DejaVu Sans Mono"; run.font.size = Pt(8)
            i += 1; continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i]); i += 1
            rows = table_rows(block)
            table = doc.add_table(rows=1, cols=len(rows[0]))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            # Fixed page-wide columns. Authoring keeps cells concise; narrow cells wrap.
            widths = [6.83/len(rows[0])] * len(rows[0])
            if len(widths) >= 5:
                widths[0] *= 1.25
                widths[1:] = [(6.83-widths[0])/(len(widths)-1)]*(len(widths)-1)
            for ci, width in enumerate(widths):
                table.columns[ci].width = Inches(width)
            for ri, values in enumerate(rows):
                cells = table.rows[0].cells if ri == 0 else table.add_row().cells
                for ci, value in enumerate(values):
                    cells[ci].width = Inches(widths[ci])
                    cells[ci].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p = cells[ci].paragraphs[0]; p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.space_before = Pt(3); p.paragraph_format.line_spacing = 1
                    inline(p, value)
                    for run in p.runs:
                        run.font.size = Pt(8 if len(widths) >= 5 else 9)
                        if ri == 0: run.bold = True
                trpr = table.rows[ri]._tr.get_or_add_trPr()
                no_split = OxmlElement("w:cantSplit"); trpr.append(no_split)
                if ri == 0:
                    repeat = OxmlElement("w:tblHeader"); trpr.append(repeat)
            doc.add_paragraph()
            continue
        if line.startswith("#"):
            match = re.match(r"^(#{1,6})\s+(.*)$", line)
            if match:
                level = len(match[1])
                if i == 0:
                    p = doc.add_paragraph(style="Title"); inline(p, match[2])
                else:
                    if include_figures and match[2].strip().lower() == "figures":
                        doc.add_page_break()
                    p = doc.add_paragraph(style=f"Heading {min(level, 3)}"); inline(p, match[2])
                i += 1; continue
        if line.startswith("!["):
            match = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", line)
            if match:
                path = (source.parent / match[2]).resolve()
                if ROOT.resolve() not in path.parents or not path.is_file():
                    raise ValueError(f"Figure link is not a local work-package file: {match[2]}")
                if images_inserted:
                    doc.add_page_break()
                doc.add_picture(str(path), width=Inches(6.8))
                images_inserted += 1
            i += 1; continue
        bullet = re.match(r"^(?:[-*]\s+|\d+\.\s+)(.*)", line)
        if bullet:
            numbered = line[0] not in "-*"
            p = doc.add_paragraph(style=None if numbered else "List Bullet")
            # Preserve authored reference numbers; Word's shared automatic list
            # would otherwise continue a preceding Methods/checklist list.
            inline(p, line if numbered else bullet[1]); i += 1; continue
        paragraph = [line.lstrip("> ")]; i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,6}\s|\||```|[-*]\s|\d+\.\s|!\[)", lines[i].strip()):
            paragraph.append(lines[i].strip().lstrip("> ")); i += 1
        p = doc.add_paragraph(); inline(p, " ".join(paragraph))
    if include_figures and not images_inserted:
        doc.add_page_break(); doc.add_heading("Figures", level=1)
        stems = ["figure1_development_heterogeneity", "figure2_historical_external_contrasts", "figure3_new_cpu_csp_source_diagnostic"]
        caption_lines = [v.strip() for v in (ROOT / "FIGURE_CAPTIONS.md").read_text().splitlines() if v.startswith("**Figure")]
        for index, (stem, caption) in enumerate(zip(stems, caption_lines)):
            if index:
                doc.add_page_break()
            doc.add_picture(str(ROOT / "figures" / f"{stem}.png"), width=Inches(6.8))
            p = doc.add_paragraph(); inline(p, caption)
            for run in p.runs: run.font.size = Pt(10)
    doc.save(output)


def export_pdf(docx_path: Path) -> Path:
    if not SOFFICE.is_file():
        raise RuntimeError("Existing LibreOffice executable is unavailable; do not claim PDF was verified")
    with tempfile.TemporaryDirectory(prefix="redesign-soffice-") as temp:
        subprocess.run([str(SOFFICE), f"-env:UserInstallation=file://{temp}", "--headless", "--convert-to", "pdf",
                        "--outdir", str(docx_path.parent), str(docx_path)], check=True, capture_output=True, text=True, timeout=150)
    out = docx_path.with_suffix(".pdf")
    if not out.is_file():
        raise RuntimeError(f"PDF conversion did not create {out.name}")
    return out


def pdf_review(path: Path) -> dict:
    with fitz.open(path) as pdf:
        empty = []
        outside = []
        texts = []
        for i, page in enumerate(pdf):
            texts.append(page.get_text())
            if not page.get_text().strip() and not page.get_images():
                empty.append(i+1)
            for block in page.get_text("dict")["blocks"]:
                if block["type"] != 0: continue
                for line in block["lines"]:
                    x0,y0,x1,y1 = line["bbox"]
                    if x0 < -0.5 or y0 < -0.5 or x1 > page.rect.width+0.5 or y1 > page.rect.height+0.5:
                        outside.append({"page": i+1, "bbox": [x0,y0,x1,y1]})
        text = "\n".join(texts)
        return {"pages": len(pdf), "empty_pages": empty, "text_outside_page": outside,
                "visible_work_status": "RESEARCH WORK VERSION" in text,
                "source_content_character_count": len(text), "sha256": sha(path)}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--figures-only", action="store_true")
    args = parser.parse_args()
    inputs = figures()
    checks = {}
    if not args.figures_only:
        for stem, include in (("REVISED_MANUSCRIPT", True), ("SUPPORTING_INFORMATION", False)):
            source = ROOT / f"{stem}.md"
            if not source.is_file():
                raise FileNotFoundError(f"Awaiting authored {source.name}; figures were saved without fabricated text")
            inputs.append(source)
            destination = ROOT / f"{stem}.docx"
            markdown_docx(source, destination, include)
            pdf = export_pdf(destination)
            checks[pdf.name] = pdf_review(pdf)
    files = sorted((ROOT / "figures").rglob("*"))
    outputs = [p for p in files if p.is_file()] + [ROOT / "FIGURE_CAPTIONS.md"]
    if not args.figures_only:
        outputs += [ROOT / f"{name}.{suffix}" for name in ("REVISED_MANUSCRIPT", "SUPPORTING_INFORMATION") for suffix in ("docx", "pdf")]
    record = {"schema_version": 1, "built_at_utc": datetime.now(timezone.utc).isoformat(),
              "status": "FIGURES_ONLY" if args.figures_only else "WORK_ARTIFACTS_BUILT_PENDING_VISUAL_REVIEW",
              "scope": "saved numeric tables, frozen intervals and authored text only; no science inference or model fits",
              "new_model_fits": 0, "target_inference_runs": 0,
              "generator_sha256": sha(Path(__file__).resolve()),
              "render_versions": {"matplotlib": matplotlib.__version__, "numpy": np.__version__,
                                  "pandas": pd.__version__, "pymupdf": fitz.VersionBind},
              "input_sha256": {str(p.relative_to(ROOT)): sha(p) for p in inputs},
              "output_sha256": {str(p.relative_to(ROOT)): sha(p) for p in outputs}, "pdf_checks": checks,
              "figure_source_rules": {"Figure 1": "nine saved historical participant rows; no new statistical intervals",
                                      "Figure 2": "copied frozen participant-bootstrap intervals; no pooling",
                                      "Figure 3": "nine new CPU source LOSO confusion rows; excludes resubstitution"}}
    out = ROOT / "evidence" / "work_artifact_build_receipt.json"
    out.write_text(json.dumps(record, indent=2)+"\n")
    print(json.dumps({"status": record["status"], "receipt": str(out), "pdf_checks": checks}, indent=2))


if __name__ == "__main__":
    main()
