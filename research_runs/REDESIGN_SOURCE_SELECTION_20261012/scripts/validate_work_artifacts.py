#!/usr/bin/env python3
"""Check rendering inputs, document structure and plotted saved values.

This checks work-package artifacts only. It does not execute models, recompute
scientific intervals, or establish independent target-cohort validation.
The optional visual flag records a completed agent layout review, not author
approval or a substitute for the numerical and provenance audit.
"""
import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re

from docx import Document
import fitz
import numpy as np
import pandas as pd


ROOT = Path(__file__).resolve().parents[1]


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--visual-review-confirmed", action="store_true")
    args = parser.parse_args()
    p = ROOT
    build_path = p / "evidence/work_artifact_build_receipt.json"
    build = json.loads(build_path.read_text())
    checks = []

    def check(name, passed, detail=None):
        checks.append({"name": name, "passed": bool(passed), "detail": detail})

    for name in ["REVISED_MANUSCRIPT", "SUPPORTING_INFORMATION"]:
        check(name+" current Markdown matches rendered input", sha(p/(name+".md")) == build["input_sha256"][name+".md"])
        doc = Document(p/(name+".docx"))
        with fitz.open(p/(name+".pdf")) as pdf:
            text = "\n".join(page.get_text() for page in pdf)
            check(name+" PDF has work status on every page", all("RESEARCH WORK VERSION" in page.get_text() for page in pdf))
        check(name+" no blank/outside-page text", not build["pdf_checks"][name+".pdf"]["empty_pages"] and not build["pdf_checks"][name+".pdf"]["text_outside_page"])
        check(name+" current DOCX hash matches build receipt", sha(p/(name+".docx")) == build["output_sha256"][name+".docx"])
        check(name+" current PDF hash matches build receipt", sha(p/(name+".pdf")) == build["output_sha256"][name+".pdf"])
        check(name+" expected tables", len(doc.tables) == (1 if name == "REVISED_MANUSCRIPT" else 10), {"table_count": len(doc.tables)})
        if name == "REVISED_MANUSCRIPT":
            check("Main document has exactly three figure images", len(doc.inline_shapes) == 3)
            check("Main PDF has every numbered reference", all(re.search(rf"(?m)^{i}\. ", text) for i in range(1, 21)))
            for i in range(1, 4):
                check(f"Figure {i} caption occurs once in main PDF", len(re.findall(rf"Figure {i}\.", text)) == 1)
        else:
            for i in range(1, 11):
                check(f"Table S{i} caption occurs once in SI PDF", len(re.findall(rf"Table S{i}\.", text)) == 1)
            check("S4 physiological section occurs once", text.count("S4. Historical physiological description") == 1)
            check("SI prospective grid labelled not executed", "Planned core validation — NOT EXECUTED" in text)

    original = pd.read_csv(p/"results/historical_internal_paired.csv")
    plotted = pd.read_csv(p/"figures/source_data/figure1_development_contrasts.csv")
    check("Figure1 all nine participant identities", plotted.subject.tolist() == original.subject.tolist())
    for col in plotted.columns[1:]:
        check("Figure1 frozen values "+col, np.allclose(plotted[col], original[col], atol=1e-9, rtol=0))
    original = pd.read_csv(p/"results/cpu-csp-20261013T004218Z-2828deb4/participant_confusion_and_metrics.csv")
    original = original[original.evaluation == "new_source_loso"].sort_values("subject")
    plotted = pd.read_csv(p/"figures/source_data/figure3_new_cpu_csp_source_loso.csv")
    check("Figure3 source LOSO identity only", plotted.subject.tolist() == original.subject.tolist())
    for col in plotted.columns[1:]:
        check("Figure3 copied values "+col, np.allclose(plotted[col], original[col], atol=1e-9, rtol=0))
    audit = json.loads((p/"evidence/historical_reaggregation.json").read_text())
    plotted = pd.read_csv(p/"figures/source_data/figure2_external_complete_pipeline_contrasts.csv")
    for i, key in enumerate(["q14_physionet", "q15_cho", "q15_lee"]):
        item = audit[key]
        frozen = item.get("frozen_primary_resampling_replayed", item.get("frozen_primary_exact_zero_policy_replayed"))
        row = plotted.iloc[i]
        check("Figure2 "+key+" point", abs(row.shared_minus_broad_balanced_accuracy_pp - item["shared_minus_broad"]["mean_difference_pp"]) < 1e-9)
        check("Figure2 "+key+" original frozen CI", np.allclose([row.frozen_bootstrap_95_ci_lower_pp, row.frozen_bootstrap_95_ci_upper_pp], frozen["participant_bootstrap_percentile_95_ci_pp"], atol=1e-9, rtol=0))
    for stem in ["figure1_development_heterogeneity", "figure2_historical_external_contrasts", "figure3_new_cpu_csp_source_diagnostic"]:
        for ext in ["png", "pdf", "svg"]:
            check(stem+"."+ext+" exists", (p/"figures"/f"{stem}.{ext}").is_file())
    check("Figures source units are explicit", plotted.scope.eq("historical_complete_frozen_pipeline_contrast_not_pooled").all())
    passed = all(c["passed"] for c in checks)
    record = {
        "schema_version": 1, "reviewed_at_utc": datetime.now(timezone.utc).isoformat(),
        "status": "PASS_FOR_WORK_ARTIFACT_RENDERING_NOT_SCIENTIFIC_COMPLETION" if passed and args.visual_review_confirmed else "NUMERIC_PASS_VISUAL_PENDING" if passed else "FAIL",
        "check_count": len(checks), "passed_count": sum(c["passed"] for c in checks), "checks": checks,
        "visual_review": {
            "confirmed_by_agent": args.visual_review_confirmed,
            "scope": "All 17 pages reviewed in contact sheets; detailed main external-table and figure-page views; charts inspected at source size. Layout review only, not semantic revalidation of all prose.",
            "main_pages": build["pdf_checks"]["REVISED_MANUSCRIPT.pdf"]["pages"],
            "si_pages": build["pdf_checks"]["SUPPORTING_INFORMATION.pdf"]["pages"],
            "repairs": ["Removed Figure2 title/header overlap", "Increased Figure1 final-width type and moved means/medians below axis labels", "Removed y ticks in Figure1 panels B/C and increased the common left label margin so minus signs are clear", "SI author removed premature duplicate S4/TableS8 block; S1–S10 captions occur once in order"],
            "remaining_layout_issues": [],
            "preview_directory_outside_repository": "/workspace/mi-eeg-redesign-qa-20261012/artifact_previews",
        },
        "build_receipt_sha256": sha(build_path), "output_sha256": build["output_sha256"],
        "new_model_fits_during_artifact_generation": 0, "new_target_inference_during_artifact_generation": 0,
        "scientific_status": "Independent target evaluation and redesigned neural grid not executed",
        "validator_sha256": sha(Path(__file__).resolve()),
    }
    (p/"evidence/work_artifact_review.json").write_text(json.dumps(record, indent=2)+"\n")
    print(json.dumps({"status": record["status"], "passed_count": record["passed_count"], "check_count": len(checks)}))
    if not passed:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
