"""Check the working delivery and frozen bytes, without fitting/inference.

This checks declared content and artifact bindings; it cannot close scientific
confirmation gates or certify every hidden implementation error absent.
"""
from pathlib import Path
import csv,hashlib,json,re,subprocess
from datetime import datetime,timezone
import fitz
from docx import Document

BASE=Path(__file__).resolve().parents[1]
ROOT=BASE.parents[1]
PARENT='39ab4fe0492f4d2ab4c679c752634d3f74bb7bc8'
CHECKS={};DETAILS={}
def sha(p):
    h=hashlib.sha256()
    with Path(p).open('rb') as f:
        for b in iter(lambda:f.read(8*1024*1024),b''):h.update(b)
    return h.hexdigest()
def check(name,value,detail=None):
    CHECKS[name]=bool(value)
    if detail is not None:DETAILS[name]=detail

def citations(s):
    found=set()
    for x in re.findall(r'\[((?:\d+[–\-,]?\s*)+)\]',s):
        for item in x.split(','):
            item=item.strip()
            if '–' in item or '-' in item:
                a,b=re.split('[–-]',item);found.update(range(int(a),int(b)+1))
            else:found.add(int(item))
    return found

def main():
    old=ROOT/'research_runs/PLOS_ONE_POSTSUBMISSION_REVIEW_20261010/evidence'
    original=json.loads((old/'source_protection_baseline.json').read_text())['scientific_files']
    mismatches=[p for p,h in original.items() if not (ROOT/p).is_file() or sha(ROOT/p)!=h]
    check('16417_scientific_files_unchanged',len(original)==16417 and not mismatches,mismatches)
    submitted=json.loads((old/'submitted_package_preservation.json').read_text())['all_submitted_package_files']
    wrong=[p for p,v in submitted.items() if not (ROOT/p).is_file() or sha(ROOT/p)!=v['sha256']]
    check('73_submitted_v1_1_files_unchanged',len(submitted)==73 and not wrong,wrong)
    prefix='research_runs/PLOS_ONE_POSTSUBMISSION_REVIEW_20261010/'
    tree=subprocess.check_output(['git','ls-tree','-rz',PARENT,'--',prefix],cwd=ROOT)
    old_work={line.split(b'\t',1)[1].decode():line.split(b'\t',1)[0].split()[2].decode() for line in tree.split(b'\0') if line}
    diff=[]
    for p,h in old_work.items():
        actual=subprocess.check_output(['git','hash-object',str(ROOT/p)],cwd=ROOT,text=True).strip() if (ROOT/p).is_file() else None
        if actual!=h:diff.append(p)
    check('prior_revision_all_git_blobs_unchanged',bool(old_work) and not diff,{'files':len(old_work),'mismatches':diff})
    changed=subprocess.check_output(['git','diff','--name-only',PARENT],cwd=ROOT,text=True).splitlines()
    check('tracked_changes_only_new_research_directory',all(x.startswith(str(BASE.relative_to(ROOT))+'/') for x in changed))
    receipt=json.loads((BASE/'evidence/cpu_csp_preanalysis_commit_receipt.json').read_text())
    wrong=[p for p,h in receipt['protocol_files'].items() if sha(ROOT/p)!=h]
    check('executed_csp_preanalysis_bytes_unchanged',not wrong,wrong)
    run=BASE/'results/cpu-csp-20261013T004218Z-2828deb4'
    val=json.loads((run/'independent_validation.json').read_text())
    check('real_csp_validator_scope_and_counters',val['status']=='PASS_STORED_SCORES_PARTITIONS_RAW_LABELS_VERIFIED' and val['new_source_fits']==10 and val['new_neural_fits']==0 and val['target_fits']==0 and val['new_result_is_independent_confirmatory_evidence'] is False)
    metrics=list(csv.DictReader((run/'participant_confusion_and_metrics.csv').open()))
    loso=[float(x['balanced_accuracy_percent']) for x in metrics if x['evaluation']=='new_source_loso']
    resub=[float(x['balanced_accuracy_percent']) for x in metrics if x['evaluation']=='new_all_source_resubstitution']
    check('csp_scores_recount',len(loso)==len(resub)==9 and abs(sum(loso)/9-val['fresh_loso_mean_participant_ba_percent'])<1e-10 and abs(sum(resub)/9-val['fresh_all_source_resubstitution_mean_participant_ba_percent'])<1e-10)
    md=(BASE/'REVISED_MANUSCRIPT.md').read_text();si=(BASE/'SUPPORTING_INFORMATION.md').read_text()
    hist=json.loads((BASE/'evidence/historical_reaggregation.json').read_text())
    for key in ['q14_physionet','q15_cho','q15_lee']:
        a=hist[key];z=a['frozen_primary_exact_zero_policy_replayed'] if key=='q14_physionet' else a['frozen_primary_resampling_replayed']
        ci=z['participant_bootstrap_percentile_95_ci_pp']
        values=[f"{a['BROAD_EEGNET']['mean_balanced_accuracy']*100:.3f}",f"{a['MU_BETA_SHARED']['mean_balanced_accuracy']*100:.3f}",f"{a['shared_minus_broad']['mean_difference_pp']:+.3f}",f"[{ci[0]:+.3f},{ci[1]:+.3f}]"]
        # Negative lower limit is used as unicode minus in the authored table.
        check(key+'_main_table_frozen_values',all(x.replace('-','−') in md or x in md for x in values),values)
    for name,x in [('loss',hist['q5_mean_loss']['mean_balanced_accuracy']*100),('rank',hist['q8_mean_rank']['mean_balanced_accuracy']*100),('fixed',hist['fixed20']['mean_balanced_accuracy']*100),('fresh_loso',val['fresh_loso_mean_participant_ba_percent']),('fresh_resub',val['fresh_all_source_resubstitution_mean_participant_ba_percent'])]:
        check(name+'_displayed_value',f'{x:.2f}%' in md)
    refs=json.loads((BASE/'references.json').read_text())
    nums=[r['work_reference_number'] for r in refs]
    before=md.split('## References')[0]
    cited=citations(before)
    check('references_continuous_and_used',nums==list(range(1,len(refs)+1)) and cited==set(nums),{'cited':sorted(cited),'references':nums})
    check('new_dreyer_formal_metadata_bound',refs[17]['title']==json.loads((BASE/'evidence/dreyer_bibliographic_metadata.json').read_text())['message']['title'][0] and refs[17]['pages']=='580')
    for n in range(1,4):
        check(f'figure_{n}_cited_and_exported',f'Fig {n}' in md and all((BASE/('figures/'+{1:'figure1_development_heterogeneity',2:'figure2_historical_external_contrasts',3:'figure3_new_cpu_csp_source_diagnostic'}[n]+'.'+ext)).is_file() for ext in ['png','svg','pdf']))
    for n in range(1,11):
        check(f'supplement_table_s{n}_once',len(re.findall(rf'^Table S{n}\.',si,re.M))==1)
    check('no_duplicate_physiology_section',si.count('## S4.')==1)
    check('manuscript_explicit_work_scope','not submission-ready' in md and 'have not been executed' in md and 'zero new neural fits' in md)
    check('negative_and_uncertain_results_retained',all(x in md for x in ['−1.513','0.603670','no corrected established gain','not absolute contralateral ERD','no p-values']))
    check('source_and_target_nominal_design_not_results',json.loads((BASE/'config/core_experiment_v1.json').read_text())['new_neural_fits_started']==0 and json.loads((BASE/'evidence/dataset_candidate_audit.json').read_text())['independent_confirmatory_evaluation_completed'] is False)
    real_probe=json.loads((BASE/'evidence/core_source_input_probe.json').read_text())
    check('real_source19_probe_bound',real_probe['status']=='PASS_REAL_SOURCE_INPUT_EXTRACTION_NO_MODEL_FITS' and real_probe['real_source_trials']==2592 and real_probe['new_neural_fits']==0 and real_probe['target_predictions']==0 and real_probe['preprocessor_sha256']==sha(BASE/'scripts/core_source_inputs.py') and real_probe['config_sha256']==sha(BASE/'config/core_experiment_v1.json') and all(v==[2592,19,320] for v in real_probe['view_shapes'].values()))
    build=json.loads((BASE/'evidence/work_artifact_build_receipt.json').read_text())
    check('current_authored_text_bound_to_exports',all(sha(BASE/name)==build['input_sha256'][name] for name in ['REVISED_MANUSCRIPT.md','SUPPORTING_INFORMATION.md']))
    wrong_exports=[name for name,h in build['output_sha256'].items() if not (BASE/name).exists() or sha(BASE/name)!=h]
    check('all_artifact_outputs_bound',not wrong_exports,wrong_exports)
    for stem in ['REVISED_MANUSCRIPT','SUPPORTING_INFORMATION']:
        doc=Document(BASE/(stem+'.docx'))
        text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        text=text.replace('\u200b','')
        check(stem+'_docx_work_status',bool(re.search(r'not[^.]{0,180}submission-ready',text.lower(),re.S)))
        pdf=fitz.open(BASE/(stem+'.pdf'));pt='\n'.join(p.get_text() for p in pdf)
        check(stem+'_pdf_has_text_and_work_status',len(pdf)>0 and len(pt)>3000 and bool(re.search(r'not[^.]{0,180}submission-ready',pt.lower(),re.S)))
        check(stem+'_pdf_no_empty_pages',all(len(p.get_text().strip())>30 for p in pdf),{'pages':len(pdf)})
    check('no_raw_EEG_bundled',not [p.name for p in BASE.rglob('*') if p.is_file() and p.suffix.lower() in {'.mat','.edf','.bdf','.gdf','.npz'}])
    software=json.loads((BASE/'evidence/all_software_tests.json').read_text())
    check('153_final_software_tests_passed',software['tests']==153 and software['failures']==software['errors']==software['skipped']==0)
    for name in ['core_source_inputs.py','train_source_selection_core.py','infer_selection_core.py','analyze_redesign_scores.py','acquire_dreyer_runs.py']:
        check(name+'_tested_version_bound',software['software_source_sha256'][name]==sha(BASE/'scripts'/name))
    for f in ['manuscript_independent_review.json','manuscript_export_independent_review.json','future_core_independent_review.json']:
        peer=json.loads((BASE/'evidence'/f).read_text())
        check(f+'_scoped_peer_pass',str(peer.get('status','')).startswith('PASS') and not peer.get('failed_checks') and peer.get('failed',0)==0)
    failed=[k for k,v in CHECKS.items() if not v]
    output={'checked_at_utc':datetime.now(timezone.utc).isoformat(),'status':'PASS_TECHNICAL_WORK_DELIVERY_SCIENTIFIC_GATES_OPEN' if not failed else 'FAIL_WORK_DELIVERY','scope':'Saved byte/hash/content/figure/export preservation checks; no EEG fits/inference, no independent target confirmation or submission-ready certification','check_count':len(CHECKS),'checks':CHECKS,'failed_checks':failed,'details':DETAILS,'new_real_source_csp_fits':10,'new_real_neural_fits':0,'new_target_predictions':0,'target_fits':0,'independent_confirmatory_evaluation_completed':False,'submission_ready':False}
    (BASE/'evidence/final_work_delivery_validation.json').write_text(json.dumps(output,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:output[k] for k in ['status','check_count','failed_checks','submission_ready']}))
    return bool(failed)
if __name__=='__main__': raise SystemExit(main())
