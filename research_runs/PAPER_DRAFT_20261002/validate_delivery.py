"""Check delivered documents against archived evidence; no model execution."""
from pathlib import Path
import hashlib
import json
import re
import subprocess
import zipfile
import xml.etree.ElementTree as ET

import fitz
import numpy as np
import pandas as pd
from docx import Document

OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[1]
content=json.loads((OUT/'manuscript_content.json').read_text())
numbers=json.loads((OUT/'paper_numbers.json').read_text())
snapshot=json.loads((OUT/'evidence/source_snapshot.json').read_text())
inventory=pd.read_csv(OUT/'tables/completed_internal_inventory.csv')
qc=pd.read_csv(OUT/'tables/early_binary_qc_inventory.csv')
internal=pd.read_csv(OUT/'tables/internal_subjects.csv')
external=pd.read_csv(OUT/'tables/external_subjects.csv')
assert len(internal)==9 and internal.subject.is_unique
assert len(external)==109 and external.subject.is_unique
assert len(inventory)==61 and inventory.arm.is_unique
assert len(qc)==34
assert content['new_fits']==numbers['new_fits']==numbers['new_checkpoint_inference']==0
assert snapshot['new_model_fits']==snapshot['new_checkpoint_inference']==0

checks={}
def record(path,digest):
    p=Path(path)
    if not p.is_absolute():p=ROOT/p
    p=p.resolve()
    assert p.is_file(),str(p)
    if str(p) in checks:assert checks[str(p)]==digest,str(p)
    checks[str(p)]=digest
for path,digest in snapshot['sources'].items():record(path,digest)
reference_detail_path=OUT/'evidence/reference_details.json'
record(reference_detail_path,hashlib.sha256(reference_detail_path.read_bytes()).hexdigest())
for path,digest in json.loads((OUT/'evidence/inventory_provenance.json').read_text())['input_sha256'].items():record(path,digest)
for name in ['q4_q11_review','independent_numbers']:
    review=json.loads((OUT/f'evidence/{name}.json').read_text())
    for path,value in review['input_artifacts'].items():record(value.get('path',path),value['sha256'])
review=json.loads((OUT/'evidence/q12_q14_review.json').read_text())
for item in review['key_evidence']:record(item['absolute_path'],item['sha256'])
methods=json.loads((OUT/'evidence/methods_and_references.json').read_text())
for name in ['source_code_inventory','validation_receipts']:
    for item in methods[name]:record(item['path'],item['sha256'])
for config_path in sorted((ROOT/'results/Q13-E004').glob('*/run_config.json')):
    config=json.loads(config_path.read_text())
    assert set(config['selected_epochs_by_target'].values())=={20}
    record(config_path,hashlib.sha256(config_path.read_bytes()).hexdigest())
for path,digest in checks.items():
    assert hashlib.sha256(Path(path).read_bytes()).hexdigest()==digest,path

blocks=content['blocks']
tables=[b for b in blocks if b['type']=='table']
figures=[b for b in blocks if b['type']=='figure']
assert len(tables)==7 and len(figures)==4 and len(content['references'])==11
assert all(r['verified'] for r in methods['references'])
doc=Document(OUT/'manuscript_en.docx')
assert len(doc.tables)==7 and len(doc.inline_shapes)==4
for expected,actual in zip(tables,doc.tables):
    assert len(actual.rows)==len(expected['rows'])+1
    assert len(actual.columns)==len(expected['headers'])
    for expected_row,actual_row in zip([expected['headers']]+expected['rows'],actual.rows):
        assert [c.text for c in actual_row.cells]==expected_row

pdf=fitz.open(OUT/'manuscript_en.pdf')
full_text='\n'.join(page.get_text() for page in pdf)
assert '\ufffd' not in full_text
for value in ['33.72','42.67','61.81','62.39','0.125','0.06446','200,000','20261002','S6.']:
    assert value in full_text,value
bounds=[];blank=[]
for n,page in enumerate(pdf,1):
    text=page.get_text('dict')
    spans=[s for b in text['blocks'] if 'lines' in b for line in b['lines'] for s in line['spans']]
    if len(page.get_text().strip())<100:blank.append(n)
    for span in spans:
        x0,y0,x1,y1=span['bbox']
        if x0<35 or x1>page.rect.width-35 or y0<25 or y1>page.rect.height-20:
            bounds.append({'page':n,'bbox':[x0,y0,x1,y1]})
assert not blank and not bounds,(blank,bounds)

markdown=(OUT/'manuscript_en.md').read_text()
assert '[-3.21,+1.42]' not in markdown and 'strictly monotone' not in markdown
assert 'Q15' not in markdown and 'Cho2017' not in markdown and 'Lee2019' not in markdown
assert 'strictly increasing (order-preserving)' in markdown
assert 'all counts use fixed 20 epochs' in markdown
assert '2,304 trials, 20 epochs, and 720 optimizer updates' in markdown
assert '200,000-draw reanalysis with seed 20261002' in markdown
tex=(OUT/'manuscript.tex').read_text()
assert r'\includegraphics' not in tex and r'\bibliography{' not in tex and r'\input{' not in tex
assert tex.count(r'\nextgroupplot')==7
assert tex.count(r'\begin{longtable}')==tex.count(r'\end{longtable}')==7
assert r'\fill[cell00]' in tex and r'\fill[cell58]' in tex
unescaped=re.sub(r'\\[{}%]', '', tex)
balance=0
for char in unescaped:
    if char=='{':balance+=1
    elif char=='}':balance-=1
    assert balance>=0
assert balance==0
changes=subprocess.check_output(['git','diff','--name-only','HEAD'],cwd=ROOT,text=True).splitlines()
assert not changes,changes

ci_source=pd.read_csv(ROOT/'results/Q10-V001/paired_contrasts.csv')
shared=ci_source.loc[ci_source.contrast.eq('Q9-E001/MU_BETA_SHARED minus Q8-E001')].iloc[0]
assert np.isclose(shared.subject_bootstrap_95ci_high*100,1.41460905349794)
q13=numbers['q13_matched_contrast'];q14=numbers['q14_primary']
ledger=[
 {'claim':'Mean-loss/rank mean four-class BA','values_percent':[100*numbers['q5_mean'],100*numbers['q8_mean']], 'inputs':['results/Q5-E001/subject_seed_metrics.csv','research_runs/Q8-E001/results/subject_seed_metrics.csv'],'status':'exploratory nine-person development'},
 {'claim':'Rank minus CE CI','interval_pp':numbers['q8_original_statistics']['bootstrap_mean_delta_95CI_pp'],'bootstrap_draws':20000,'seed':20260923,'input':'research_runs/Q8-E001/analysis/Q5_vs_Q8_paired_statistics.json','note':'Preserve archived nine-person sign sensitivity; conventional tie-excluding result separate'},
 {'claim':'Shared minus internal broad CI','interval_pp':[100*shared.subject_bootstrap_95ci_low,100*shared.subject_bootstrap_95ci_high],'bootstrap_draws':20000,'seed':20260924,'input':'results/Q10-V001/paired_contrasts.csv'},
 {'claim':'Fixed20 minus matched historical CE schedule','difference_pp':numbers['q13_fixed20_minus_ce_pp'],'interval_pp':[-100*q13['subject_bootstrap_95ci_high'],-100*q13['subject_bootstrap_95ci_low']],'sign_flip_p':q13['exact_sign_flip_p_exploratory'],'bootstrap_draws':10000,'seed':20260926,'input':'results/Q13-E006/postrun_statistics/paired_contrasts.csv'},
 {'claim':'External frozen binary primary','difference_pp':100*q14['mean_difference'],'interval_pp':[100*x for x in q14['subject_bootstrap_percentile_95_ci']],'sign_test_p':q14['two_sided_exact_sign_test_p_excluding_ties'],'bootstrap_draws':20000,'seed':20260924,'tie_rule':'Exact-zero differences excluded; numerical-tolerance sensitivity separately disclosed','input':'results/Q14-E002R2V1/validation_report.json','status':'Interval spans zero; no superiority established'},
 {'claim':'Source-count durations','fixed_epochs':20,'input':'results/Q13-E004/Q8_SRC2/run_config.json','note':'All k counts fixed20; two single-session conditions have equal2304-trial720-update budgets'},
]
(OUT/'evidence/claim_ledger.json').write_text(json.dumps(ledger,indent=2)+'\n')
report={'status':'delivery_document_and_evidence_checks_passed','fits_started':0,'new_checkpoint_inference':0,
 'source_files_rehashed':len(checks),'all_source_hashes_match':True,'tracked_source_changes':changes,
 'pdf_pages':len(pdf),'pdf_export_engine':'ReportLab','latex_compilation':'unverified; native editor and compiler timed out',
 'pdf_text_integrity':True,'pdf_text_bounds_checked':True,'blank_pages':blank,
 'visual_review':'All pages and four figures reviewed; after final text updates, changed figure and affected page tails rechecked',
 'docx_tables':len(doc.tables),'docx_figures':len(doc.inline_shapes),'docx_table_cells_match_shared_content':True,
 'references_verified':11,'tex_chart_panels':7,'tex_robustness_heatmap_cells':54,
 'tex_structural_checks':'Passing basic structure and standalone dependency checks; not a compiler confirmation',
 'inventory_rows':len(inventory),'early_binary_qc_rows':len(qc),'research_chat_history_read':False,
 'author_and_journal_details':'Pending; working author-review draft','external_publication_or_submission':False}
(OUT/'evidence/delivery_validation.json').write_text(json.dumps(report,indent=2)+'\n')
print(json.dumps(report,ensure_ascii=False))
