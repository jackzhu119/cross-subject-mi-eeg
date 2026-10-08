"""Create consistent English Markdown, DOCX, PDF and standalone LaTeX drafts.

Only archived results are read. No EEG model, training, or prediction is called.
The PDF export is a document-layout export, not evidence of LaTeX compilation.
"""
from pathlib import Path
import csv
import hashlib
import html
import json
import re
import pandas as pd

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether, PageBreak

OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[1]
NUM=json.loads((OUT/'paper_numbers.json').read_text())
REF=json.loads((OUT/'evidence/methods_and_references.json').read_text())['references']
for detail in json.loads((OUT/'evidence/reference_details.json').read_text()):
    reference=next(r for r in REF if r['key']==detail['key'])
    assert reference['doi']==detail['doi']
    for field in ['volume','issue']:
        if detail.get(field):reference[field]=detail[field]
ORDER=['Brunner2008','Tangermann2012','Lawhern2018','Ramoser2000','Ang2012','Ang2008','Schirrmeister2017','Sagawa2020','Schalk2009Dataset','Schalk2004','Goldberger2000']
REFMAP={r['key']:r for r in REF}
REFS=[REFMAP[k] for k in ORDER]
assert all(r['verified'] for r in REFS)
TITLE='Source-only model selection and shared spectral representations in cross-subject motor-imagery EEG decoding'
SUBTITLE='An exploratory internal study with a frozen binary external evaluation'
BLOCKS=[]
def heading(text,level=1):BLOCKS.append({'type':'heading','level':level,'text':text})
def para(text):BLOCKS.append({'type':'paragraph','text':text})
def table(label,headers,rows,note):BLOCKS.append({'type':'table','label':label,'headers':headers,'rows':rows,'note':note})
def figure(name,label,caption):BLOCKS.append({'type':'figure','name':name,'label':label,'caption':caption})
def equation(text,latex):BLOCKS.append({'type':'equation','text':text,'latex':latex})
def pct(x):return f'{100*x:.2f}'
def pp(x):return f'{x:+.2f}'

q5,q6,q8,shared=(NUM[k] for k in ('q5_mean','q6_mean','q8_mean','shared_mean'))
q14=NUM['q14_primary'];ci14=[100*x for x in q14['subject_bootstrap_percentile_95_ci']]
duration=NUM['q13_matched_contrast'];duration_ci=[-100*duration['subject_bootstrap_95ci_high'],-100*duration['subject_bootstrap_95ci_low']]
models={r['label']:r for r in NUM['four_class_models']}

heading('Abstract')
para(f'''Objective. Cross-subject motor-imagery EEG decoding requires training and model selection without calibration from the held-out person. We examined whether failures associated with source-only training-duration selection could be distinguished from the effects of spectral representations. Methods. Completed experiments were re-audited using archived protocols, trial predictions, source partitions, and experiment-specific validation records. Internal four-class analyses used nine BNCI2014_001 participants with both sessions held out per outer fold. We compared raw validation-loss selection with within-fold rank aggregation, fixed-duration and matched-runtime controls, spatial-spectral ablations, and source-only robustness interventions. A separate, frozen binary source model was evaluated on 109 PhysioNet participants. Results. The internal mean balanced accuracy increased from {pct(q5)}% under mean-loss selection to {pct(q8)}% under mean-rank selection; most of the aggregate improvement arose from two participants. In a matched-runtime control, fixed 20-epoch training exceeded the historical source-selected duration schedule by {NUM['q13_fixed20_minus_ce_pp']:.2f} percentage points (pp; bootstrap 95% interval [{duration_ci[0]:.2f}, {duration_ci[1]:.2f}], exact sign-flip p = {duration['exact_sign_flip_p_exploratory']:.3f}). Shared mu/beta input did not improve the internal mean over the rank-selected broad model. In external binary transfer, broad and shared models achieved {pct(NUM['q14_means']['BROAD_EEGNET_BA'])}% and {pct(NUM['q14_means']['MU_BETA_SHARED_BA'])}%; the primary paired difference was {100*q14['mean_difference']:+.3f} pp (95% interval [{ci14[0]:.3f}, {ci14[1]:.3f}], sign-test p = {q14['two_sided_exact_sign_test_p_excluding_ties']:.4f}). Conclusion. These exploratory results identify training-duration selection and participant heterogeneity as important audit targets. They do not establish general superiority of rank selection or shared spectral representations. External interpretation requires disclosure of a metadata-driven sampling-rate amendment.''')
para('Keywords: motor imagery; EEG; cross-subject decoding; EEGNet; source-only model selection; spectral parameter sharing; reproducibility.')

heading('1. Introduction')
para('Motor-imagery EEG provides a setting in which a decoder must distinguish imagined actions from weak, variable scalp signals. A calibration-free decoder should transfer from previously recorded people to a new person without fitting to that person’s labels or distributional statistics. This requirement differs from within-person evaluation and from adaptation procedures that use unlabeled target recordings. The BCI Competition IV dataset 2a remains a useful, small, well-documented benchmark for studying this distinction [1,2]. Its limited participant count also makes aggregate performance especially sensitive to a few people.')
para('Spatial filtering and frequency decomposition offer complementary inductive biases. Common spatial patterns (CSP) learn discriminative projections from class covariance structure [4], while filter-bank CSP extends that approach across frequency bands [5,6]. Compact convolutional models such as EEGNet combine temporal filtering, depthwise spatial filtering, and separable convolutions [3]. Broader CNN work illustrates the flexibility of learned EEG representations [7]. None of these architectural ideas removes the need to choose training duration using source participants whose validation curves may differ in scale, shape, and optimum epoch.')
para('A decoder that predicts predominantly one class can have balanced accuracy near chance even when the same architecture performs well with a different training duration. Such a failure can be mistaken for evidence that the input distribution requires normalization or a new representation. Conversely, a positive average after an intervention may reflect rescue of one or two participants rather than consistent improvement. Source-only fitting prevents per-fit target leakage, but it does not make a method independently confirmatory when the research program was developed after examining the same benchmark.')
para('We therefore treat model selection, representation, and robustness as distinct experimental questions. The first question is whether heterogeneous source-validation losses can lead to short training schedules and whether a rank-based aggregation changes those schedules. The second is whether fixed mu/beta views processed by one parameter-shared EEGNet provide an advantage beyond broad inputs, explicit spatial projections, or architectural controls. The third is whether those conclusions survive source-composition and matched-runtime checks. Finally, we evaluate a separately frozen binary model on an independent public cohort [9,10], while preserving its protocol-amendment history.')
para('The contribution is an evidence-grounded experimental audit, rather than a claim of a universally improved decoder. We report null and adverse conditions alongside positive point estimates; keep the held-out person as the inference unit; separate four-class development from binary external transfer; and distinguish archived checkpoint reconstruction from current prediction-table reanalysis. No new model was trained to prepare this manuscript.')

heading('2. Materials and methods')
heading('2.1. Datasets, populations, and task separation',2)
para('BNCI2014_001 corresponds to BCI Competition IV dataset 2a [1,2]. Nine participants completed two recording sessions, each containing six runs of 48 trials: 12 each of left-hand, right-hand, feet, and tongue imagery. Twenty-two EEG channels and three EOG channels were sampled at 250 Hz. Original acquisition used a left-mastoid reference, right-mastoid ground, 0.5–100 Hz band-pass, and 50-Hz notch; the offline operations below are additional processing of those recordings. The primary neural input uses EEG channels only. The archived metadata contain 5,184 trials, 576 per participant and 1,296 per class. All 488 expert artifact-flagged trials remain in the primary evaluation population. This include-all endpoint differs from the original competition’s artifact-free scoring; source-clean conditions exclude flagged source-training trials but retain every target trial.')
para('The external EEG Motor Movement/Imagery Database version 1.0.0 [9] supplies a separate binary task. We include all 109 subjects and only imagery runs 4, 8, and 12; T1 and T2 denote left- and right-fist imagery. Executed-movement, bilateral-task, and rest events do not enter this contrast. The final audited inventory contains 327 official-checksum-verified EDF files and 4,918 unique trials. The nominal 160-Hz acquisition description is supplemented by the actual headers: three subjects have 128-Hz recordings. Four-class and binary balanced accuracy have different chance levels and are never pooled into a common score.')
table('Table 1. Evaluated populations and computational endpoints.',
 ['Analysis','People','Classes','Unique trials','Primary tensor'],
 [['Internal BNCI','9','4','5,184','22 × 750 at 250 Hz'],['Binary BNCI source','9','2','2,592','22 × 480 at 160 Hz'],['External PhysioNet','109','2','4,918','22 × 480 at 160 Hz']],
 'Both sessions of each BNCI outer target are held out. The binary source models are separately fitted and are not a relabeling of the four-class checkpoints. The 34,426 external model/seed prediction rows represent seven saved predictions per unique trial, not independent observations.')
heading('2.2. Offline preprocessing and source-only partitions',2)
para('For internal BNCI analyses, fourth-order zero-phase Butterworth filters are applied within each native recording run before epoch extraction. The half-open interval [2.5,5.5) s after the MAT trial start corresponds to [0.5,3.5) s after the cue and yields 750 samples. There is no baseline subtraction, additional notch, or additional re-reference in the primary neural pipeline. MNE-loaded volt values are multiplied by 10^6 for network inputs in microvolts. These offline zero-phase operations use future samples and are not a causal online decoder.')
para('The broad mean-loss and mean-rank EEGNet conditions use 4–40 Hz. Restricted broad input uses 8–30 Hz, and the shared representation uses fixed mu (8–13 Hz) and beta (13–30 Hz) views. The bands are predefined experimental choices, not participant-specific estimates or physiological source separation. Any fitted scaler, CSP/PCA projection, or whitening transform is estimated on the applicable source-training partition only.')
para('Each primary four-class outer split holds out both sessions of one participant. The eight source people contribute 4,608 trials and the target contributes 576. Sorted source IDs define four consecutive, two-person inner-validation groups. Each inner fit trains on six people/3,456 trials and validates on two people/1,152 trials. This separation is by person, not by random trial. Final evaluation retains three fixed seeds, averaged within person before cross-person aggregation. Model evaluation freezes BatchNorm statistics; target examples do not update them.')
heading('2.3. Training-duration selection and diagnostic controls',2)
para('The mean-loss rule selects the earliest epoch minimizing the equally weighted mean validation cross-entropy across the four inner folds. Candidate epochs are 1–40. The mean-rank rule first ranks epochs within each fold by validation cross-entropy, assigns average ranks to ties, averages those ranks across folds, and selects the earliest minimum. Its operational purpose is invariance to strictly increasing (order-preserving) transformations within a fold; that property does not guarantee better target prediction. The rank follow-up reuses the 36 frozen original inner trajectories and trains 27 new final models after the schedule has been frozen.')
equation('Selected epoch = earliest argmin over e of (1/K) Σ_k rank_k[L_k(e)].',r'e^*=\min\operatorname*{arg\,min}_{e\in\{1,\ldots,40\}}\frac{1}{K}\sum_{k=1}^{K}\operatorname{rank}_{k}\{L_k(e)\},\quad K=4.')
para('The rank candidate was motivated by earlier results on the same BNCI participants, so the internal comparison remains exploratory despite source-only per-fit selection. A four-cell diagnostic for S3 crosses raw versus source-normalized input with two versus 16 fixed epochs. Here raw means filtered, microvolt-scaled EEG without fitted source normalization. Two cells reuse the archived original conditions and two are new archived diagnostic conditions; this is one participant, not a population-level replication. Source normalization estimates per-channel mean and scale from source-training samples only.')
para('Later controls compare fixed 20-epoch training with raw-loss-selected schedules. An additive matched-runtime experiment uses the same historical source-selected epoch counts but retrains the 27 final models on the same actual runtime and raw files as the fixed-duration comparator. It does not reselect epochs under the newer runtime. Source-count controls use k = 2, 4, or 6 source participants, four predefined identity subsets per count, and the k = 8 anchor; all counts use fixed 20 epochs. The 0train-versus-1test source-session contrast uses equal budgets: 2,304 trials, 20 epochs, and 720 optimizer updates per fit. Source-count contrasts and comparisons of one source session with the both-session anchor change the number of training examples and optimizer updates; those contrasts do not isolate a single causal factor.')
heading('2.4. Neural representation and spatial-spectral controls',2)
para('The four-class model is Braindecode 1.5.1 EEGNet with 22 input channels, 750 samples, F1 = 8, depth multiplier D = 2, F2 = 16, temporal kernel length 64 samples, and dropout 0.25. It has 2,932 trainable parameters in the archived implementation. Training uses Adam, learning rate 0.001, batch size 64, cross-entropy, zero weight decay, and no scheduler or default augmentation. The selection seed is 20260923 and the three final seeds are 20260924, 20260925, and 20260926. This is a configured EEGNet study, not an exact reproduction of the original paper’s preprocessing or splits [3].')
para('For shared mu/beta input, both band views pass through the same EEGNet in one concatenated mini-batch. Their class logits are averaged before cross-entropy in training and before softmax in inference. Network parameters and BatchNorm are shared; dropout acts per view. There are no learned band weights, attention mechanism, or separate independently trained branches. Parameter count matches one EEGNet, while forward-pass work and preprocessing differ.')
equation('Shared prediction: p = softmax[(fθ(Xμ) + fθ(Xβ))/2].',r'p(y\mid X)=\operatorname{softmax}\left\{\frac{f_{\theta}(X_{\mu})+f_{\theta}(X_{\beta})}{2}\right\}.')
para('Controls include individual frequency bands; fixed source-selected durations; Welch log-bandpower with source-fitted shallow classifiers; eight source-only CSP or PCA time-series projections per band; log-Euclidean covariance features; source-only normalization; and source-clean training. The CSP/PCA neural controls both use 16 projected channels and source-fitted standardization, so their width is matched to each other but differs from the 22-channel shared-logit model. Architectural ablations use two independent networks, early band stacking, four shared bands, and a broadband capacity comparator. The independent two-band model has 5,864 parameters and its comparator has 5,914, a disclosed 0.85% mismatch. The traditional filter-bank implementation is FBCSP-inspired, not a literal reproduction of Ang et al. [5,6].')
heading('2.5. Source-only robustness interventions',2)
para('The six completed robustness conditions comprise source-pooled whitening, source-balanced empirical risk minimization (ERM), source-group distributionally robust optimization (GroupDRO), channel dropout, gain perturbation, and their combination. Source whitening uses equal-person covariance aggregation, a ridge of 5% of the mean eigenvalue, and scale preservation; no target covariance or target alignment is fitted. GroupDRO treats source person as group [8]: persistent log weights increase by 0.05 times detached per-group cross-entropy, are normalized by log-sum-exp, and weight the group loss. Its declared matched comparator is balanced ERM, preserving the batching/sampling intervention. Training-only channel dropout uses probability 0.1 with surviving-channel scale 1/0.9; gain perturbation uses U[0.8,1.2], applied before dropout in the combined condition. Validation and target inputs are not augmented.')
para('Whitening and augmentation conditions are compared with the previously frozen broad rank-selected baseline. That historical baseline was executed on a different runtime from the robustness batch, so those cross-batch deltas combine intervention and runtime variation. GroupDRO versus balanced ERM is a cleaner within-runtime objective comparison. We retain this distinction beside the numerical table rather than attributing every delta to its named method.')
heading('2.6. Separate binary source freeze and external inference',2)
para('The binary BNCI source population contains 2,592 left/right trials, 288 per person. Filtering and epoching occur at 250 Hz before deterministic polyphase resampling by 16/25 to 160 Hz, giving 22 × 480 samples. The broad binary model uses 8–30 Hz; the shared model uses the fixed mu/beta views. All-nine-source duration selection uses the validation groups [1,2], [3,4], [5,6], and [7,8,9], with equal fold weighting. Frozen durations are 18 broad-model epochs and 17 shared-model epochs. Checkpoints, channels, labels, contrast, seeds, and preprocessing are frozen before external access. A source-only CSP4 + LDA comparator uses reg = None and LDA solver = svd; it does not inherit the internal CSP8 OAS/shrinkage-LDA configuration.')
para('A partial external run stopped at S88 when its 128-Hz header violated the original 160-Hz assumption. A versioned, metadata-driven amendment inspected remaining headers without using their predictions or scores. S88, S92, and S100 are filtered and epoched at their native rate, then resampled by 5/4 from 384 to 480 samples. The other 106 subjects retain native 160-Hz processing. The 87 completed earlier subject outputs are copied byte-for-byte and 22 additional subjects complete the cohort. No target statistics, fitted normalization, adaptation, seed selection, or class-performance-dependent branch is introduced. Because amendment followed partial external execution, the combined result is not a pristine execution of the original freeze.')
para('The original aggregate validation failed when re-serialized CSV probabilities were subjected to an inappropriate exact floating-point equality check. A later cross-host runtime mismatch was a separate limitation and motivated a versioned portability amendment. Both events remain in the record. The versioned validation-only attempt verified official EDF checksums and event identities and recomputed saved probability/argmax consistency, subject metrics, and statistics. It passed with zero new fits and zero new model-inference rows. This audit must not be described as regenerating all external predictions from checkpoints on the later validation host.')
heading('2.7. Endpoints, uncertainty, and evidence verification',2)
equation('Balanced accuracy = (1/C) Σ_c TP_c/(TP_c + FN_c).',r'\mathrm{BA}_{i,s}=\frac{1}{C}\sum_{c=1}^{C}\frac{\mathrm{TP}_{i,s,c}}{\mathrm{TP}_{i,s,c}+\mathrm{FN}_{i,s,c}},\qquad\overline{\mathrm{BA}}=\frac{1}{N}\sum_{i=1}^{N}\frac{1}{3}\sum_{s=1}^{3}\mathrm{BA}_{i,s}.')
para('The held-out person is the inference unit. Seeds are paired and averaged within person; predefined source-identity subsets are averaged within person and source count. Shallow conditions have one prediction per trial rather than three independent neural seeds. Paired differences are expressed in balanced-accuracy percentage points (pp). Person-bootstrap percentile intervals resample people, not trials, sessions, seeds, or source subsets. External intervals condition on the frozen source models and do not quantify source-cohort sampling uncertainty. Internal analyses remain exploratory because benchmark outcomes informed the research program and outer training sets overlap.')
para('We preserve archived experiment-specific statistics and identify each interval’s resampling procedure. Rank-selection intervals use 20,000 person-bootstrap draws with seed 20260923. The archived shared-minus-broad representation contrast and robustness intervals use 20,000 draws with seed 20260924. Source-composition and matched-runtime intervals use 10,000 draws with seed 20260926 and exhaustive two-sided sign flips over nine participant differences; familywise exploratory p-values use Holm correction within the frozen families. The external primary contrast uses 20,000 draws with seed 20260924 and an exact two-sided binomial sign test excluding exact zero differences. Its primary hypothesis is shared-minus-broad; contextual comparisons with CSP do not replace it. A supplementary numerical audit uses 200,000 draws with seed 20261002 and is labelled separately. We do not equate an interval excluding zero with agreement across all small-sample tests.')
para('One historical rank-selection script labelled a nine-person binomial calculation as a sign test while retaining two unchanged participants in its denominator. We describe its p = 0.1797 as an archived nine-person sign sensitivity. A conventional tie-excluding calculation uses seven nonzero changes and gives p = 0.015625; the paired t-test gives p = 0.111. This distinction is reported transparently rather than selecting the favorable procedure or rewriting the frozen result. Multiplicity corrections from later declared families are not retrospectively imposed on unrelated earlier statistics.')
para('For this manuscript, archived prediction tables were independently reaggregated by participant and seed, with checks for missing/duplicate trial identities and agreement with reported balanced accuracy. Available historical scientific validators provide deeper checkpoint/raw-data reconstruction for specific batches; the early spectral batch’s top-level validator establishes orchestration only. Validation depth is therefore recorded per experiment. Some text prediction hashes correspond to archived CRLF line endings while the cloud checkout uses LF; those cases are documented as exact newline-conversion matches, not asserted to be identical original bytes. No checkpoint inference or new fit was performed in manuscript preparation.')

heading('3. Results')
heading('3.1. Traditional baselines and heterogeneous duration selection',2)
para(f'The all-trial four-class broad CSP8 + LDA baseline achieved 39.04% balanced accuracy; broad CSP + SVM and the full filter-bank LDA comparator achieved 37.50% and 38.12%. Raw mean-loss-selected EEGNet achieved {pct(q5)}%, below the broad CSP mean. These descriptive comparisons do not establish superiority of a model class. The mean-loss schedule selected only one or two epochs for S2, S3, and S8, while source-validation optima differed substantially across folds.')
para(f'Source-only per-channel normalization increased the mean to {pct(q6)}%, a {100*(q6-q5):+.2f} pp difference. The archived paired 95% interval spans −1.12 to +13.49 pp and the median change is only +0.23 pp. The intervention has heterogeneous effects, including negative transfer in S8. It does not establish a general normalization benefit.')
para('The S3 diagnostic separates normalization from duration. Raw input at two epochs gave 26.10% balanced accuracy and source normalization at two epochs gave 25.69%. At 16 epochs, raw and normalized inputs gave 67.71% and 65.86%. The short-run raw condition placed approximately 94% of predictions in one class. Duration can reproduce the observed rescue for this participant; this four-cell diagnostic does not prove that duration explains all cross-person failures.')
para(f'Mean-rank selection increased equal-person mean balanced accuracy to {pct(q8)}%, a {NUM["q8_delta_pp"]:+.2f} pp change over raw mean-loss selection. The archived bootstrap interval is [+1.05,+19.96] pp and the median change is {NUM["q8_median_delta_pp"]:+.2f} pp. Seven participants improved and two were unchanged. S3 and S8 account for 87.57% of the summed improvement. Thus the positive average includes a marked concentration of benefit, and the prior same-dataset development and disagreement among small-sample statistics limit a general claim.')
figure('figure1_selection_and_s3','Figure 1','Training-duration selection and participant-level internal results. A: archived source-selected epochs. B: each participant’s all-trial balanced accuracy after averaging three fixed seeds; the dotted line is four-class chance (25%). C: the S3 four-cell diagnostic; error bars are standard deviations across three seeds, not uncertainty across people. Raw and source-normalized input are compared at two fixed durations. The diagnostic involves one person.')
core=[['Broad CSP8 + LDA','4–40','39.04','Single source fit/fold'],['Filter-bank CSP + LDA','4–40 bank','38.12','Fixed filter-bank control'],['EEGNet, mean loss','4–40',pct(q5),'Source-selected duration'],['EEGNet, source norm','4–40',pct(q6),'Source-only fitted scale'],['EEGNet, mean rank','4–40',pct(q8),'Reused source curves'],['Shared mu/beta','8–13 / 13–30',pct(shared),'Representation-specific rank'],['Broad, fixed 20','4–40',pct(NUM['q13_fixed20_mean']),'Later runtime'],['Broad, matched CE schedule','4–40',pct(NUM['q13_matched_ce_mean']),'Same runtime as fixed 20']]
table('Table 2. Selected completed four-class conditions; all nine participants.', ['Condition','Bands (Hz)','Mean BA (%)','Interpretation'],core,'All-trial primary population. Neural means average three fixed seeds within person first. Rows span different batches; only explicitly matched comparisons isolate runtime variation. These rows summarize declared controls, not a target-selected ranking. Complete representation and robustness inventories are supplied in the supplement.')
heading('3.2. Matched-runtime duration and source-composition controls',2)
para(f'In the additive matched-runtime experiment, retraining the historical raw-loss duration schedule gave {pct(NUM["q13_matched_ce_mean"])}% balanced accuracy versus {pct(NUM["q13_fixed20_mean"])}% for fixed 20 epochs. The fixed-minus-historical-schedule mean was {NUM["q13_fixed20_minus_ce_pp"]:+.2f} pp, with bootstrap 95% interval [{duration_ci[0]:+.2f},{duration_ci[1]:+.2f}] pp. The exact sign-flip p-value was {duration["exact_sign_flip_p_exploratory"]:.3f}. The schedule comparison persists on a matched runtime, but its nine-person inferential evidence remains sensitive to the test; it is not proof of a universally optimal fixed duration or of a newly selected schedule under that runtime.')
para('Mean balanced accuracy increased descriptively with source count. Broad fixed-duration models gave 32.96%, 38.89%, 43.16%, and 43.99% at k = 2, 4, 6, and 8; shared models gave 34.74%, 38.89%, 42.48%, and 43.98%. Source-identity variation within each count remained substantial. The two-source broad-minus-eight-source difference was −11.03 pp, but its Holm-adjusted exploratory p was 0.164. Every source-count/selection/session family retained null corrected findings. One-session-versus-the-other contrasts were small and heterogeneous. These trajectories are consistent with training-information sensitivity and do not separate source diversity from sample count or optimizer-update budget.')
figure('figure2_source_count_and_runtime','Figure 2','Source composition and matched-runtime controls. A: strong lines are equal-person means; faint lines are individual target-person trajectories after averaging the four predefined identity subsets and three seeds at k = 2, 4, 6. All counts use 20 epochs; k = 8 reuses the fixed-20 anchor. B: all nine fixed-20-minus-historical-CE-schedule effects on the matched runtime. The interval is the sign-reversed archived matched contrast. Source count changes data volume and optimizer updates; these are exploratory sensitivities.')
heading('3.3. Spectral, spatial, and architectural interventions',2)
para(f'The internal shared mu/beta condition achieved {pct(shared)}% versus {pct(q8)}% for the broad mean-rank baseline, a −0.48 pp mean difference. The archived Q10-V001 exploratory paired interval is [−3.21,+1.41] pp (20,000 draws, seed 20260924), with exact sign-flip p = 0.914. There is no demonstrated primary spectral-sharing gain. A restricted broad band and the individual bands provide descriptive frequency controls, while fixed-duration representation controls preserve their declared schedules.')
para(f'The source-only CSP8 and PCA8 neural projections achieved {pct(models["Q10 CSP8 + EEGNet"]["mean_ba"])}% and {pct(models["Q10 PCA8 + EEGNet"]["mean_ba"])}%. The architectural controls achieved {pct(models["Q11 independent two-band"]["mean_ba"])}% for independent two-band networks, {pct(models["Q11 capacity-matched broad"]["mean_ba"])}% for the broadband capacity comparator, {pct(models["Q11 early stack"]["mean_ba"])}% for early stacking, and {pct(models["Q11 four-band shared"]["mean_ba"])}% for four-band sharing. These completed alternatives do not supply evidence for a general shared-band advantage. Their differences in preprocessing, projected width, BatchNorm behavior, and compute prevent a single-factor physiological interpretation. Complete model means and participant-level values accompany the manuscript.')
heading('3.4. Source-only robustness has no corrected established gain',2)
para('Source-pooled whitening, source-balanced ERM, and gain perturbation had positive point estimates relative to the historical broad rank baseline, whereas dropout conditions were lower. None established a corrected improvement. The matched GroupDRO-minus-balanced-ERM difference was −6.45 pp (unadjusted bootstrap interval [−11.47,−1.79] pp; exact sign-flip p = 0.046875, Holm-adjusted p = 0.140625). This adverse estimate is a follow-up lead, not evidence that distributionally robust optimization generally harms EEG decoding. Whitening and gain perturbation intervals cross zero; their cross-runtime historical comparator further limits attribution.')
robustrows=[['Whitening','44.34','Q8 historical','+1.67','[−0.39,+3.67]','0.3594'],['Balanced ERM','43.18','Q8 historical','+0.51','[−0.71,+1.77]','0.4766'],['GroupDRO','36.73','Balanced ERM','−6.45','[−11.47,−1.79]','0.1406'],['Channel dropout','41.00','Q8 historical','−1.67','[−4.51,+0.64]','0.7617'],['Gain perturbation','43.42','Q8 historical','+0.75','[−0.34,+1.86]','0.7617'],['Dropout + gain','40.92','Q8 historical','−1.75','[−5.05,+0.91]','0.7617']]
table('Table 3. Exploratory source-only robustness (nine participants).',['Condition','BA (%)','Control','Δ (pp)','95% CI (pp)','Holm p'],robustrows,'Intervals are the archived, unadjusted person-bootstrap intervals; p-values use the declared within-family Holm correction. All comparisons with Q8 cross runtimes. GroupDRO versus balanced ERM is the matched-batching, within-runtime contrast. No interval/p-value is treated as a target-selection rule.')
heading('3.5. Frozen binary external transfer has a small uncertain primary effect',2)
para('In binary BNCI development, broad, shared, and CSP models achieved 68.21%, 67.35%, and 61.54% balanced accuracy. These development outcomes cannot be pooled with the four-class analyses. The separately frozen all-nine-source models then provided predictions for all 109 external subjects. The final validation confirms 4,918 trial identities and 34,426 saved prediction rows. No target fit was made.')
para(f'External equal-person balanced accuracy was {pct(NUM["q14_means"]["BROAD_EEGNET_BA"])}% for broad EEGNet, {pct(NUM["q14_means"]["MU_BETA_SHARED_BA"])}% for shared mu/beta, and {pct(NUM["q14_means"]["CSP4_LDA_BA"])}% for CSP4 + LDA. The predeclared shared-minus-broad primary difference was {100*q14["mean_difference"]:+.3f} pp, with 95% bootstrap interval [{ci14[0]:+.3f},{ci14[1]:+.3f}] pp. Sixty-three people favored shared, 43 favored broad, and three tied; the exact tie-excluding sign-test p was {q14["two_sided_exact_sign_test_p_excluding_ties"]:.5f}. Neither the interval nor this test establishes external superiority.')
para('Individual performance varied widely. The shared model ranged from 39.33% to 97.07% balanced accuracy, with median 58.73%. Three-seed average differences also varied in direction across people. The three resampled subjects had mean shared-minus-broad Δ = +0.514 pp, versus +0.575 pp for the 106 native-160-Hz subjects. The small three-person subgroup is descriptive, not a powered subgroup test. This result is reported with the sampling-rate and validation-portability amendments, rather than as an unmodified original external protocol.')
table('Table 4. Separate binary external evaluation (109 participants).',['Frozen model','Mean BA (%)','Subject SD (pp)','Median BA (%)'],[['Broad EEGNet','61.81','12.56','58.23'],['Shared mu/beta','62.39','12.92','58.73'],['CSP4 + LDA','54.46','7.90','52.08']],'Deep-model seeds are averaged within each person. The primary contrast is shared minus broad: +0.573 pp, 95% CI [−0.097,+1.242], p = 0.06446. CSP comparisons are contextual and do not replace the primary hypothesis. Sampling-rate amendment and versioned portability validation are disclosed in Methods.')
figure('figure3_external_primary','Figure 3','Independent-cohort binary transfer. A: empirical cumulative distributions of three-seed subject mean balanced accuracy (one CSP value per person); the vertical dotted line denotes binary chance (50%). B: all 109 primary paired effects, with zero and the mean indicated. The interval and p-value are the frozen person-bootstrap and tie-excluding sign-test outputs. The cohort contains three deterministically resampled subjects under the disclosed metadata amendment.')

heading('4. Discussion')
heading('4.1. Model selection is part of the generalization problem',2)
para('The internal experiments indicate that a source-only decoder can fail through its training-duration decision even before a representation intervention is considered. The S3 fixed-duration diagnostic reproduced a large recovery with raw input and did not reproduce it through normalization at the shorter duration. Rank aggregation changed source-selected epochs and rescued S3/S8 in the archived follow-up. A matched-runtime control retained a substantial average penalty for the historical short schedule. Together these findings motivate inspecting source-validation curves, selected duration, class recall, and prediction concentration whenever a cross-person model appears to collapse.')
para('These results support a specific diagnostic interpretation, not a universal duration prescription. Rank aggregation discards fold-specific loss scale but also discards magnitude information and weights every fold equally. The external study did not independently compare mean-rank with mean-loss duration selection, so it is not an external confirmation of rank selection. A prospective duration comparison on new participants, with runtime matched and training-update budgets controlled, is required to establish that contribution beyond this research program.')
heading('4.2. Spectral sharing is a constrained representation, not established superiority',2)
para('Fixed mu/beta views and shared parameters are plausible ways to constrain a compact decoder. However, the internal primary representation comparison was slightly negative and the independent binary contrast was small and uncertain. Source-fitted spatial projections and architectural controls did not create a coherent superiority pattern. This evidence does not justify a claim that the method identifies invariant rhythms, separates physiological sources, or solves subject shift. Nor does a null comparison prove equivalence; the intervals and sample sizes leave a range of small effects compatible with the data.')
para('The shared model’s parameter economy must also be distinguished from compute economy. It processes two band views and shares their BatchNorm statistics, whereas an independent-branch model has different capacity and normalization behavior. The near capacity-matched broad comparator has a small disclosed parameter mismatch. A future controlled study should measure runtime and isolate these design factors rather than crediting every difference to band physiology. We retain the implemented descriptive name “shared mu/beta EEGNet” instead of labelling it a novel attention or adaptation architecture.')
heading('4.3. Participant heterogeneity and statistical limits',2)
para('The same average can conceal near-chance people, high-performing people, adverse changes, and large rescues. The concentration of the rank-selection gain in S3/S8 and the directional variability of external differences make person-level tables essential. Multiple seeds help characterize stochastic variation but do not increase the number of independent people. Source subsets and sessions are also repeated measurements within the same target person. Outer LOSO training sets overlap, adding dependence beyond a simple independent-participant model.')
para('Bootstrap intervals, sign tests, sign flips, and paired t-tests answer related questions under different assumptions. With nine skewed differences they need not agree. We report those discrepancies and the archived tie-denominator convention explicitly. The broad internal program contains many hypotheses, so internal unadjusted findings are exploratory; the later frozen families retain their Holm-corrected null outcomes. The independent external primary result is stronger in population size and source-freeze design, but its interval includes zero and its metadata amendment reduces the claim of pristine confirmation.')
heading('4.4. Reproducibility and scope of external evidence',2)
para('The public record retains configurations, source identities, checkpoints where published, trial predictions, transform receipts, statistics, technical failures, and versioned amendments. A successful schedule/checkpoint bookkeeping check is distinct from scientific replay. Likewise, a later audit of saved probabilities and official event identities is distinct from regenerating predictions on a different host. Preserving those distinctions and newline-custody notes avoids overstating reproducibility while making the available evidence inspectable.')
para('The external cohort uses different acquisition hardware, references, channels selected from a broader montage, task labels, and a binary decision. Its source-to-external performance change cannot be attributed to one isolated shift mechanism. All pipelines remain offline; no clinical population, assistive-control task, or online BCI was tested. We make no claim of clinical benefit or universal calibration-free reliability.')

heading('5. Conclusion')
para('Completed source-only experiments show that training-duration decisions can reproduce substantial cross-subject failure and rescue in a small, previously explored benchmark. Person-level and matched-runtime analyses qualify the aggregate gains. Fixed shared mu/beta representations, spatial controls, and source-only robustness interventions do not establish a general advantage. A separately frozen binary external evaluation yields a small positive shared-minus-broad point estimate with an interval spanning zero. The resulting evidence favors transparent selection diagnostics, complete negative-result reporting, and prospective participant-level validation over a claim of a finished universally superior spatial-spectral decoder.')

heading('Data, code, ethics, and author declarations')
para('The source datasets are public resources [1,2,9], with PhysioNet providing the external distribution platform [11]; their original licenses and usage terms apply. Code, archived experiment records, and published results are available at https://github.com/jackzhu119/cross-subject-mi-eeg. This draft reviewed GitHub main at commit adb2d406b4300e2c8e4d5112969291c0331f3ed5; the manuscript evidence bundle records local input hashes, audit scope, and reproducible table/figure scripts. Raw EEG is not redistributed in the manuscript bundle. Some large checkpoint files may be omitted by recorded publication policies; availability must be verified before submission.')
para('This work analyzes existing public recordings and recruited no new participants. Ethical approval and consent for original collection are governed by the dataset creators’ reports; the authors must confirm any institutional requirements for secondary analysis. Author names, affiliations, corresponding author, funding, competing interests, contributions, and any required AI-assistance disclosure remain to be supplied or confirmed by the authors. No approval number, consent statement, funding award, or competing-interest declaration is invented here.')

heading('Supplementary material')
heading('S1. Full internal model inventory and participant effects',2)
para('Table S1 reports all completed nine-person Q4–Q11 four-class conditions, including shallow spectral/geometric controls and source-clean/source-normalized variants. It is descriptive: higher target mean does not select a method for a confirmatory claim. Individual participant values, fixed seeds, primary all-trial populations, and source file hashes are available in the accompanying CSV/JSON evidence. The complete inventory includes 61 Q4–Q14 condition records, with repeated endpoints and the four single-person diagnostic cells marked; these rows are not independent replications or a cumulative fit count. Robustness conditions are reported in Table 3, source-composition means in the complete CSV, and their statistical families in Table S2.')
inventory=pd.read_csv(OUT/'tables/completed_internal_inventory.csv')
early=inventory.loc[inventory.arm.str.match(r'Q(?:4|5|6|8|9|10|11)-') & inventory.n_subjects.eq(9)]
table('Table S1. Completed nine-person Q4–Q11 four-class conditions.',['Archived condition','Mean BA (%)','Subject SD (pp)'],[[r.arm.replace('_',' '),pct(r.mean_ba),f'{100*r.subject_sd_ba:.2f}'] for r in early.itertuples()], 'Neural seeds are averaged within person; shallow arms have one evaluated model per fold. Q4-A001 reuses fitted CSP feature caches but refits scaler/LDA; its k8/k72 predictions reproduce prior Q4 endpoints and do not add independent outcomes. Archived condition IDs identify protocol order, not a sorted leaderboard. The four S3 diagnostic cells are shown in Figure 1C, not included as nine-person rows.')
heading('S2. Robustness heterogeneity',2)
figure('figureS1_robustness_heterogeneity','Figure S1','All nine participant-level robustness differences in fixed order. Color and annotations are balanced-accuracy percentage points. The GroupDRO row uses balanced ERM as control; every other row uses the historical Q8 rank baseline and therefore crosses runtime. This is descriptive, without target-based exclusion or significance stars.')
heading('S3. Duration, source-count, and source-session statistical families',2)
import pandas as pd
q13table=pd.read_csv(OUT/'tables/q13_contrasts.csv')
table('Table S2. Archived exploratory Q13/E006 paired contrasts.',['Contrast','Δ (pp)','95% CI (pp)','Holm p'],[[str(r.contrast).replace('Q9_MU_BETA_SHARED','Shared').replace('Q9_SHARED','Shared').replace('Q8_E006_RAW_CE','Matched CE').replace('Q8_RAW_CE_REUSE_Q5','Historical CE').replace('Q8_BROAD','Broad').replace('Q8_FIXED20','Fixed20').replace('_',' '),f'{100*r.mean_paired_ba_difference:+.2f}',f'[{100*r.subject_bootstrap_95ci_low:+.2f},{100*r.subject_bootstrap_95ci_high:+.2f}]',f'{r.holm_within_family_p_exploratory:.4f}'] for r in q13table.itertuples()], 'Each row uses nine people after within-person averaging of fixed seeds and, where applicable, source subsets. The historical CE contrast crosses runtimes; the separately declared matched-CE-minus-fixed20 row matches them. Intervals are unadjusted, while p-values are corrected within each declared family. These signs follow the archived contrasts; Figure 2B reverses the matched contrast for readability.')
heading('S4. Audit and reporting boundaries',2)
para('The manuscript-number audit independently reaggregates stored predictions and records its actual CSV hashes. Historical source scientific replay counts are reported only where a passing scientific receipt supports them. The Q9 top-level batch report alone is orchestration-only. The external versioned portability validator adds no new inference. CRLF-to-LF matches are labelled explicitly. The paper bundle includes the source-snapshot manifest, per-review input hashes, an independent numerical audit, generated tables, figure sources, and draft-render checks. None of these files is a new training authorization or an external outcome-selected method freeze.')
para('The original rank-follow-up archive reports a bootstrap interval [+1.048,+19.959] pp (20,000 draws, seed 20260923); the independent 200,000-draw reanalysis with seed 20261002 gives [+1.067,+19.952] pp. The main text preserves the original interval. The saved nine-person sign sensitivity retains zero changes in its denominator; the conventional tie-excluding sensitivity is identified separately. Disclosing both implementation and resampling differences prevents an apparent numerical disagreement from being silently concealed.')
para('The external primary sign test preserves the archived exact-zero rule. S10 has a saved shared-minus-broad difference of approximately −1.11 × 10^−16 and therefore counts as negative: 63 positive, 43 negative, and three zero differences, p = 0.06446. A numerical-tolerance sensitivity at 10^−14 classifies that difference as a tie, giving 63 positive, 42 negative, four ties, and p = 0.05044. The primary rule is unchanged; neither calculation establishes superiority at the conventional 0.05 threshold.')
heading('S5. Earlier binary calibration, artifact, and EOG checks',2)
para('Earlier pipeline checks use only left/right imagery and must remain separate from the four-class analysis. P2-E001 contains 2,346 expert-clean trials. Its within-session leave-run-out and cross-session models use labeled trials from the evaluated person; only its cross-subject LOSO excludes both target sessions. Within-session/LOSO each evaluate all 2,346 trials, while cross-session evaluates 1,183 later-session trials. The single-person P2-SMOKE-S1B is a technical smoke, not a population replication. The all-trial P2-E002-ALLTRIALS condition changes both source and target populations to include 2,592 trials; some retained metadata uses an earlier clean-name field, so effective populations, rather than those labels, determine interpretation.')
para('P3-E001 compares CSP4, Welch PSD88, and their 92-feature fusion under source-all versus source-expert-clean training. The same 2,346 expert-clean target trials define its primary endpoint; all 2,592 and the 246 flagged target trials are secondary strata. P4-E001B is the successful EOG-processing retry: regression coefficients are fitted on clean source data, then applied to three synchronous target EOG channels. It uses no target parameter fit but requires those additional sensors. Its clean-target balanced accuracy is 62.57% versus 62.85% uncorrected, so the result does not establish a decoding benefit or selective ocular-artifact removal. The failed original P4 attempt remains in the technical record.')
qc=pd.read_csv(OUT/'tables/early_binary_qc_inventory.csv')
qcprimary=qc.loc[qc.primary_endpoint.eq(True)&qc.n_subjects.eq(9)]
def qc_label(r):
    label=str(r.model).replace('WelchPSD88','PSD88').replace('+shrinkageLDA',' + LDA')
    if r.experiment=='P4-E001B':label='EOG regression' if r.condition=='source_train_fitted_EOG_regression' else 'Uncorrected CSP4 + LDA'
    return f'{r.experiment} / {label}'
table('Table S3. Earlier binary QC primary endpoints; all nine participants.',['Stage / model','Split / source policy','Target trials','BA (%)'],[[qc_label(r),str(r.mode).replace('_',' ')+'; '+str(r.condition).replace('_',' '),str(r.n_test_trials),pct(r.mean_ba)] for r in qcprimary.itertuples()], 'All rows are binary, with 50% chance. Within-session and cross-session rows use target-person labeled calibration. P3/P4 rows use expert-clean targets; P2 all-trial rows use all trials. CSV supplements retain all/flagged secondary endpoints, target-information notes, and the single-person smoke. These modes and populations are not pooled into a calibration-free score.')
heading('S6. Selection-curve analyses and technical history',2)
para('Completed Q8-A001/A002/A003 analyses inspect archived source-validation curves without additional model fitting. Leave-one-inner-fold-out sensitivity and six candidate aggregation rules are hypothesis-generating analyses; they do not constitute six target-evaluated decoder conditions. P1-E001 is a single-subject acquisition audit. P2-SMOKE-S1 and the original P4-E001 retain failed run receipts, followed by successful retries. Q13-E002/E003 and Q14-E003 remain unrun and receive no performance values. The external partial 87-person attempt is not counted as an additional cohort alongside the completed 109-person result. The separate completed-experiment inventory links each record to its evidential role and explicitly marks reuse, validation depth, and technical failure.')

def ref_text(r):
    authors='; '.join(r['authors'])
    venue=r.get('journal',r.get('venue',r.get('publisher',r.get('conference',r.get('type','')))))
    volume=str(r.get('volume',''))
    if r.get('issue'):volume+='('+str(r['issue'])+')'
    pages=r.get('pages_or_article',r.get('pages',''))
    details=', '.join(x for x in [volume,pages] if x)
    if r.get('version'):details='version '+str(r['version'])
    location=venue+(', '+details if details else '')
    return f'{authors} ({r["year"]}). {r["title"]}. {location}. {r.get("url","")}'
REFERENCES=[ref_text(r) for r in REFS]
(OUT/'manuscript_content.json').write_text(json.dumps({'title':TITLE,'subtitle':SUBTITLE,'blocks':BLOCKS,'references':REFERENCES,'new_fits':0},indent=2,ensure_ascii=False)+'\n')

def markdown():
    lines=[f'# {TITLE}','',f'*{SUBTITLE}*','','Working draft for author review — 2 October 2026. Author details pending.','']
    for b in BLOCKS:
        kind=b['type']
        if kind=='heading':lines.extend(['#'*(b['level']+1)+' '+b['text'],''])
        elif kind in ('paragraph','equation'):lines.extend([b['text'],''])
        elif kind=='table':
            lines.extend(['**'+b['label']+'**','','| '+' | '.join(b['headers'])+' |','| '+' | '.join(['---']*len(b['headers']))+' |'])
            lines.extend('| '+' | '.join(row)+' |' for row in b['rows'])
            lines.extend(['',b['note'],''])
        elif kind=='figure':lines.extend([f'![{b["label"]}](figures/{b["name"]}.png)','',f'**{b["label"]}.** {b["caption"]}',''])
    lines.extend(['## References',''])
    lines.extend(f'{i}. {r}' for i,r in enumerate(REFERENCES,1))
    (OUT/'manuscript_en.md').write_text('\n'.join(lines)+'\n')

def word():
    doc=Document();s=doc.sections[0];s.top_margin=s.bottom_margin=Inches(.8)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(7)
    doc.add_heading(TITLE,0);doc.add_paragraph(SUBTITLE,'Subtitle')
    doc.add_paragraph('Working draft for author review — 2 October 2026. Author details pending.')
    for b in BLOCKS:
        if b['type']=='heading':doc.add_heading(b['text'],b['level'])
        elif b['type'] in ('paragraph','equation'):doc.add_paragraph(b['text'])
        elif b['type']=='table':
            p=doc.add_paragraph();p.add_run(b['label']).bold=True
            t=doc.add_table(rows=1,cols=len(b['headers']));t.style='Light Shading Accent 1'
            for cell,text in zip(t.rows[0].cells,b['headers']):cell.text=text
            for row in b['rows']:
                for cell,text in zip(t.add_row().cells,row):cell.text=text
            doc.add_paragraph(b['note'])
        elif b['type']=='figure':
            doc.add_picture(str(OUT/'figures'/f'{b["name"]}.png'),width=Inches(6.3))
            p=doc.add_paragraph();p.add_run(b['label']+'. ').bold=True;p.add_run(b['caption'])
    doc.add_heading('References',1)
    for i,r in enumerate(REFERENCES,1):doc.add_paragraph(f'[{i}] {r}')
    footer=s.footer.paragraphs[0];footer.text='Working manuscript • Source-only MI EEG • 2 October 2026'
    doc.save(OUT/'manuscript_en.docx')

def pdf():
    base=Path('/usr/share/fonts/truetype/dejavu')
    for name,file in [('Paper','DejaVuSerif.ttf'),('PaperBold','DejaVuSerif-Bold.ttf'),('PaperItalic','DejaVuSerif-Italic.ttf'),('Caption','DejaVuSans.ttf')]:
        pdfmetrics.registerFont(TTFont(name,str(base/file)))
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='PaperBody',fontName='Paper',fontSize=9.8,leading=14.1,alignment=TA_JUSTIFY,spaceAfter=7))
    styles.add(ParagraphStyle(name='PaperTitle',fontName='PaperBold',fontSize=19,leading=24,alignment=TA_CENTER,spaceAfter=10))
    styles.add(ParagraphStyle(name='PaperSubtitle',fontName='PaperItalic',fontSize=10,leading=14,alignment=TA_CENTER,spaceAfter=12))
    styles.add(ParagraphStyle(name='PaperH1',fontName='PaperBold',fontSize=13,leading=17,spaceBefore=13,spaceAfter=7,keepWithNext=True))
    styles.add(ParagraphStyle(name='PaperH2',fontName='PaperBold',fontSize=10.5,leading=15,spaceBefore=9,spaceAfter=5,keepWithNext=True))
    styles.add(ParagraphStyle(name='PaperTableTitle',fontName='PaperBold',fontSize=10.5,leading=15,spaceBefore=9,spaceAfter=5))
    styles.add(ParagraphStyle(name='PaperCaption',fontName='Caption',fontSize=8,leading=11,spaceAfter=9))
    styles.add(ParagraphStyle(name='PaperCell',fontName='Caption',fontSize=7.5,leading=10))
    story=[]
    def p(text,style='PaperBody'):return Paragraph(html.escape(text),styles[style])
    story.extend([p(TITLE,'PaperTitle'),p(SUBTITLE,'PaperSubtitle'),p('Working draft for author review • 2 October 2026 • Author details pending','PaperCaption')])
    for b in BLOCKS:
        k=b['type']
        if k=='heading':
            if b['text']=='Supplementary material':story.append(PageBreak())
            story.append(p(b['text'],'PaperH1' if b['level']==1 else 'PaperH2'))
        elif k in ('paragraph','equation'):story.append(p(b['text']))
        elif k=='table':
            n=len(b['headers']);width=6.65*inch
            if n==6:weights=[1.45,.62,1.05,.65,1.6,.6]
            elif n==5:weights=[1.3,.55,.55,.95,1.4]
            elif n==4:weights=[2.3,1.9,.7,.6] if 'S3.' in b['label'] else ([1.9,.8,1.8,1.4] if 'Table 2.' in b['label'] else [2.9,1,1.45,1])
            else:weights=[3.3,1.1,1.3]
            colwidth=[width*w/sum(weights) for w in weights]
            data=[[p(t,'PaperCell') for t in row] for row in [b['headers']]+b['rows']]
            t=Table(data,colWidths=colwidth,repeatRows=1,hAlign='LEFT')
            t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e7edf4')),('LINEBELOW',(0,0),(-1,0),.7,colors.HexColor('#607080')),('LINEBELOW',(0,-1),(-1,-1),.6,colors.HexColor('#607080')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f7f9fb')])]))
            story.extend([p(b['label'],'PaperTableTitle'),t,Spacer(1,5),p(b['note'],'PaperCaption')])
        elif k=='figure':
            image=Image(str(OUT/'figures'/f'{b["name"]}.png'));ratio=image.imageHeight/image.imageWidth
            image.drawWidth=6.65*inch;image.drawHeight=image.drawWidth*ratio
            story.append(KeepTogether([image,Spacer(1,5),p(b['label']+'. '+b['caption'],'PaperCaption')]))
    story.append(p('References','PaperH1'))
    for i,r in enumerate(REFERENCES,1):story.append(KeepTogether([p(f'[{i}] {r}','PaperCaption')]))
    def footer(c,doc):
        c.saveState();c.setFont('Caption',7);c.setFillColor(colors.HexColor('#607080'))
        c.drawString(.82*inch,.42*inch,'Source-only MI EEG • Working manuscript');c.drawRightString(7.48*inch,.42*inch,str(doc.page));c.restoreState()
    SimpleDocTemplate(str(OUT/'manuscript_en.pdf'),pagesize=(8.3*inch,11.7*inch),leftMargin=.82*inch,rightMargin=.83*inch,topMargin=.72*inch,bottomMargin=.7*inch,title=TITLE,author='Author details pending').build(story,onFirstPage=footer,onLaterPages=footer)

def tex_escape(s):
    replacements={'\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_','{':r'\{','}':r'\}','~':r'\textasciitilde{}','^':r'\textasciicircum{}'}
    s=''.join(replacements.get(c,c) for c in s)
    for a,b in [('μ',r'$\mu$'),('β',r'$\beta$'),('Δ',r'$\Delta$'),('θ',r'$\theta$'),('Σ',r'$\Sigma$'),('×',r'$\times$'),('−','-'),('–','--'),('—','---'),('’',"'"),('“','``'),('”',"''"),('≤',r'$\leq$'),('≥',r'$\geq$'),('→',r'$\to$'),('•',r'$\cdot$')]:s=s.replace(a,b)
    return s

def tex_inline_figure(name):
    from tex_figures import tex_inline_figure as render
    return render(name,OUT,NUM)

def latex():
    lines=[r'\documentclass[11pt]{article}',r'\usepackage[T1]{fontenc}',r'\usepackage[utf8]{inputenc}',r'\usepackage[margin=0.9in]{geometry}',r'\usepackage{amsmath,booktabs,array,longtable,graphicx,hyperref,pgfplots}',r'\usepgfplotslibrary{groupplots}',r'\pgfplotsset{compat=1.18}',r'\hypersetup{hidelinks}',r'\setlength{\emergencystretch}{3em}',r'\title{'+tex_escape(TITLE)+r'\\\large '+tex_escape(SUBTITLE)+'}',r'\author{Author details pending}',r'\date{Working draft for author review, 2 October 2026}',r'\begin{document}',r'\maketitle']
    for b in BLOCKS:
        k=b['type']
        if k=='heading':
            if b['text']=='Abstract':lines.append(r'\section*{Abstract}')
            else:lines.append((r'\section*{' if b['level']==1 else r'\subsection*{')+tex_escape(b['text'])+'}')
        elif k=='paragraph':lines.extend([tex_escape(b['text']),''])
        elif k=='equation':lines.extend([r'\begin{equation*}',b['latex'],r'\end{equation*}'])
        elif k=='table':
            n=len(b['headers'])
            if n==6:weights=[1.45,.62,1.05,.65,1.6,.6]
            elif n==5:weights=[1.3,.55,.55,.95,1.4]
            elif n==4:weights=[2.3,1.9,.7,.6] if 'S3.' in b['label'] else [2.9,1,1.45,1]
            else:weights=[3.3,1.1,1.3]
            spec='@{}'+''.join(r'>{\raggedright\arraybackslash}p{'+f'{.88*w/sum(weights):.6f}'+r'\linewidth}' for w in weights)+'@{}'
            header=' & '.join(tex_escape(t) for t in b['headers'])+r' \\'
            lines.extend([r'\par\medskip\noindent\textbf{'+tex_escape(b['label'])+'}',r'\begingroup\small\setlength{\tabcolsep}{4pt}',r'\begin{longtable}{'+spec+'}',r'\toprule',header,r'\midrule\endfirsthead',r'\toprule',header,r'\midrule\endhead',r'\bottomrule\endfoot',r'\bottomrule\endlastfoot'])
            lines.extend(' & '.join(tex_escape(t) for t in row)+r' \\' for row in b['rows'])
            lines.extend([r'\end{longtable}',r'\noindent\footnotesize '+tex_escape(b['note']),r'\par\endgroup\medskip'])
        elif k=='figure':
            lines.extend([r'\begin{figure}[htbp]',r'\centering',tex_inline_figure(b['name']),r'\caption*{\textbf{'+tex_escape(b['label'])+'}. '+tex_escape(b['caption'])+'}',r'\end{figure}'])
    # caption* is explicitly supplied; no external .bib or figure assets are needed.
    lines.insert(5,r'\usepackage{caption}')
    lines.extend([r'\section*{References}',r'\begin{enumerate}'])
    lines.extend(r'\item '+tex_escape(r) for r in REFERENCES)
    lines.extend([r'\end{enumerate}',r'\end{document}'])
    (OUT/'manuscript.tex').write_text('\n'.join(lines)+'\n')

markdown();word();pdf();latex()
word_count=sum(len(re.findall(r"\b[\w’-]+\b",b.get('text',''))) for b in BLOCKS if b['type']=='paragraph')
print(json.dumps({'english_body_words':word_count,'references':len(REFERENCES),'figures':sum(b['type']=='figure' for b in BLOCKS),'tables':sum(b['type']=='table' for b in BLOCKS),'formats':['md','docx','pdf','tex'],'new_fits':0}))
