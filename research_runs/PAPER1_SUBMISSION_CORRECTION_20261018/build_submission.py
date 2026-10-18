#!/usr/bin/env python3
"""Build a separate, verified editorial submission-fix directory from immutable Paper1 inputs.
Do not modify scientific data, older manuscript releases, or related Paper2 text.
"""
from __future__ import annotations
import argparse, hashlib, json, re, subprocess, zipfile
from pathlib import Path
from docx import Document
import fitz


def shafile(path):
    h=hashlib.sha256()
    with open(path,'rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''):h.update(b)
    return h.hexdigest()


def md_section(text,heading,next_heading,new):
    pat=re.escape(heading)+r'\n\n.*?(?=\n\n'+re.escape(next_heading)+')'
    out,n=re.subn(pat,heading+'\n\n'+new,text,count=1,flags=re.S)
    if n!=1:raise RuntimeError('Missing section '+heading)
    return out


def patch_docx(source,out,subs):
    doc=Document(source)
    for i,value in subs.items():doc.paragraphs[i].text=value
    out.parent.mkdir(parents=True,exist_ok=True);doc.save(out)
    return doc


def convert(out):
    for name in ['Main_Manuscript_REVISED.docx','Cover_Letter_REVISED.docx','S1_Appendix_REVISED.docx','Ethics_Explanation_Optional.docx']:
        p=out/name
        subprocess.run(['libreoffice','-env:UserInstallation=file:///tmp/p1_submission_correction_lo','--headless','--convert-to','pdf','--outdir',str(out),str(p)],check=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=90)
        assert (out/(p.stem+'.pdf')).is_file()


def build(src,out,config):
    c=json.loads(Path(config).read_text());out.mkdir(parents=True,exist_ok=True)
    ethics=c['ethics'];supp=c['supp_ethics'];cover=c['cover']
    m=src/'manuscript_en.md';s=src/'supplementary_en.md'
    assert shafile(m)=='6f7b4e04d76db3afc34b1a2ba7cb949fef8b6e81250a60906e53ff6ca150394d'
    x=m.read_text();x=md_section(x,'### Research ethics and access scope','### Evidence and reproducibility',ethics)
    (out/'manuscript_en_REVISED.md').write_text(x)
    y=s.read_text();a=y.index('The author reports only secondary access to participant codes');b=y.index('\n\n',a)
    y=y[:a]+supp+y[b:];(out/'supplementary_en_REVISED.md').write_text(y)
    (out/'Cover_Letter_REVISED.md').write_text('# Cover Letter\n\n'+'\n\n'.join(cover[1:])+'\n')
    (out/'Related_Manuscript_Disclosure.md').write_text(c['related_note'])
    (out/'UPLOAD_GUIDE_zh.md').write_text(c['guide'])
    (out/'Ethics_Explanation_REVISED.md').write_text('# '+c['author_statement'][0]+'\n\n'+'\n\n'.join(c['author_statement'][1:])+'\n')
    p1=Document(src/'main_manuscript.docx');q1=patch_docx(src/'main_manuscript.docx',out/'Main_Manuscript_REVISED.docx',{41:ethics})
    p2=Document(src/'S1_Appendix.docx');q2=patch_docx(src/'S1_Appendix.docx',out/'S1_Appendix_REVISED.docx',{43:supp})
    patch_docx(src/'submission/Cover_Letter.docx',out/'Cover_Letter_REVISED.docx',dict(enumerate(cover)))
    an=Document(src/'submission/Ethics_Author_Statement.docx')
    for i,p in enumerate(an.paragraphs):p.text=c['author_statement'][i] if i<len(c['author_statement']) else ''
    an.save(out/'Ethics_Explanation_Optional.docx')
    assert len(q1.tables)==len(p1.tables)==5 and len(q2.tables)==len(p2.tables)==5
    assert len(q1.paragraphs)==len(p1.paragraphs) and len(q2.paragraphs)==len(p2.paragraphs)
    for i,(p,q) in enumerate(zip(q1.paragraphs,p1.paragraphs)):
        if i!=41:assert p.text==q.text,('unexpected main edit',i)
    for i,(p,q) in enumerate(zip(q2.paragraphs,p2.paragraphs)):
        if i!=43:assert p.text==q.text,('unexpected S1 edit',i)
    for x,y in [(q1,p1),(q2,p2)]:
        for t,u in zip(x.tables,y.tables):
            assert [[cell.text for cell in row.cells] for row in t.rows]==[[cell.text for cell in row.cells] for row in u.rows]
    assert 'Research Administration Office' not in ethics and 'research administration office' not in supp.lower()
    assert 'PONE-D-26-52050' in '\n'.join(cover)
    convert(out)
    for name,pages in [('Main_Manuscript_REVISED.pdf',24),('Cover_Letter_REVISED.pdf',1),('S1_Appendix_REVISED.pdf',14)]:
        with fitz.open(out/name) as pdf:
            assert len(pdf)==pages,(name,len(pdf))
            assert all(pg.get_text().strip() for pg in pdf)
            if name=='Main_Manuscript_REVISED.pdf':
                assert sum('Article 32(1)' in pg.get_text() for pg in pdf)==1
                assert all('research administration office' not in pg.get_text().lower() for pg in pdf)
    # Do not post the in-review Paper2 manuscript to a public GitHub repository.
    receipt={'status':'PASS','source_commit':c['old_source_commit'],'scope':'ethics and related manuscript disclosure only',
             'no_new_training':True,'no_new_inference':True,'original_result_files_untouched':True,
             'main_tables_preserved':5,'S1_tables_preserved':5,'main_pages':24,'supplementary_pages':14,'cover_pages':1,
             'paper2_pdf_in_public_repo':False,
             'bytes_and_sha256':{p.name:{'bytes':p.stat().st_size,'sha256':shafile(p)} for p in out.glob('*') if p.is_file()}}
    (out/'VALIDATION.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n')
    included=[f for f in sorted(out.glob('*')) if f.is_file() and f.name not in ('Paper1_PLOS_Submission_Fix.zip','Paper1_PLOS_Submission_Fix.zip.sha256')]
    artifact=out/'Paper1_PLOS_Submission_Fix.zip'
    with zipfile.ZipFile(artifact,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for f in included:z.write(f,f.name)
    with zipfile.ZipFile(artifact) as z:
        assert z.testzip() is None
        assert all(hashlib.sha256(z.read(f.name)).hexdigest()==shafile(f) for f in included)
    (out/'Paper1_PLOS_Submission_Fix.zip.sha256').write_text(shafile(artifact)+'  '+artifact.name+'\n')
    print(json.dumps({'status':'PASS','updated_file_count':len(included),'zip_sha256':shafile(artifact),'zip_bytes':artifact.stat().st_size},ensure_ascii=False))

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--source',type=Path,required=True);ap.add_argument('--out',type=Path,required=True);ap.add_argument('--config',type=Path,required=True)
    a=ap.parse_args();build(a.source,a.out,a.config)
