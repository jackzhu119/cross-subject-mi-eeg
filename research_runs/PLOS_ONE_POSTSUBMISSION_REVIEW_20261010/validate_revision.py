"""Read-only checks of a postsubmission editorial candidate against submitted v1.1.

Only the two validation reports are written. This script never imports or executes
EEG analysis, fits a model, performs checkpoint inference, contacts a journal, or
publishes a release. Missing exports and unbound form receipts are failures, not
inferred successes. Run again after export/package preparation is complete.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from lxml import etree
from PIL import Image
import fitz

P = Path(__file__).resolve().parent
ROOT = P.parents[1]
SUBMITTED = ROOT / "research_runs/PLOS_ONE_SUBMISSION"
OUTPUTS = {"evidence/revision_validation.json", "evidence/revision_validation.md"}
WNS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CITATION = re.compile(r"\[(\d+(?:\s*,\s*\d+)*)\]")
NUMBER = re.compile(r"(?<![A-Za-z0-9_])[-+−–]?\d+(?:,\d{3})*(?:\.\d+)?(?:[eE][-+]?\d+)?")
GROUPING = re.compile(r"\[1,\s*2\],\s*\[3,\s*4\],\s*\[5,\s*6\],\s*(?:and\s*)?\[7,\s*8,\s*9\]")


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def norm(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("−", "-").replace("–", "-")).strip()


def numbers(text: str) -> list[str]:
    # Q/P project codes are identifiers, not scientific quantities. Electrode
    # names in text are never edited; excluding them here prevents false numeric
    # mismatches when semantic names replace internal condition identifiers.
    text = re.sub(r"\b[QP]\d+(?:[–-][QP]?\d+)?(?:[-/]E\d+)?\b", "", text)
    return [v.replace(",", "").replace("−", "-").replace("–", "-") for v in NUMBER.findall(text)]


def block_strings(block: dict):
    for key in ("text", "caption", "label", "note"):
        if block.get(key):
            yield str(block[key])
    for row in [block.get("headers", [])] + block.get("rows", []):
        for cell in row:
            yield str(cell)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--no-package", action="store_true", help="Pre-package check; absence of the full candidate ZIP is recorded as pending, not a complete delivery pass.")
    args = parser.parse_args()
    checks, details, pending, pdf_info = {}, {}, [], {}

    def check(name: str, value: bool, detail=None):
        checks[name] = bool(value)
        if detail is not None:
            details[name] = detail

    def read_json(relative: str, base: Path = P):
        path = base / relative
        if not path.is_file():
            check("present_" + str(path.relative_to(ROOT)), False, "Missing required input/export.")
            pending.append(str(path.relative_to(ROOT)))
            return None
        try:
            return json.loads(path.read_text())
        except (OSError, ValueError) as exc:
            check("readable_" + relative, False, type(exc).__name__)
            return None

    def guarded(name, operation):
        try:
            operation()
        except Exception as exc:
            check(name, False, {"error_type": type(exc).__name__, "message": str(exc)[:300]})

    baseline = read_json("evidence/source_protection_baseline.json")
    protected_count = 0
    if baseline:
        science = baseline.get("scientific_files", {})
        protected_count = len(science)
        check("protected_scientific_inventory_count", protected_count == 16417, protected_count)
        bad = []
        for name, expected in science.items():
            file = ROOT / name
            if not file.is_file() or digest(file) != expected:
                bad.append(name)
        check("all_16417_frozen_scientific_files_byte_identical", not bad, {"files_checked": protected_count, "mismatches": bad})
        original_baseline = read_json("evidence/source_protection_baseline.json", SUBMITTED)
        check("scientific_baseline_not_replaced", original_baseline is not None and science == original_baseline.get("scientific_files"))

    preservation = read_json("evidence/submitted_package_preservation.json")
    submitted_count = 0
    if preservation:
        inventory = preservation.get("all_submitted_package_files", {})
        submitted_count = len(inventory)
        check("submitted_package_inventory_count", submitted_count == 73, submitted_count)
        mismatches = []
        for name, rec in inventory.items():
            file = ROOT / name
            if not file.is_file() or file.stat().st_size != rec["bytes"] or digest(file) != rec["sha256"]:
                mismatches.append(name)
        check("submitted_v1_1_all_73_files_byte_identical", not mismatches, mismatches)
        check("submitted_v1_1_tag_and_artifact_pinned", preservation.get("submitted_tag") == "plos-one-submission-v1.1" and preservation.get("submitted_artifact_commit") == "1d4b0d1ef4d8bef19dd077a07e1a4ecdb1bde336")

    submitted_content = read_json("evidence/editorial_transformations.json", SUBMITTED)
    content = read_json("evidence/editorial_transformations.json")
    if content is None and (P / "revision_content.json").is_file():
        content = read_json("revision_content.json")
    main_blocks = content.get("main_blocks", []) if content else []
    appendix_blocks = content.get("appendix_blocks", []) if content else []
    blocks = main_blocks + appendix_blocks
    text = "\n".join(s for b in blocks for s in block_strings(b))
    main_text = "\n".join(s for b in main_blocks for s in block_strings(b))
    table_cells = 0
    if content and submitted_content:
        for section in ("main_blocks", "appendix_blocks"):
            old_tables = [b for b in submitted_content[section] if b["type"] == "table"]
            new_tables = [b for b in content[section] if b["type"] == "table"]
            check(section + "_original_table_count_retained", len(new_tables) >= len(old_tables), {"submitted": len(old_tables), "candidate": len(new_tables)})
            # Existing tables remain in original order; new diagnostic tables may
            # be appended. Every numeric cell is checked at its row/column grain.
            for ti, old in enumerate(old_tables):
                if ti >= len(new_tables):
                    continue
                new = new_tables[ti]
                shape = len(old["rows"]) == len(new["rows"]) and len(old["headers"]) == len(new["headers"])
                check(f"{section}_table_{ti+1}_shape", shape, old["label"])
                for ri, old_row in enumerate(old["rows"]):
                    if ri >= len(new["rows"]):
                        continue
                    check(f"{section}_table_{ti+1}_row_{ri+1}_numeric_cells", len(old_row) == len(new["rows"][ri]) and all(numbers(str(a)) == numbers(str(b)) for a, b in zip(old_row, new["rows"][ri])), {"submitted_numeric_cells": [numbers(str(x)) for x in old_row], "candidate_numeric_cells": [numbers(str(x)) for x in new["rows"][ri]]})
                    table_cells += len(old_row)
            old_eq = [b for b in submitted_content[section] if b["type"] == "equation"]
            new_eq = [b for b in content[section] if b["type"] == "equation"]
            check(section + "_equation_count_retained", len(new_eq) >= len(old_eq))
            for ei, old in enumerate(old_eq):
                check(f"{section}_equation_{ei+1}_latex_identical", ei < len(new_eq) and new_eq[ei].get("latex") == old.get("latex"))

        invariant_numbers = ["33.72", "42.67", "87.57", "+0.573", "-0.097", "+1.242", "-1.513", "-2.249", "-0.804", "+0.204", "-0.515", "+0.969", "51.99", "52.77", "5184", "2592", "10520", "10800", "4918"]
        canonical = norm(text).replace(",", "")
        for value in invariant_numbers:
            check("frozen_displayed_quantity_" + value, value in canonical)
        check("source_validation_unequal_fixed_groups_disclosed", "2/2/2/3" in text and "equal fold" in text.lower())
        check("Q16_descriptive_boundary_retained", "descriptive" in main_text.lower() and ("sensor-level" in main_text.lower() or "sensor level" in main_text.lower()) and ("post-outcome" in main_text.lower() or "after decoder outcomes" in main_text.lower()) and ("no p-values" in main_text.lower() or "without p-values" in main_text.lower() or "no p values" in main_text.lower()))
        check("complete_pipeline_causal_boundary_retained", "complete" in main_text.lower() and "pipeline" in main_text.lower() and "14" in main_text and "19" in main_text and "BatchNorm" in main_text)

    abstract_path = P / "PLOS_ONE_Abstract.txt"
    if abstract_path.is_file():
        abstract = abstract_path.read_text().strip()
        words = len(abstract.split())
        check("abstract_official_word_limit_no_citations", 0 < words <= 300 and not CITATION.search(abstract), {"words": words, "journal_limit": 300, "requested_target": "approximately 180–230"})
        check("abstract_requested_word_range", 180 <= words <= 230, words)
        check("abstract_matches_structured_manuscript", any(b.get("text") == abstract for b in main_blocks))
        check("abstract_adverse_and_uncertain_external_effects_retained", all(x in norm(abstract) for x in ["+0.573", "-1.513", "+0.204"]))
    else:
        check("abstract_present", False)
        pending.append("PLOS_ONE_Abstract.txt")

    refs = read_json("references_plos.json") or []
    old_refs = read_json("references_plos.json", SUBMITTED) or []
    if refs:
        nums = [r.get("plos_number", r.get("number")) for r in refs]
        check("reference_numbering_continuous_unique", nums == list(range(1, len(refs) + 1)))
        keys = [r.get("key") for r in refs]
        check("reference_keys_unique_and_original_references_retained", len(keys) == len(set(keys)) and set(r["key"] for r in old_refs) <= set(keys))
        seen, uses = [], Counter()
        for b in blocks:
            # Numeric confusion matrices, interval cells and other table entries
            # are data, not bracketed reference citations.
            for s in (str(b[key]) for key in ("text", "caption", "note") if b.get(key)):
                # Fixed source-validation subject groups are not citations.
                s = GROUPING.sub("SOURCE_VALIDATION_GROUPS", s)
                for match in CITATION.finditer(s):
                    for n in map(int, match.group(1).replace(" ", "").split(",")):
                        uses[n] += 1
                        if n not in seen:
                            seen.append(n)
        check("no_dangling_or_unused_references", set(uses) == set(nums), {"uncited": sorted(set(nums) - set(uses)), "dangling": sorted(set(uses) - set(nums))})
        check("references_first_appearance_order", seen == list(range(1, len(refs) + 1)), seen)
        for r in refs:
            key = str(r.get("key"))
            check("reference_" + key + "_metadata_present", bool(r.get("title") and r.get("authors") and r.get("year") and r.get("vancouver")))
            doi = r.get("doi")
            check("reference_" + key + "_DOI_syntax", not doi or bool(re.fullmatch(r"10\.\d{4,9}/\S+", str(doi))))
            explicit_record = r.get("verified") is True and bool(r.get("verified_via") or r.get("verification_sources") or r.get("metadata_source_sha256"))
            scoped_record = bool(r.get("verification_scope") and (r.get("formal_publisher_url") or r.get("url")) and (r.get("metadata_source") or r.get("verification_sources")))
            check("reference_" + key + "_verification_record_present", explicit_record or scoped_record, "Saved metadata/source/scope coverage, not a fresh network DOI resolution or assertion that every full text was accessible.")

    documents = {}
    for name, source in [("PLOS_ONE_Manuscript.docx", main_blocks), ("SupportingInformation/S1_Appendix.docx", appendix_blocks), ("PLOS_ONE_Cover_Letter.docx", None)]:
        path = P / name
        if not path.is_file():
            check(name + "_present", False)
            pending.append(name)
            continue
        def inspect_docx(name=name, path=path, source=source):
            doc = Document(path)
            documents[name] = doc
            protection = doc.settings.element.find(".//{%s}documentProtection" % WNS)
            check(name + "_unlocked", protection is None or protection.get("{%s}enforcement" % WNS) in {"0", "false", "off"})
            check(name + "_heading_depth_max_three", all(not x.style.name.startswith("Heading ") or int(x.style.name.split()[-1]) <= 3 for x in doc.paragraphs))
            check(name + "_single_column", all(s._sectPr.find("{%s}cols" % WNS) is None or s._sectPr.find("{%s}cols" % WNS).get("{%s}num" % WNS, "1") == "1" for s in doc.sections))
            if source is not None:
                ln = doc._element.find(".//{%s}lnNumType" % WNS)
                check(name + "_continuous_line_numbers_setting", ln is not None and ln.get("{%s}restart" % WNS) == "continuous")
                check(name + "_double_spaced", doc.styles["Normal"].paragraph_format.line_spacing == 2)
                check(name + "_page_number_field", any("PAGE" in p._p.xml for s in doc.sections for p in s.footer.paragraphs))
                tables = [b for b in source if b["type"] == "table"]
                check(name + "_rendered_table_count", len(doc.tables) == len(tables), {"actual": len(doc.tables), "source": len(tables)})
                for i, (actual, block) in enumerate(zip(doc.tables, tables), 1):
                    check(f"{name}_rendered_table_{i}_all_text_cells", [[c.text for c in row.cells] for row in actual.rows] == [[str(c) for c in block["headers"]]] + [[str(c) for c in row] for row in block["rows"]])
                rendered = [p.text for p in doc.paragraphs]
                for i, block in enumerate(source):
                    if block["type"] in {"paragraph", "heading", "equation"}:
                        check(f"{name}_rendered_source_block_{i}", (str(block["text"]) in rendered or (block["type"] == "equation" and str(block["text"]) + "  (" + str(block.get("equation_number")) + ")" in rendered)))
            if name == "PLOS_ONE_Manuscript.docx":
                check("main_document_no_embedded_figures", len(doc.inline_shapes) == 0)
                rendered = "\n".join(p.text for p in doc.paragraphs)
                check("main_author_ORCID_and_email", all(x in rendered for x in ["Ziyuan Zhu", "0009-0005-1153-4926", "zzy2630816871@gmail.com"]))
                check("main_truthful_AI_scope", all(x in rendered for x in ["ChatGPT", "Codex"]) and "model version" in rendered.lower() and "review" in rendered.lower())
                check("main_secondary_ethics_author_boundary", "No formal approval or exemption documentation" in rendered and ("author’s assessment" in rendered or "author's assessment" in rendered) and "institutional or committee determination" in rendered)
                for i in range(1, 7):
                    check(f"main_Fig_{i}_caption_once", sum(p.text.startswith(f"Fig {i}. ") for p in doc.paragraphs) == 1)
                indexed_paragraphs = []
                for index, node in enumerate(doc.element.body):
                    if node.tag.endswith("}p"):
                        indexed_paragraphs.append((index, "".join(node.xpath(".//w:t/text()"))))
                labels = [f"Fig {i}" for i in range(1, 7)] + [f"Table {i}" for i in range(1, len(doc.tables) + 1)]
                for label in labels:
                    appearances = [(index, value) for index, value in indexed_paragraphs if re.search(re.escape(label) + r"(?!\d)", value)]
                    captions = [(index, value) for index, value in appearances if value.startswith(label + ".")]
                    check("main_" + label.replace(" ", "_") + "_caption_immediately_after_first_citation", bool(appearances and captions and not appearances[0][1].startswith(label + ".") and captions[0][0] == appearances[0][0] + 1))
                reference_paragraphs = [p.text for p in doc.paragraphs]
                check("main_rendered_reference_list_matches_metadata", all(f"{r.get('plos_number', r.get('number'))}. {r['vancouver']}" in reference_paragraphs for r in refs))
            check(name + "_no_obsolete_draft_label", not re.search(r"Scientific draft for author review|NOT_SUBMITTED_PORTAL_ACTIONS_REMAIN|not yet submitted|JNE under review|NOT FOR SUBMISSION", "\n".join(p.text for p in doc.paragraphs), re.I))
        guarded(name + "_inspection_completed", inspect_docx)

    for i in range(1, 7):
        name = f"Figures/Fig{i}.tif"
        path, submitted = P / name, SUBMITTED / name
        if not path.is_file():
            check(name + "_present", False)
            pending.append(name)
            continue
        def inspect_figure(path=path, submitted=submitted, i=i):
            check(f"Fig{i}_submitted_TIFF_bytes_unchanged", digest(path) == digest(submitted))
            with Image.open(path) as im:
                dpi = tuple(float(value) for value in im.info.get("dpi", (0, 0)))
                width = im.width / dpi[0] * 2.54 if dpi[0] else 0
                height = im.height / dpi[1] * 2.54 if dpi[1] else 0
                check(f"Fig{i}_TIFF_RGB_single_page_LZW", im.format == "TIFF" and im.mode == "RGB" and im.n_frames == 1 and im.tag_v2.get(259) == 5)
                check(f"Fig{i}_journal_resolution_and_size", all(300 <= x <= 600.1 for x in dpi) and 6.68 <= width <= 19.051 and height <= 22.231 and path.stat().st_size < 10_000_000, {"pixels": list(im.size), "dpi": list(dpi), "width_cm": width, "height_cm": height, "bytes": path.stat().st_size})
            svg_name = f"FigureSources/Fig{i}.svg"
            check(f"Fig{i}_SVG_source_bytes_unchanged", digest(P / svg_name) == digest(SUBMITTED / svg_name))
        guarded(f"Fig{i}_inspection_completed", inspect_figure)
    workflow = P / "FigureSources/Fig1.svg"
    if workflow.is_file():
        check("Fig1_baseline_relative_Q16_label", "Baseline-relative μ/β power" in workflow.read_text() and "ERD / ERS" not in workflow.read_text())
    check("all_supporting_information_files_below_20MB", all(file.stat().st_size < 20_000_000 for file in (P / "SupportingInformation").glob("*") if file.is_file()), {file.name: file.stat().st_size for file in (P / "SupportingInformation").glob("*") if file.is_file()})
    blank_form = P / "AuthorForms/PLOS_Human_Participants_Checklist_2026_blank.pdf"
    check("official_blank_human_form_original_bytes", blank_form.is_file() and digest(blank_form) == "ebf2a9913fe2926669dfb60c1ba5f9f98130b00b9e858bf9b1d243b3707162e4")

    old_inventory = read_json("evidence/derived_data_inventory.json", SUBMITTED)
    candidate_archive = P / "SupportingInformation/S1_Data.zip"
    if candidate_archive.is_file() and old_inventory:
        def inspect_data():
            with zipfile.ZipFile(candidate_archive) as new, zipfile.ZipFile(SUBMITTED / "SupportingInformation/S1_Data.zip") as old:
                names = new.namelist()
                check("S1_Data_no_duplicate_members_crc_valid", len(names) == len(set(names)) and new.testzip() is None)
                check("S1_Data_original_66_scientific_members_count", len(old_inventory["files"]) == 66)
                bad = []
                for name, rec in old_inventory["files"].items():
                    if name not in names or hashlib.sha256(new.read(name)).hexdigest() != rec["sha256"] or new.read(name) != old.read(name):
                        bad.append(name)
                check("S1_Data_original_66_scientific_member_bytes_identical", not bad, bad)
                additions = sorted(set(names) - set(old.namelist()))
                check("S1_Data_added_diagnostics_separate_namespace", all(name.startswith(("postsubmission_diagnostics/", "postsubmission/", "diagnostics/")) for name in additions), additions)
                administrative = [n for n in old.namelist() if n not in old_inventory["files"]]
                check("S1_Data_original_administrative_members_retained", all(n in names for n in administrative), administrative)
                check("S1_Data_no_raw_EEG_or_weights_added", not any(n.lower().endswith((".mat", ".edf", ".gdf", ".fif", ".pt", ".pth", ".joblib", ".pkl")) for n in additions))
        guarded("S1_Data_inspection_completed", inspect_data)
    else:
        check("S1_Data_archive_present", False)
        pending.append("SupportingInformation/S1_Data.zip")

    line_text = {}
    for name in ["PLOS_ONE_Manuscript_Review.pdf", "SupportingInformation/S1_Appendix.pdf", "PLOS_ONE_Cover_Letter.pdf"]:
        path = P / name
        if not path.is_file():
            check(name + "_present", False)
            pending.append(name)
            continue
        def inspect_pdf(name=name, path=path):
            with fitz.open(path) as pdf:
                outside, gutters, page_texts, numbered_pages = [], [], [], []
                for page in pdf:
                    page_texts.append(page.get_text())
                    lines = [line for b in page.get_text("dict")["blocks"] if "lines" in b for line in b["lines"]]
                    body = [line for line in lines if line["bbox"][0] >= 60 and line["bbox"][1] < page.rect.height - 50]
                    numbered_pages.append(any("".join(s["text"] for s in line["spans"]).strip() == str(page.number + 1) and line["bbox"][1] > page.rect.height - 60 and page.rect.width * .35 < line["bbox"][0] < page.rect.width * .65 for line in lines))
                    for line in lines:
                        spans = line["spans"]
                        for span in spans:
                            x0, y0, x1, y1 = span["bbox"]
                            if x0 < -1 or y0 < -1 or x1 > page.rect.width + 1 or y1 > page.rect.height + 1:
                                outside.append({"page": page.number + 1, "text": span["text"], "bbox": span["bbox"]})
                        value = "".join(s["text"] for s in spans).strip()
                        if line["bbox"][0] < 50 and value.isdigit():
                            gutters.append(int(value))
                            if name.startswith("PLOS_ONE_Manuscript") and body:
                                match = min(body, key=lambda other: abs(other["bbox"][1] - line["bbox"][1]))
                                if abs(match["bbox"][1] - line["bbox"][1]) < 2:
                                    line_text[int(value)] = "".join(s["text"] for s in match["spans"])
                check(name + "_no_text_outside_page", not outside, outside[:20])
                check(name + "_searchable_text_present", len("".join(page_texts)) > 500)
                check(name + "_PDF_unencrypted", not pdf.is_encrypted)
                check(name + "_actual_page_numbers", all(numbered_pages), {"pages": len(pdf), "numbered_pages": sum(numbered_pages)})
                check(name + "_no_replacement_character", "\ufffd" not in "".join(page_texts))
                if name.startswith("PLOS_ONE_Manuscript"):
                    check("main_review_pdf_actual_continuous_gutter_numbers", len(gutters) > 500 and gutters == list(range(gutters[0], gutters[-1] + 1)), {"count": len(gutters), "first": gutters[0] if gutters else None, "last": gutters[-1] if gutters else None})
                    check("main_review_pdf_no_embedded_figures", all(not page.get_images() for page in pdf))
                pdf_info[name] = {"pages": len(pdf), "bytes": path.stat().st_size, "sha256": digest(path), "gutter_line_numbers": len(gutters)}
        guarded(name + "_inspection_completed", inspect_pdf)

    form = read_json("AuthorForms/form_completion_receipt.json")
    if form:
        for key, path in [("manuscript_docx_sha256", "PLOS_ONE_Manuscript.docx"), ("manuscript_review_pdf_sha256", "PLOS_ONE_Manuscript_Review.pdf"), ("completed_form_sha256", "AuthorForms/PLOS_Human_Participants_Checklist_2026_completed.pdf")]:
            check("human_form_current_binding_" + key, (P / path).is_file() and form.get(key) == digest(P / path))
        ranges = form.get("methods_line_ranges", {})
        access = read_json("evidence/research_access_dates.json")
        for key in ["secondary_analysis_ethics_author_assessment", "documented_access_dates", "direct_identifier_access_author_report"]:
            interval = ranges.get(key, [])
            valid = len(interval) == 2 and all(isinstance(n, int) for n in interval) and interval[1] >= interval[0]
            actual_text = " ".join(line_text.get(i, "") for i in range(interval[0], interval[1] + 1)) if valid else ""
            if key == "secondary_analysis_ethics_author_assessment":
                matches = "No formal approval or exemption documentation" in norm(actual_text) and "institutional or committee determination" in norm(actual_text)
            else:
                expected = (access or {}).get("methods_access_paragraph" if key == "documented_access_dates" else "methods_identifiability_paragraph", "")
                matches = bool(expected) and norm(actual_text) == norm(expected)
            check("human_form_actual_PDF_line_binding_" + key, valid and matches, {"range": interval, "extracted_text": actual_text})
        def inspect_form():
            with fitz.open(P / "AuthorForms/PLOS_Human_Participants_Checklist_2026_completed.pdf") as pdf:
                actual = {(page.number + 1, w.field_name): w.field_value for page in pdf for w in page.widgets()}
                check("human_form_receipt_matches_actual_widgets", len(actual) == len(form.get("fields", [])) and all(actual.get((rec["page"], rec["field"])) == rec["value"] for rec in form.get("fields", [])))
                check("human_form_no_formal_ethics_exemption_invented", actual.get((1, "NA")) == "X" and not actual.get((1, "Uploaded")))
                check("human_form_no_prospective_recruitment", actual.get((2, "Completed")) == "X" and not actual.get((2, "Please state the line numbers in the Methods where this is reported 2")))
                check("human_form_retrospective_access_completed", actual.get((2, "Please state the line numbers in the Methods where this is reported 2_2")) == "X" and not actual.get((2, "Completed_2")))
        guarded("human_form_actual_widgets_inspection_completed", inspect_form)

    status = read_json("evidence/source_version.json")
    if status:
        check("candidate_current_target_PLOS_ONE", status.get("target_journal") == "PLOS ONE")
        values = json.dumps(status, ensure_ascii=False)
        check("submitted_manuscript_id_and_date_recorded", "PONE-D-26-50597" in values and "2026-10-09" in values)
        check("submission_author_reported_portal_not_accessed", "AUTHOR_REPORTED" in values.upper() and status.get("live_submission_portal_accessed") is False)
        approval = str(status.get("candidate_approval", status.get("new_formatted_version_author_approval", "")))
        check("new_candidate_approval_pending_not_fabricated", "PENDING" in approval.upper() or "REVIEW_REQUIRED" in approval.upper(), approval)
        check("candidate_not_submitted_or_formally_released", status.get("candidate_submitted") is False and status.get("formal_release_created") is False)
        counters = {key: status.get(key) for key in ["new_model_fits", "new_checkpoint_inference", "new_q16_scientific_execution"]}
        check("candidate_declares_no_new_model_or_Q16_execution", all(v == 0 for v in counters.values()), counters)
        check("candidate_not_obsolete_unsubmitted_manuscript_status", "NOT_SUBMITTED_PORTAL_ACTIONS_REMAIN" not in values)

    stale = re.compile(r"NOT_SUBMITTED_PORTAL_ACTIONS_REMAIN|not yet submitted|JNE under review|Scientific draft for author review", re.I)
    active_admin = ["README.md", "PLOS_ONE_Submission_Checklist_zh.md", "PLOS_ONE_Submission_Fields.md", "Changes_and_Preservation.md", "REPRODUCIBILITY.md"]
    for name in active_admin:
        file = P / name
        check("active_metadata_" + name + "_present_current", file.is_file() and not stale.search(file.read_text()))
    for name in ["LICENSE", "LICENSE-DATA.md", "LICENSING.md"]:
        check("author_owned_license_scope_preserved_" + name, (P / name).is_file() and (P / name).read_bytes() == (ROOT / name).read_bytes())

    candidates = sorted(f for f in P.glob("*Package.zip") if "Submission_Package" not in f.name)
    if candidates:
        archive = candidates[0]
        def inspect_package():
            with zipfile.ZipFile(archive) as z:
                check("full_candidate_archive_crc_unique_paths", z.testzip() is None and len(z.namelist()) == len(set(z.namelist())) and all(not name.startswith("/") and ".." not in Path(name).parts for name in z.namelist()))
                names = z.namelist()
                required = ["PLOS_ONE_Manuscript.docx", "PLOS_ONE_Manuscript_Review.pdf", "SupportingInformation/S1_Appendix.docx", "SupportingInformation/S1_Appendix.pdf", "SupportingInformation/S1_Data.zip", "references_plos.json"] + [f"Figures/Fig{i}.tif" for i in range(1, 7)]
                for relative in required:
                    matches = [name for name in names if name == relative or name.endswith("/" + relative)]
                    check("package_current_byte_readback_" + relative, len(matches) == 1 and (P / relative).is_file() and hashlib.sha256(z.read(matches[0])).hexdigest() == digest(P / relative))
                details["full_candidate_archive"] = {"name": archive.name, "bytes": archive.stat().st_size, "sha256": digest(archive), "members": len(names), "validation_reports_self_binding_excluded": sorted(OUTPUTS)}
        guarded("candidate_archive_inspection_completed", inspect_package)
    elif args.no_package:
        pending.append("Full candidate ZIP not checked (--no-package); delivery acceptance remains pending.")
    else:
        check("full_candidate_archive_present", False)
        pending.append("Full postsubmission candidate Package.zip")

    failures = [name for name, passed in checks.items() if not passed]
    report = {
        "status": "failed" if failures else ("pending_package" if pending else "passed"),
        "scope": "Saved-byte preservation, declared editorial-content consistency, document/export/form bindings and package readback. No experimental re-execution, journal action, fresh literature-network resolution or universal implementation-correctness claim.",
        "checked_at_utc": datetime.now(timezone.utc).isoformat(),
        "check_count": len(checks), "checks": checks, "failed_checks": failures,
        "pending_inputs_or_exports": sorted(set(pending)), "details": details,
        "protected_scientific_files": protected_count, "submitted_package_files": submitted_count,
        "original_table_cells_checked": table_cells, "pdfs": pdf_info,
        "new_model_fits": 0, "new_checkpoint_inference": 0, "new_Q16_experiments": 0,
        "journal_submission_performed": False, "email_sent": False, "release_created": False,
        "candidate_author_approval": "PENDING_AUTHOR_REVIEW",
        "limitations": ["Numerical preservation checks do not prove all historical model/adapter implementation errors absent.", "Reference saved metadata verification records require the separate primary-source literature audit.", "This delivery check does not determine whether missing source-audit experiment records have been supplied or scientifically verified.", "PDF coordinate and gutter checks support layout review but do not replace visual author review or portal-generated PDF review."]
    }
    (P / "evidence").mkdir(exist_ok=True)
    (P / "evidence/revision_validation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    md = ["# Postsubmission candidate validation", "", f"Status: **{report['status']}**", "", report["scope"], "", f"Checks: {len(checks)}; failed: {len(failures)}. Protected scientific files: {protected_count}; preserved submitted package files: {submitted_count}; original table cells checked: {table_cells}.", "", "## Failed checks", ""]
    md += ["- " + name for name in failures] or ["None."]
    md += ["", "## Pending inputs/exports", ""] + (["- " + name for name in report["pending_inputs_or_exports"]] or ["None."])
    md += ["", "## Actual PDF exports", "", "| File | Pages | Bytes | SHA-256 |", "| --- | ---: | ---: | --- |"]
    md += [f"| {name} | {rec['pages']} | {rec['bytes']} | {rec['sha256']} |" for name, rec in pdf_info.items()]
    md += ["", "## Limits and required author review", ""] + ["- " + item for item in report["limitations"]]
    md += ["", "No new model fitting, checkpoint inference, Q16 experiment, journal submission, editor contact, email delivery or formal release was performed by this validation script. Candidate approval remains pending author review.", ""]
    (P / "evidence/revision_validation.md").write_text("\n".join(md))
    print(json.dumps({"status": report["status"], "checks": len(checks), "failures": failures, "pending": report["pending_inputs_or_exports"], "pdf_pages": {name: rec["pages"] for name, rec in pdf_info.items()}}, ensure_ascii=False))
    return int(bool(failures) or bool(pending))


if __name__ == "__main__":
    raise SystemExit(main())
