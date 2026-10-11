"""Export the editorial candidate; never load EEG, models, or frozen scientific executables."""
from pathlib import Path
import ast,copy,json,re,hashlib,zipfile,shutil,difflib,subprocess,sys,html
from datetime import datetime,timezone
from docx import Document
from docx.shared import Inches,Pt,Mm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
OUT=Path(__file__).resolve().parent; ROOT=OUT.parents[1]; SRC=ROOT/'research_runs/PAPER_FINAL_20261006'; OLD=ROOT/'research_runs/PLOS_ONE_SUBMISSION'
DATA=json.loads((OUT/'revision_content.json').read_text()); TITLE=DATA['title']; SHORT='Source-only selection and complete spectral-sharing pipelines in motor-imagery EEG'
AUTHOR=json.loads((SRC/'manuscript_content.json').read_text())['authors'][0];AUTHOR=copy.deepcopy(AUTHOR);AUTHOR['orcid']='0009-0005-1153-4926'
refs=[r['vancouver'] for r in DATA['references']];main=DATA['main_blocks'];appendix=DATA['appendix_blocks']
def save(n,t):(OUT/n).write_text(t.rstrip()+'\n')
def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
# Reuse only known document formatting helpers from the immutable submitted exporter.
tree=ast.parse((OLD/'build_plos_package.py').read_text())
funcs=[x for x in tree.body if isinstance(x,ast.FunctionDef) and x.name in ['caption_text','document','mdblocks']]
code=ast.unparse(ast.Module(body=funcs,type_ignores=[]))
code=code.replace("doc.add_paragraph(b['text'])", "doc.add_paragraph(b['text'] + ('  (' + str(b['equation_number']) + ')' if b.get('equation_number') else ''))")
exec(compile(code,'submitted_document_formatting_helpers','exec'),globals())
front='# '+TITLE+'\n\n'+AUTHOR['name']+'\n\n'+AUTHOR['affiliation']+'\n\nCorresponding author: '+AUTHOR['email']+'\n\nORCID: https://orcid.org/'+AUTHOR['orcid']+'\n\nPostal address: '+AUTHOR.get('postal_address','')+'\n\n'
save('PLOS_ONE_Manuscript.md',front+mdblocks(main)+'\n\n## References\n\n'+(OUT/'references_plos.md').read_text()+'\n\n## Supporting information captions\n\nS1 Appendix. Supplementary results, detailed methods and post-submission diagnostic context.\n\nS1 Data. Frozen derived outputs with separately identified archived-output diagnostic additions.')
save('SupportingInformation/S1_Appendix.md',mdblocks(appendix)+'\n\n## References\n\n'+(OUT/'references_plos.md').read_text())
save('PLOS_ONE_Abstract.txt',main[1]['text'])
save('PLOS_ONE_Title_and_Keywords.md','# Title\n\n'+TITLE+'\n\n# Short title\n\n'+SHORT+'\n\n# Keywords\n\nMotor imagery; EEG; source-only model selection; cross-subject transfer; balanced accuracy; spectral sharing; reproducibility')
document(TITLE,main,'PLOS_ONE_Manuscript.docx')
document('S1 Appendix',appendix,'SupportingInformation/S1_Appendix.docx',images=True)
# Cover text is a draft for a permitted revision, not a new-submission declaration.
cover=[]
for line in (OUT/'PLOS_ONE_Cover_Letter.md').read_text().splitlines():
 if line and not line.startswith('#'):cover.append({'type':'paragraph','text':line})
document('Draft revision cover letter — PLOS ONE',cover,'PLOS_ONE_Cover_Letter.docx',bibliography=False,cover=True)
# Diff report uses original and candidate full text; it is a review aid, not a submitted manuscript.
old=(OLD/'PLOS_ONE_Manuscript.md').read_text();new=(OUT/'PLOS_ONE_Manuscript.md').read_text()
save('Revision_Changes_Diff.html',difflib.HtmlDiff(wrapcolumn=100).make_file(old.splitlines(),new.splitlines(),'Preserved v1.1 manuscript','Post-submission candidate',context=True,numlines=3))
doc=Document();doc.add_heading('Post-submission candidate: change review',0);doc.add_paragraph('PONE-D-26-50597; internal candidate only. Red strike-through marks removed text; green underline marks added text. Primary frozen outputs are unchanged. This is a review aid, not a revised submission.')
for row in difflib.ndiff(old.splitlines(),new.splitlines()):
 if row.startswith('? ') or not row[2:].strip():continue
 if row.startswith('  '):continue
 p=doc.add_paragraph();r=p.add_run(row[2:]);r.font.size=Pt(10)
 if row.startswith('- '):r.font.color.rgb=RGBColor(150,0,0);r.font.strike=True
 else:r.font.color.rgb=RGBColor(0,100,50);r.font.underline=True
doc.save(OUT/'Revision_Changes_Tracked.docx')
# S1 Data: all original members unchanged; additions have explicit postsubmission namespace.
QA=OUT/'evidence'; originalzip=OLD/'SupportingInformation/S1_Data.zip';dest=OUT/'SupportingInformation/S1_Data.zip'
with zipfile.ZipFile(originalzip) as zi,zipfile.ZipFile(dest,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=9) as zo:
 for n in zi.namelist():zo.writestr(n,zi.read(n))
 for n in ['m1_csp_audit.json','new_report_csp_review.json','new_report_csp_per_person_confusion.csv','audit_archived_csp_outputs.py']:
  p=QA/n
  if p.exists():zo.writestr('postsubmission_diagnostics/'+n,p.read_bytes())
 zo.writestr('postsubmission_diagnostics/README.txt','Read-only reconstructions from original frozen CSP predictions. No new model fits, inference or primary results. Equal-person BA and pooled class frequencies use different aggregation. New source LOSO and mean-centering numbers in the supplied report lack execution artifacts and are not verified endpoints. Original 68 archive members remain byte-identical.\n')
# Actual PDF conversion with an isolated LibreOffice profile.
soffice=shutil.which('soffice') or '/opt/codex/runtimes/codex-primary-runtime/dependencies/bin/override/soffice'
profile=Path('/tmp/plos-postsubmission-lo');profile.mkdir(exist_ok=True)
for sub,files in [('', ['PLOS_ONE_Manuscript.docx','PLOS_ONE_Cover_Letter.docx','Revision_Changes_Tracked.docx']),('SupportingInformation',['S1_Appendix.docx'])]:
 cmd=[soffice,'-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to','pdf','--outdir',str(OUT/sub)]+[str(OUT/sub/f) for f in files]
 subprocess.run(cmd,check=True,capture_output=True,text=True,timeout=240)
(OUT/'PLOS_ONE_Manuscript.pdf').replace(OUT/'PLOS_ONE_Manuscript_Review.pdf')
save('evidence/export_receipt.json',json.dumps({'generated_at_utc':datetime.now(timezone.utc).isoformat(),'scope':'editorial candidate export only','new_fits':0,'new_inference':0,'new_q16_experiments':0,'main_docx_sha256':digest(OUT/'PLOS_ONE_Manuscript.docx'),'main_pdf_sha256':digest(OUT/'PLOS_ONE_Manuscript_Review.pdf'),'abstract_word_count':len(main[1]['text'].split()),'main_figure_embedding':'standalone figures; captions only in submission manuscript','original_figures_byte_preserved':True},indent=2))
print('exported_candidate',TITLE,len(refs),'references')
