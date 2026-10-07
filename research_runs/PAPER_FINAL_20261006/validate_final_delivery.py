"""Check manuscript consistency and saved evidence; no EEG or model execution."""
from pathlib import Path
import hashlib
import importlib.metadata
import json
import re
import subprocess
from datetime import datetime, timezone

import fitz
import pandas as pd
from docx import Document

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
def load(name): return json.loads((OUT / name).read_text())
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
data = load('manuscript_content.json')
q15 = load('evidence/q15_numbers.json')
internal = load('evidence/internal_review.json')
review = load('evidence/final_scientific_review.json')
checks = {}
def check(name, condition):
    checks[name] = bool(condition)
    if not condition: raise AssertionError(name)

check('manuscript_preparation_no_fit_or_checkpoint_inference', data['new_fits'] == data['new_checkpoint_inference'] == 0)
check('completed_q15_independent_arithmetic', q15['passed'])
check('original_source_fits15_new_source0_target0', q15['original_source_fits'] == 15 and q15['new_source_fits'] == q15['target_fits'] == 0)
check('earlier_saved_prediction_review', internal['arm_count_recomputed'] == 38 and internal['prediction_rows_rescored_sum'] == 809434)
check('final_scientific_review', review['status'] == 'scientific_and_numerical_review_passed' and not review['checks_failed'])
refs = load('references.json')['references']
check('28_verified_unique_references', len(refs) == len({r['key'] for r in refs}) == 28 and all(r['verified'] for r in refs))
check('28_manuscript_references_cited', len(data['references']) == len(set(data['reference_keys'])) == 28)
reviewed_hashes = review['input_sha256']
for name, digest in reviewed_hashes.items():
    check('review_binding_' + name, sha(OUT / name) == digest)
for name in ['manuscript_content.json', 'references.json', 'supplementary_methods.md', 'evidence/q16_analysis/summary.json', 'evidence/q16_analysis/independent_validation.json']:
    check('current_review_covers_' + name, name in reviewed_hashes)
check('author_correspondence', data['authors'][0]['name'] == 'Ziyuan Zhu' and data['authors'][0]['email'] == 'zzy2630816871@gmail.com' and data['authors'][0]['corresponding'])
check('published_affiliation_usage', data['authors'][0]['affiliation'] == load('evidence/affiliation_verification.json')['recommended_affiliation_en'])
author_declarations = load('evidence/author_declarations.json')
check('no_funding_author_confirmed', author_declarations['funding'] == 'none')
check('no_competing_interests_author_confirmed', author_declarations['competing_interests'] == 'none')
check('no_approval_or_exemption_required_author_confirmed', author_declarations['ethics_approval_required_for_this_secondary_analysis'] is False and author_declarations['ethics_exemption_required_for_this_secondary_analysis'] is False)
check('author_final_review_approval_confirmed', author_declarations['final_manuscript_author_approval_asserted'] and author_declarations['actual_contributions_author_confirmed'] and author_declarations['originality_and_exclusive_submission_author_confirmed'])
check('no_ethics_identifier_or_committee_determination_invented', author_declarations['approval_identifier'] is None and author_declarations['exemption_identifier'] is None and not author_declarations['committee_issued_determination_asserted'] and not author_declarations['institutional_policy_independently_verified'])
tables = [b for b in data['blocks'] if b['type'] == 'table']
figures = [b for b in data['blocks'] if b['type'] == 'figure']
check('nine_figures_twelve_tables', len(figures) == 9 and len(tables) == 12)
check('table_columns_complete', all(len(row) == len(b['headers']) for b in tables for row in b['rows']))
for b in figures:
    for ext in ['png', 'pdf', 'svg']: check(b['name'] + '_' + ext, (OUT / 'figures' / (b['name'] + '.' + ext)).stat().st_size > 1000)

doc = Document(OUT / 'manuscript_en.docx')
check('DOCX_tables_figures_complete', len(doc.tables) == len(tables) and len(doc.inline_shapes) == len(figures))
doc_text = '\n'.join(p.text for p in doc.paragraphs)
markdown = (OUT / 'manuscript_en.md').read_text()
tex = (OUT / 'manuscript.tex').read_text()
pdf = fitz.open(OUT / 'manuscript_en.pdf')
pdf_text = '\n'.join(p.get_text(clip=fitz.Rect(0, 0, p.rect.width, p.rect.height - 40)) for p in pdf)
def normalized(s): return re.sub(r'\s+', '', s).replace('\u00ad', '')
for b in data['blocks']:
    if b['type'] == 'paragraph':
        check('DOCX_paragraph_' + hashlib.sha256(b['text'].encode()).hexdigest()[:12], b['text'] in doc_text)
        check('Markdown_paragraph_' + hashlib.sha256(b['text'].encode()).hexdigest()[:12], b['text'] in markdown)
        check('PDF_paragraph_' + hashlib.sha256(b['text'].encode()).hexdigest()[:12], normalized(b['text']) in normalized(pdf_text))
check('PDF_all_pages_nonempty', all(len(p.get_text().strip()) > 30 for p in pdf))
check('PDF_author_email', data['authors'][0]['name'] in pdf_text and data['authors'][0]['email'] in pdf_text)
for statement in ['This research received no funding.', 'The author declares no competing interests.', 'The author confirms that neither ethics approval nor an exemption was required for this secondary analysis.']:
    check('author_statement_all_exports_' + hashlib.sha256(statement.encode()).hexdigest()[:12], all(statement in text for text in [markdown, doc_text, tex]) and normalized(statement) in normalized(pdf_text))
check('LaTeX_standalone_source', '\\begin{document}' in tex and '\\end{document}' in tex and '\\input{' not in tex and '\\includegraphics{' not in tex and '\\bibliography{' not in tex)
check('LaTeX_all_embedded_figures', tex.count('\\begin{figure}') == len(figures) and tex.count('\\begin{tikzpicture}') == len(figures))
check('no_obsolete_author_placeholders', 'Author details pending' not in markdown and ', ;' not in markdown)
check('JNE_abstract_under300', len(re.findall(r'\b[\w’-]+\b', data['blocks'][1]['text'])) <= 300)

people = pd.read_csv(OUT / 'tables/q15_external_subjects.csv')
summary = pd.read_csv(OUT / 'tables/q15_model_summary.csv')
seeds = pd.read_csv(OUT / 'tables/q15_seed_metrics.csv')
inventory = pd.read_csv(OUT / 'tables/completed_experiment_inventory.csv')
check('Q15_106_person_rows', len(people) == 106 and people.groupby('dataset').size().to_dict() == {'Cho2017': 52, 'Lee2019_MI': 54})
check('Q15_six_model_rows742_seed_rows', len(summary) == 6 and len(seeds) == 742)
check('complete_inventory105_rows', len(inventory) == 105)
check('unique_trials_correct', summary.groupby('dataset').n_trials.first().to_dict() == {'Cho2017': 10520, 'Lee2019_MI': 10800})
check('model_summary_contains_primary_FWER', abs(summary.iloc[0].primary_holm_p - 9.99950002499875e-5) < 1e-12 and abs(summary.iloc[3].primary_holm_p - .6036698165091745) < 1e-12)

required = ['supplementary_methods.md', 'supplementary_inventory.md', 'reproducibility_readme.md', 'journal_strategy_zh.md', 'cover_letter_draft.md', 'author_information_template.md', 'submission_checklist_zh.md', '中文说明.md']
for name in required: check(name, (OUT / name).stat().st_size > 300)

q16_dir = ROOT / 'research_runs/Q16-P001-BNCI-20261006'
q16_validation = load('evidence/q16_analysis/independent_validation.json')
q16_manifest = load('evidence/q16_analysis/run_manifest.json')
check('Q16_independent_validation_passed', q16_validation['passed'] and q16_validation['status'] == 'independent_raw_and_saved_evidence_validation_passed')
check('Q16_no_new_decoder_work', q16_validation['new_decoder_fits'] == q16_validation['new_checkpoint_inference'] == q16_manifest['new_decoder_fits'] == q16_manifest['new_checkpoint_inference'] == 0)
replay = q16_validation['raw_replay']
check('Q16_18files5184events228096rows', replay['raw_files_sha256_verified'] == 18 and replay['raw_event_identities_replayed'] == 5184 and replay['all_trial_channel_band_rows_checked'] == 228096)
check('Q16_manual_central_FFT_31104pairs62208powers', replay['manual_central_channel_band_pairs_replayed'] == 31104 and replay['manual_baseline_task_power_comparisons'] == 62208)
check('Q16_freeze_immutable_commit', q16_manifest['freeze']['git_commit'] == q16_validation['freeze_and_hash_gate']['freeze_commit'] == '050e01b028aaab8e3d745934b13b2d17e9bb0a7a')
for record in q16_manifest['outputs']:
    f = ROOT / record['path']
    check('Q16_output_hash_' + f.name, sha(f) == record['sha256'] and f.stat().st_size == record['bytes'])
    check('Q16_paper_evidence_copy_' + f.name, sha(OUT / 'evidence/q16_analysis' / f.name) == record['sha256'])
for name in ['raw_source_receipt.json', 'metadata_audit.json', 'preprocessing_freeze.json', 'independent_validation.json', 'run_manifest.json']:
    check('Q16_provenance_copy_' + name, sha(q16_dir / name) == sha(OUT / 'evidence/q16_analysis' / name))
for record in q16_manifest['freeze']['frozen_code'] + q16_manifest['freeze']['frozen_protocols'] + q16_manifest['freeze']['frozen_inputs']:
    check('Q16_frozen_file_' + record['path'], sha(ROOT / record['path']) == record['sha256'])
figure_receipt = load('evidence/figures_q16.json')
check('Q16_figures_visually_reviewed', figure_receipt['status'] == 'rendered_and_visually_reviewed')
for name, digest in figure_receipt['artifact_sha256'].items():
    check('Q16_figure_hash_' + name, sha(OUT / name) == digest)
check('Q16_all_people_and_no_color_clipping', figure_receipt['central_profile_points'] == 108 and figure_receipt['association_figure']['scatter_points'] == 54 and figure_receipt['visual_review']['fixed_color_clipped_values'] == 0)

for filename, expected_tables, expected_figures in [('manuscript_main_en.docx', 8, 6), ('supplementary_materials.docx', 6, 3)]:
    variant = Document(OUT / filename)
    check('split_export_' + filename, len(variant.tables) == expected_tables and len(variant.inline_shapes) == expected_figures)
    check('split_export_' + filename + '_nonempty_pdf', all(len(p.get_text().strip()) > 30 for p in fitz.open(OUT / filename.replace('.docx', '.pdf'))))
check('scientific_completion_markers_absent', all(marker not in markdown for marker in ['Q16_PENDING', 'PENDING_Q16', 'WORKFLOW_PENDING']))
original_paths = [p for p in subprocess.check_output(['git', 'diff', '--name-only', '-z', '124e1b02895b13657b11365d8360c2515800792b'], cwd=ROOT).decode('utf-8').split('\0') if p]
check('original_scientific_files_preserved', all(p.startswith(('research_runs/PAPER_FINAL_20261006/', 'research_runs/PAPER_FINAL_20261006-PUBLICATION/', 'research_runs/Q16-P001-BNCI-20261006/')) or p in {'README.md', 'scripts/q16_common.py', 'scripts/q16_metadata_audit.py', 'scripts/q16_bnci_analysis.py', 'scripts/validate_q16_independent.py'} for p in original_paths))

# Bind the exact reviewed sources and delivered exports, without leaking secrets.
input_hashes = {}
for name in ['manuscript_content.json', 'build_content.py', 'build_revision_content.py', 'export_manuscript.py', 'references.json', 'tables/q15_model_summary.csv', 'evidence/q15_numbers.json', 'evidence/internal_review.json', 'evidence/final_scientific_review.json', 'evidence/author_declarations.json', 'evidence/q16_analysis/independent_validation.json', 'evidence/q16_analysis/run_manifest.json', 'manuscript_en.md', 'manuscript_en.docx', 'manuscript_en.pdf', 'manuscript.tex', 'manuscript_main_en.docx', 'manuscript_main_en.pdf', 'supplementary_materials.docx', 'supplementary_materials.pdf']:
    input_hashes[name] = sha(OUT / name)
versions = {name: importlib.metadata.version(name) for name in ['numpy', 'pandas', 'scipy', 'matplotlib', 'python-docx', 'reportlab', 'PyMuPDF']}
report = {'status': 'passed', 'checked_at_utc': datetime.now(timezone.utc).isoformat(), 'checks': checks, 'failed_checks': [], 'pdf_pages': len(pdf),
          'visual_review': 'Title, workflow, quantitative result panels, Q16 sensor maps and supplementary association panels inspected; PDF text is complete.',
          'pdf_export_engine': 'ReportLab', 'native_latex_compilation': load('evidence/native_latex_check.json'),
          'input_sha256': input_hashes, 'runtime_distributions': versions,
          'new_fits': 0, 'new_checkpoint_inference': 0, 'raw_EEG_loaded_by_delivery_checker': False,
          'new_Q16_raw_signal_analysis_separately_validated': True,
          'original_scientific_files_modified': False, 'core_author_declarations_confirmed': True,
          'remaining_author_confirmations': [], 'author_finalization_record': 'evidence/author_finalization/author_confirmations.json', 'author_supplied_final_confirmations': True,
          'journal_submission_performed': False}
source = {'paper_base_commit': '7af1a137e2676a018e1e880ab076de6cae4ce30b',
          'previous_reviewed_paper_commit': 'ac75a339c8db2861ff8e7d072e50690c79c602c4',
          'q15_scientific_code_revision': '271af288a2f3863430ab80e3145c2dee9bd5571d',
          'q15_verified_results_commit': 'bc48b257eb44f412ad069f50d0f1a72a33c3c520',
          'q15_results_branch': 'q15/run-20261005T050511Z-migration-from-r2-b82ad79b',
          'paper_branch': 'paper/zero-calibration-q16-20261006',
          'q16_freeze_commit': '050e01b028aaab8e3d745934b13b2d17e9bb0a7a',
          'previous_completed_q15_paper_commit': 'e4d0303c39ab7104a99556cac23ae6175f03e091',
          'source_fit_count': 15, 'target_fit_count': 0, 'new_model_fits': 0, 'new_checkpoint_inference': 0,
          'input_evidence_sha256': input_hashes, 'raw_data_redistributed': False}
(OUT / 'evidence/source_snapshot.json').write_text(json.dumps(source, indent=2) + '\n')
check('review_bindings_still_match_after_generated_delivery_outputs', all(sha(OUT / name) == digest for name, digest in reviewed_hashes.items()))
(OUT / 'evidence/delivery_validation.json').write_text(json.dumps(report, indent=2, ensure_ascii=False) + '\n')
print(json.dumps({'status': report['status'], 'checks': len(checks), 'pages': len(pdf), 'new_fits': 0}))
