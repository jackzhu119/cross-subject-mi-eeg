"""Create consistent English Markdown, DOCX, PDF and standalone LaTeX drafts.

Only archived results are read. No EEG model, training, or prediction is called.
The PDF export is a document-layout export, not evidence of LaTeX compilation.
"""
from pathlib import Path
import csv
import hashlib
import html
import json
import re
import pandas as pd

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether, PageBreak

OUT=Path(__file__).resolve().parent
DATA=json.loads((OUT/'manuscript_content.json').read_text())
TITLE=DATA['title'];SUBTITLE=DATA['subtitle'];BLOCKS=DATA['blocks'];REFERENCES=DATA['references']
NUM=json.loads((OUT/'paper_numbers.json').read_text())
AUTHOR=DATA['authors'][0]['name']
AFFILIATION=DATA['authors'][0]['affiliation']
EMAIL=DATA['authors'][0]['email']

def markdown():
    lines=[f'# {TITLE}','',f'*{SUBTITLE}*','',AUTHOR + ' (corresponding author)', '', AFFILIATION, '', 'Correspondence: ' + EMAIL, '', 'Scientific draft for author review — 7 October 2026.','']
    for b in BLOCKS:
        kind=b['type']
        if kind=='heading':lines.extend(['#'*(b['level']+1)+' '+b['text'],''])
        elif kind in ('paragraph','equation'):lines.extend([b['text'],''])
        elif kind=='table':
            lines.extend(['**'+b['label']+'**','','| '+' | '.join(b['headers'])+' |','| '+' | '.join(['---']*len(b['headers']))+' |'])
            lines.extend('| '+' | '.join(row)+' |' for row in b['rows'])
            lines.extend(['',b['note'],''])
        elif kind=='figure':lines.extend([f'![{b["label"]}](figures/{b["name"]}.png)','',f'**{b["label"]}.** {b["caption"]}',''])
    lines.extend(['## References',''])
    lines.extend(f'{i}. {r}' for i,r in enumerate(REFERENCES,1))
    (OUT/'manuscript_en.md').write_text('\n'.join(lines)+'\n')

def word(filename='manuscript_en.docx'):
    doc=Document();s=doc.sections[0];s.top_margin=s.bottom_margin=Inches(.8)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(7)
    doc.add_heading(TITLE,0);doc.add_paragraph(SUBTITLE,'Subtitle')
    doc.add_paragraph(AUTHOR + ' (corresponding author)'); doc.add_paragraph(AFFILIATION); doc.add_paragraph('Correspondence: ' + EMAIL); doc.add_paragraph('Scientific draft for author review — 7 October 2026.')
    for b in BLOCKS:
        if b['type']=='heading':doc.add_heading(b['text'],b['level'])
        elif b['type'] in ('paragraph','equation'):doc.add_paragraph(b['text'])
        elif b['type']=='table':
            p=doc.add_paragraph();p.add_run(b['label']).bold=True
            t=doc.add_table(rows=1,cols=len(b['headers']));t.style='Light Shading Accent 1'
            for cell,text in zip(t.rows[0].cells,b['headers']):cell.text=text
            for row in b['rows']:
                for cell,text in zip(t.add_row().cells,row):cell.text=text
            doc.add_paragraph(b['note'])
        elif b['type']=='figure':
            doc.add_picture(str(OUT/'figures'/f'{b["name"]}.png'),width=Inches(6.3))
            p=doc.add_paragraph();p.add_run(b['label']+'. ').bold=True;p.add_run(b['caption'])
    doc.add_heading('References',1)
    for i,r in enumerate(REFERENCES,1):doc.add_paragraph(f'[{i}] {r}')
    footer=s.footer.paragraphs[0];footer.text='Working manuscript • Source-only MI EEG • 7 October 2026'
    doc.save(OUT/filename)

def pdf(filename='manuscript_en.pdf'):
    base=Path('/usr/share/fonts/truetype/dejavu')
    for name,file in [('Paper','DejaVuSerif.ttf'),('PaperBold','DejaVuSerif-Bold.ttf'),('PaperItalic','DejaVuSerif-Italic.ttf'),('Caption','DejaVuSans.ttf')]:
        pdfmetrics.registerFont(TTFont(name,str(base/file)))
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='PaperBody',fontName='Paper',fontSize=9.8,leading=14.1,alignment=TA_JUSTIFY,spaceAfter=7))
    styles.add(ParagraphStyle(name='PaperTitle',fontName='PaperBold',fontSize=19,leading=24,alignment=TA_CENTER,spaceAfter=10))
    styles.add(ParagraphStyle(name='PaperSubtitle',fontName='PaperItalic',fontSize=10,leading=14,alignment=TA_CENTER,spaceAfter=12))
    styles.add(ParagraphStyle(name='PaperH1',fontName='PaperBold',fontSize=13,leading=17,spaceBefore=13,spaceAfter=7,keepWithNext=True))
    styles.add(ParagraphStyle(name='PaperH2',fontName='PaperBold',fontSize=10.5,leading=15,spaceBefore=9,spaceAfter=5,keepWithNext=True))
    styles.add(ParagraphStyle(name='PaperTableTitle',fontName='PaperBold',fontSize=10.5,leading=15,spaceBefore=9,spaceAfter=5))
    styles.add(ParagraphStyle(name='PaperCaption',fontName='Caption',fontSize=8,leading=11,spaceAfter=9))
    styles.add(ParagraphStyle(name='PaperCell',fontName='Caption',fontSize=7.5,leading=10))
    story=[]
    def p(text,style='PaperBody'):return Paragraph(html.escape(text),styles[style])
    story.extend([p(TITLE,'PaperTitle'),p(SUBTITLE,'PaperSubtitle'),p(AUTHOR + ' (corresponding author)','PaperSubtitle'), p(AFFILIATION,'PaperSubtitle'), p('Correspondence: ' + EMAIL,'PaperSubtitle'), p('Scientific draft for author review • 7 October 2026','PaperCaption')])
    for b in BLOCKS:
        k=b['type']
        if k=='heading':
            if b['text']=='Supplementary material':story.append(PageBreak())
            story.append(p(b['text'],'PaperH1' if b['level']==1 else 'PaperH2'))
        elif k in ('paragraph','equation'):story.append(p(b['text']))
        elif k=='table':
            n=len(b['headers']);width=6.65*inch
            if n==6:weights=[1.25,1.0,1.0,1.0,2.0,.9] if b['label'].startswith(('Table 8.', 'Methods table M1.')) else [1.45,.62,1.05,.65,1.6,.6]
            elif n==5:weights=[1.3,.55,.55,.95,1.4]
            elif n==4:weights=[1.45,1.7,1.2,2.0] if 'Target information' in b['label'] else ([2.3,1.9,.7,.6] if 'S3.' in b['label'] else [2.9,1,1.45,1])
            elif n==3:weights=[1.5,2.2,3.0] if 'Matching and remaining' in b['label'] else [3.3,1.1,1.3]
            else:weights=[1]*n
            colwidth=[width*w/sum(weights) for w in weights]
            data=[[p(t,'PaperCell') for t in row] for row in [b['headers']]+b['rows']]
            t=Table(data,colWidths=colwidth,repeatRows=1,hAlign='LEFT')
            t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e7edf4')),('LINEBELOW',(0,0),(-1,0),.7,colors.HexColor('#607080')),('LINEBELOW',(0,-1),(-1,-1),.6,colors.HexColor('#607080')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f7f9fb')])]))
            story.extend([p(b['label'],'PaperTableTitle'),t,Spacer(1,5),p(b['note'],'PaperCaption')])
        elif k=='figure':
            image=Image(str(OUT/'figures'/f'{b["name"]}.png'));ratio=image.imageHeight/image.imageWidth
            image.drawWidth=6.65*inch;image.drawHeight=image.drawWidth*ratio
            story.append(KeepTogether([image,Spacer(1,5),p(b['label']+'. '+b['caption'],'PaperCaption')]))
    story.append(p('References','PaperH1'))
    for i,r in enumerate(REFERENCES,1):story.append(KeepTogether([p(f'[{i}] {r}','PaperCaption')]))
    def footer(c,doc):
        c.saveState();c.setFont('Caption',7);c.setFillColor(colors.HexColor('#607080'))
        c.drawString(.82*inch,.42*inch,'Source-only MI EEG • Working manuscript');c.drawRightString(7.48*inch,.42*inch,str(doc.page));c.restoreState()
    SimpleDocTemplate(str(OUT/filename),pagesize=(8.3*inch,11.7*inch),leftMargin=.82*inch,rightMargin=.83*inch,topMargin=.72*inch,bottomMargin=.7*inch,title=TITLE,author=AUTHOR).build(story,onFirstPage=footer,onLaterPages=footer)

def tex_escape(s):
    replacements={'\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_','{':r'\{','}':r'\}','~':r'\textasciitilde{}','^':r'\textasciicircum{}'}
    s=''.join(replacements.get(c,c) for c in s)
    for a,b in [('μ',r'$\mu$'),('β',r'$\beta$'),('Δ',r'$\Delta$'),('θ',r'$\theta$'),('Σ',r'$\Sigma$'),('±',r'$\pm$'),('ρ',r'$\rho$'),('α',r'$\alpha$'),('×',r'$\times$'),('−','-'),('–','--'),('—','---'),('’',"'"),('“','``'),('”',"''"),('≤',r'$\leq$'),('≥',r'$\geq$'),('→',r'$\to$'),('•',r'$\cdot$')]:s=s.replace(a,b)
    return s

def tex_inline_figure(name):
    if 'q16' in name or name == 'figure_workflow_zero_calibration':
        from tex_q16_figures import tex_inline_figure as render
    elif name in {'figure4_q15_external_transfer','figureS2_q15_participant_heterogeneity'}:
        from tex_q15_figures import tex_inline_figure as render
    else:
        from tex_figures import tex_inline_figure as render
    return render(name,OUT,NUM)

def latex():
    lines=[r'\documentclass[11pt]{article}',r'\usepackage[T1]{fontenc}',r'\usepackage[utf8]{inputenc}',r'\usepackage[margin=0.9in]{geometry}',r'\usepackage{amsmath,booktabs,array,longtable,graphicx,hyperref,pgfplots}',r'\usepgfplotslibrary{groupplots}',r'\pgfplotsset{compat=1.18}',r'\hypersetup{hidelinks}',r'\setlength{\emergencystretch}{3em}',r'\title{'+tex_escape(TITLE)+r'\\\large '+tex_escape(SUBTITLE)+'}',r'\author{'+tex_escape(AUTHOR)+r'\\\small '+tex_escape(AFFILIATION)+r'\\\small Correspondence: \texttt{'+tex_escape(EMAIL)+'}}',r'\date{Scientific draft for author review, 7 October 2026}',r'\begin{document}',r'\maketitle']
    for b in BLOCKS:
        k=b['type']
        if k=='heading':
            if b['text']=='Abstract':lines.append(r'\section*{Abstract}')
            else:lines.append((r'\section*{' if b['level']==1 else r'\subsection*{')+tex_escape(b['text'])+'}')
        elif k=='paragraph':lines.extend([tex_escape(b['text']),''])
        elif k=='equation':lines.extend([r'\begin{equation*}',b['latex'],r'\end{equation*}'])
        elif k=='table':
            n=len(b['headers'])
            if n==6:weights=[1.25,1.0,1.0,1.0,2.0,.9] if b['label'].startswith('Table 8.') else [1.45,.62,1.05,.65,1.6,.6]
            elif n==5:weights=[1.3,.55,.55,.95,1.4]
            elif n==4:weights=[2.3,1.9,.7,.6] if 'S3.' in b['label'] else [2.9,1,1.45,1]
            elif n==3:weights=[1.5,2.2,3.0] if 'Matching and remaining' in b['label'] else [3.3,1.1,1.3]
            else:weights=[1]*n
            spec='@{}'+''.join(r'>{\raggedright\arraybackslash}p{'+f'{.88*w/sum(weights):.6f}'+r'\linewidth}' for w in weights)+'@{}'
            header=' & '.join(tex_escape(t) for t in b['headers'])+r' \\'
            lines.extend([r'\par\medskip\noindent\textbf{'+tex_escape(b['label'])+'}',r'\begingroup\small\setlength{\tabcolsep}{4pt}',r'\begin{longtable}{'+spec+'}',r'\toprule',header,r'\midrule\endfirsthead',r'\toprule',header,r'\midrule\endhead',r'\bottomrule\endfoot',r'\bottomrule\endlastfoot'])
            lines.extend(' & '.join(tex_escape(t) for t in row)+r' \\' for row in b['rows'])
            lines.extend([r'\end{longtable}',r'\noindent\footnotesize '+tex_escape(b['note']),r'\par\endgroup\medskip'])
        elif k=='figure':
            lines.extend([r'\begin{figure}[htbp]',r'\centering',r'\resizebox{\linewidth}{!}{'+tex_inline_figure(b['name'])+'}',r'\caption*{\textbf{'+tex_escape(b['label'])+'}. '+tex_escape(b['caption'])+'}',r'\end{figure}'])
    # caption* is explicitly supplied; no external .bib or figure assets are needed.
    lines.insert(5,r'\usepackage{caption}')
    lines.extend([r'\section*{References}',r'\begin{enumerate}'])
    lines.extend(r'\item '+tex_escape(r) for r in REFERENCES)
    lines.extend([r'\end{enumerate}',r'\end{document}'])
    (OUT/'manuscript.tex').write_text('\n'.join(lines)+'\n')

if __name__ == '__main__':
    markdown();word();pdf();latex()
    word_count=sum(len(re.findall(r"\b[\w’-]+\b",b.get('text',''))) for b in BLOCKS if b['type']=='paragraph')
    report={'english_body_words':word_count,'references':len(REFERENCES),'figures':sum(b['type']=='figure' for b in BLOCKS),'tables':sum(b['type']=='table' for b in BLOCKS),'formats':['md','docx','pdf','tex'],'new_fits':0,'new_checkpoint_inference':0,'pdf_export_engine':'ReportLab','latex_compilation_status':'separate_native_check'}
    (OUT/'evidence/export_report.json').write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps(report))
