"""Submission/editorial integrity gates; reads saved artifacts, never EEG/models."""
from pathlib import Path
import hashlib,json,re,subprocess
from datetime import datetime,timezone
from docx import Document
from docx.oxml.ns import qn
import fitz
P=Path(__file__).resolve().parent
R=P.parents[1]
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def load(path):return json.loads((P/path).read_text())
checks={}
def check(name,value):
 checks[name]=bool(value)
 if not value:raise AssertionError(name)
protected=load('evidence/quality_update/protected_scientific_inputs.json')
for name,digest in protected['sha256'].items():
 if sha(R/name)!=digest:raise AssertionError('Frozen scientific bytes changed: '+name)
check('protected_scientific_bytes_unchanged',True)
check('correct_paper_branch',subprocess.check_output(['git','branch','--show-current'],cwd=R,text=True).strip()=='paper/zero-calibration-q16-20261006')
c=load('manuscript_content.json'); refs=load('references.json')['references']
check('exact_pipeline_title',c['title']=='Source-only model selection and limits of fixed spectral-sharing pipelines in cross-subject motor-imagery EEG decoding')
check('all_28_refs_cited_and_unique',len(refs)==len(c['references'])==len(set(c['reference_keys']))==28)
check('reference_order_catalog_correspondence',set(c['reference_keys'])=={r['key'] for r in refs})
used=set()
for b in c['blocks']:
 for field in ['text','caption','note']:
  for match in re.finditer(r'\[(\d+(?:,\d+)*)\]',b.get(field,'')):used.update(map(int,match.group(1).split(',')))
check('continuous_cited_reference_numbers',used==set(range(1,29)))
check('every_reference_verified',all(r.get('verified') for r in refs))
for key,doi in [('Zhong2025EEGDG','10.1109/jbhi.2024.3431230'),('Zheng2025DG','10.3390/bioengineering12050495')]:
 r=next(r for r in refs if r['key']==key)
 m=load('evidence/quality_update/crossref_'+doi.replace('/','_').replace('jbhi','JBHI')+'.json')
 check('new_reference_metadata_'+key,r['doi'].lower()==doi and r['title']==m['title'][0] and r['year']==2025 and r['volume']==m['volume'] and r['pages_or_article']==m['page'] and r['authors']==[a['given']+' '+a['family'] for a in m['author']])
check('EEGPT_official_spelling_preserved',next(r for r in refs if r['key']=='Wang2024EEGPT')['authors'][0]=='Guagnyu Wang')
check('no_empty_author_records_or_punctuation',all(all(a.strip(' ,;')==a and a for a in r['authors']) for r in refs))
main=[]
for b in c['blocks']:
 if b.get('text')=='Supplementary material':break
 main.append(b)
text='\n'.join(b.get('text','')+b.get('caption','')+b.get('note','') for b in main)
check('no_project_draft_title_metadata','Scientific draft for author review' not in (P/'manuscript_en.md').read_text())
check('unequal_source_grouping_limitation','2/2/2/3' in text and 'equal fold weighting' in text and 'Grouping sensitivity was not evaluated prospectively' in text)
check('Q16_scope_limits','no p-values' in text and 'No external physiology was calculated' in text and 'does not alone establish absolute contralateral ERD' in text and '[0.5,3.5)' in text and '[0.5,2.5)' in text)
workflow=(P/'figures/figure_workflow_zero_calibration.svg').read_text()
check('workflow_baseline_relative_not_ERD_ERS','Baseline-relative' in workflow and 'ERD / ERS' not in workflow)
check('workflow_information_roles','Event/class metadata for eligibility' in workflow and 'Ground-truth labels used for scoring only' in workflow and 'No target fitting or model selection' in workflow)
author_finalization=load('evidence/author_finalization/author_confirmations.json')
check('AI_disclosure_actual_scope_and_author_reported_review','AI assistance disclosure' in text and 'OpenAI Codex assisted' in text and 'ChatGPT (GPT-6, as reported by the author)' in text and 'personally reviewed and revised' in text and author_finalization['author_final_review_and_approval_confirmed'] and not author_finalization['historical_chatgpt_model_identity_independently_verified'] and not author_finalization['personal_verification_of_every_code_line_or_original_reference_asserted'])
check('author_correspondence_postcode_and_confirmations',c['authors'][0]['postal_address']==author_finalization['postal_address_en'] and '401331' in c['authors'][0]['postal_address'] and author_finalization['target_journal']=='Journal of Neural Engineering' and not author_finalization['journal_or_preprint_publication_reported'] and not author_finalization['other_journal_consideration_reported'])
for name in ['manuscript_en.docx','manuscript_main_en.docx','supplementary_materials.docx']:
 d=Document(P/name)
 for i,shape in enumerate(d.inline_shapes):
  sec=d.sections[0]
  check(name+'_image_'+str(i)+'_within_width',shape.width<=sec.page_width-sec.left_margin-sec.right_margin)
 check(name+'_repeated_table_headers',all(t.rows[0]._tr.find('./'+qn('w:trPr')+'/'+qn('w:tblHeader')) is not None for t in d.tables))
 check(name+'_rows_not_split',all(row._tr.find('./'+qn('w:trPr')+'/'+qn('w:cantSplit')) is not None for t in d.tables for row in t.rows))
 check(name+'_page_number_field',any('PAGE' in f._element.xml for sec in d.sections for f in [sec.footer]))
 check(name+'_no_draft_metadata',all('Scientific draft for author review' not in p.text for p in d.paragraphs) and all('Working manuscript' not in sec.footer._element.xml for sec in d.sections))
 pdf=fitz.open(P/name.replace('.docx','.pdf'))
 check(name+'_pdf_all_pages_have_content',all(len(p.get_text().strip())>30 for p in pdf))
 reporttext='\n'.join(p.get_text() for p in pdf)
 check(name+'_pdf_no_draft_metadata','Scientific draft for author review' not in reporttext and 'Working manuscript' not in reporttext)
numeric=load('evidence/quality_update/numeric_review.json')
check('saved_numeric_review_passed',numeric['status']=='passed_saved_evidence_numeric_review' and not numeric['discrepancies'])
for name,digest in numeric['input_sha256'].items():check('numeric_review_current_'+name,sha(P/name)==digest)
report={'status':'submission_revision_integrity_passed','checked_at_utc':datetime.now(timezone.utc).isoformat(),'checks':checks,'failed_checks':[],'protected_scientific_file_count':protected['protected_file_count'],'raw_EEG_loaded':False,'new_fits':0,'new_checkpoint_inference':0,'experiment_reruns':0,'native_TeX_compilation_confirmed':False}
(P/'evidence/quality_update/submission_integrity.json').write_text(json.dumps(report,indent=2)+'\n')
print(json.dumps({'status':report['status'],'checks':len(checks),'protected_files':protected['protected_file_count']}))
