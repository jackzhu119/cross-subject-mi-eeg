#!/usr/bin/env python3
"""Export authored publication text; no EEG, model, or statistical execution.

The editable main document retains its figure captions and omits image objects.
The review copy appends the already generated figures, in numerical order.
LibreOffice and the document libraries must already be installed.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import importlib.metadata
import json
from pathlib import Path
import re
import subprocess
import tempfile
import shutil
import unicodedata

import fitz
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOFFICE = Path(shutil.which("libreoffice") or shutil.which("soffice") or "/missing")
FONT = 'Liberation Serif'  # Installed Times-compatible font; avoids substitution drift.
AUTHOR_GIVEN_NAME, AUTHOR_FAMILY_NAME = 'Ziyuan', 'Zhu'
AUTHOR_DISPLAY_NAME = f'{AUTHOR_GIVEN_NAME} {AUTHOR_FAMILY_NAME}'


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def plain_inline(text: str) -> str:
    text = re.sub(r'!\[([^]]*)\]\([^)]*\)', '', text)
    text = re.sub(r'\[([^]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    text = re.sub(r'</?(?:sup|sub)>', '', text)
    return re.sub(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)',
                  lambda match: match[0][2:-2] if match[0].startswith('**') else match[0][1:-1], text)


def inline(paragraph, text: str, *, size: int | None = None) -> None:
    text = re.sub(r'<br\s*/?>', '\n', text, flags=re.I)
    text = re.sub(r'\[([^]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    pattern = r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*|<sup>.*?</sup>|<sub>.*?</sub>)'
    for token in re.split(pattern, text):
        if not token:
            continue
        bold = token.startswith('**') and token.endswith('**')
        italic = token.startswith('*') and token.endswith('*') and not bold
        code = token.startswith('`') and token.endswith('`')
        superscript = token.startswith('<sup>') and token.endswith('</sup>')
        subscript = token.startswith('<sub>') and token.endswith('</sub>')
        if bold:
            token = token[2:-2]
        elif italic or code:
            token = token[1:-1]
        elif superscript or subscript:
            token = token[5:-6]
        # Word otherwise treats long URLs and repository paths as unbroken words.
        if 'https://' in token or any(len(part) > 55 for part in token.split()):
            token = re.sub(r'([/_-])', lambda match: match[1] + '\u200b', token)
        run = paragraph.add_run(token)
        run.bold, run.italic = (True if bold else None), (True if italic else None)
        run.font.superscript, run.font.subscript = superscript, subscript
        if size is not None:
            run.font.size = Pt(size)
        if code:
            run.font.name = 'DejaVu Sans Mono'
            run.font.size = Pt(min(size or 10, 10))


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    start = OxmlElement('w:fldChar'); start.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.text = instruction
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    run._r.extend([start, instr, end])


def set_style_font(style, size: int) -> None:
    style.font.name, style.font.size = FONT, Pt(size)
    fonts = style._element.get_or_add_rPr().rFonts
    # The stock template contains theme overrides that LibreOffice otherwise
    # prioritizes over the explicitly declared typeface for linked headings.
    for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme', 'csTheme']:
        fonts.attrib.pop(qn('w:' + key), None)
    fonts.set(qn('w:ascii'), FONT)
    fonts.set(qn('w:hAnsi'), FONT)
    fonts.set(qn('w:cs'), FONT)
    fonts.set(qn('w:eastAsia'), 'Noto Serif CJK SC')


def configure_document(source: Path, *, cover: bool = False) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.8)
    section.header_distance = section.footer_distance = Inches(0.35)
    normal = doc.styles['Normal']
    set_style_font(normal, 12)
    normal.paragraph_format.line_spacing = 1.15 if cover else 2
    normal.paragraph_format.space_after = Pt(6 if cover else 0)
    normal.paragraph_format.widow_control = True
    for name, size in [('Title', 15), ('Heading 1', 13), ('Heading 2', 12), ('Heading 3', 12)]:
        style = doc.styles[name]
        set_style_font(style, size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.line_spacing = 1.2 if name == 'Title' else 2
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        if name == 'Title':
            borders = style._element.find('w:pPr/w:pBdr', namespaces=style._element.nsmap)
            if borders is not None:
                borders.getparent().remove(borders)
    for name in ['List Bullet', 'List Number']:
        set_style_font(doc.styles[name], 12)
        doc.styles[name].paragraph_format.line_spacing = 1.15 if cover else 2
        doc.styles[name].paragraph_format.space_after = Pt(0)
    title = next((line.lstrip('# ').strip() for line in source.read_text().splitlines() if line.startswith('# ')), source.stem)
    props = doc.core_properties
    props.title, props.author, props.subject = plain_inline(title), AUTHOR_DISPLAY_NAME, 'PLOS ONE manuscript materials'
    props.keywords = 'motor imagery; EEG; source-only selection; source-cohort expansion'
    props.created = props.modified = datetime.now(timezone.utc)
    # Continuous line numbering is a section setting, not text inserted into content.
    if not cover:
        numbering = OxmlElement('w:lnNumType')
        numbering.set(qn('w:countBy'), '1')
        numbering.set(qn('w:start'), '1')
        numbering.set(qn('w:restart'), 'continuous')
        numbering.set(qn('w:distance'), '180')
        section._sectPr.append(numbering)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('Page '); add_field(footer, 'PAGE')
    for run in footer.runs:
        run.font.name, run.font.size = FONT, Pt(10)
    return doc


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        values = [value.strip().replace(r'\|', '|') for value in re.split(r'(?<!\\)\|', line.strip().strip('|'))]
        if all(re.fullmatch(r':?-{3,}:?', value.replace(' ', '')) for value in values):
            continue
        rows.append(values)
    ncols = max(map(len, rows), default=0)
    if not ncols or any(len(row) != ncols for row in rows):
        raise ValueError('Malformed Markdown table; row widths must agree')
    return rows


def column_widths(rows: list[list[str]], usable: float) -> list[float]:
    count = len(rows[0])
    if count == 5 and rows[0][:2] == ['Expansion', 'Pipeline']:
        # Fixed-duration control: retain the full MU_BETA_SHARED identifier.
        return [usable * fraction for fraction in [.13, .24, .22, .24, .17]]
    if count == 1:
        return [usable]
    if count == 2:
        return [usable * .31, usable * .69]
    # Long contrast labels need extra width; the remaining numeric cells stay equal.
    first = max(len(plain_inline(row[0])) for row in rows)
    if first > 24:
        weight = .35 if count <= 4 else .28
        return [usable * weight] + [usable * (1 - weight) / (count - 1)] * (count - 1)
    return [usable / count] * count


def add_table(doc, lines: list[str]) -> None:
    rows = parse_table(lines)
    section = doc.sections[-1]
    usable = (section.page_width - section.left_margin - section.right_margin) / 914400
    widths = column_widths(rows, usable)
    table = doc.add_table(rows=1, cols=len(widths))
    table.style, table.alignment, table.autofit = 'Table Grid', WD_TABLE_ALIGNMENT.CENTER, False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    font_size = 10 if len(widths) >= 5 else 11
    for ri, row in enumerate(rows):
        cells = table.rows[0].cells if ri == 0 else table.add_row().cells
        for ci, value in enumerate(row):
            cells[ci].width, cells[ci].vertical_alignment = Inches(widths[ci]), WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[ci].paragraphs[0]
            p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(3)
            inline(p, value, size=font_size)
            if ri == 0:
                for run in p.runs:
                    run.bold = True
        properties = table.rows[ri]._tr.get_or_add_trPr()
        properties.append(OxmlElement('w:cantSplit'))
        if ri == 0:
            properties.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph()


def figure_captions(text: str) -> dict[int, str]:
    result = {}
    paragraphs = re.split(r'\n\s*\n', text)
    for index, paragraph in enumerate(paragraphs):
        cleaned = re.sub(r'^#{1,6}\s+', '', paragraph.strip())
        match = re.match(r'(?:\*\*)?(?:Figure|Fig\.?)\s*(\d+)\b', cleaned)
        if match:
            if paragraph.strip().startswith('#'):
                following = []
                for block in paragraphs[index + 1:]:
                    if block.strip().startswith('#'):
                        break
                    following.append(block)
                cleaned += ' ' + ' '.join(following)
            result[int(match[1])] = ' '.join(cleaned.splitlines())
    return result


def image_paths() -> list[tuple[int, Path]]:
    choices = {}
    for path in sorted((ROOT / 'figures').glob('*.png')):
        match = re.match(r'(?:fig(?:ure)?)[_-]?(\d+)', path.name, re.I)
        if match:
            number = int(match[1])
            if number in choices:
                raise ValueError(f'Multiple PNG assets for Figure {number}')
            choices[number] = path
    return sorted(choices.items())


def render(source: Path, destination: Path, *, append_figures: bool = False, cover: bool = False) -> list[Path]:
    doc = configure_document(source, cover=cover)
    text = source.read_text(encoding='utf-8')
    lines, index = text.splitlines(), 0
    while index < len(lines):
        line = lines[index].strip()
        if not line or line in ['---', '***'] or line.startswith('<!--'):
            index += 1; continue
        if line.startswith('```'):
            block = []; index += 1
            while index < len(lines) and not lines[index].strip().startswith('```'):
                block.append(lines[index]); index += 1
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1
            for part in block:
                inline(p, f'`{part}`', size=9)
                p.add_run('\n')
            index += 1; continue
        if line.startswith('|'):
            block = []
            while index < len(lines) and lines[index].strip().startswith('|'):
                block.append(lines[index]); index += 1
            add_table(doc, block); continue
        heading = re.match(r'^(#{1,6})\s+(.*)', line)
        if heading:
            if heading[2] == 'Abstract':
                doc.add_page_break()
            level = len(heading[1])
            p = doc.add_paragraph(style='Title' if level == 1 and index == 0 else f'Heading {min(max(level-1, 1), 3)}')
            inline(p, heading[2]); index += 1; continue
        if line.startswith('!['):
            # Figure captions remain in the text. Upload-ready main has no image objects.
            index += 1; continue
        bullet = re.match(r'^([-*]|\d+[.)])\s+(.*)', line)
        if bullet:
            numbered = bullet[1][0].isdigit()
            p = doc.add_paragraph(style=None if numbered else 'List Bullet')
            inline(p, line if numbered else bullet[2]); index += 1; continue
        paragraph = [line.lstrip('> ') + ('\n' if lines[index].endswith('  ') else ' ')]; index += 1
        while index < len(lines) and lines[index].strip() and not re.match(r'^(#{1,6}\s|\||```|[-*]\s|\d+[.)]\s|!\[)', lines[index].strip()):
            paragraph.append(lines[index].strip().lstrip('> ') + ('\n' if lines[index].endswith('  ') else ' ')); index += 1
        p = doc.add_paragraph(); inline(p, ''.join(paragraph).strip())
    figure_inputs = []
    if append_figures:
        captions, paths = figure_captions(text), image_paths()
        if not paths or set(captions) != {number for number, _ in paths}:
            raise ValueError(f'Figure caption/asset mismatch: captions={sorted(captions)}, images={[n for n, _ in paths]}')
        for number, path in paths:
            doc.add_page_break()
            section = doc.sections[-1]
            width = (section.page_width - section.left_margin - section.right_margin) / 914400
            with fitz.open(path) as image:
                rect = image[0].rect
                height = width * rect.height / rect.width
            # Keep a long caption on the same page where feasible.
            if height > 5.7:
                width *= 5.7 / height
            p = doc.add_paragraph()
            p.alignment = 1
            p.add_run().add_picture(str(path), width=Inches(width))
            p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.15
            inline(p, captions[number], size=10)
            figure_inputs.append(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)
    return figure_inputs


def pdf_export(path: Path) -> Path:
    if not SOFFICE.is_file():
        raise RuntimeError('Existing LibreOffice executable unavailable; no package installation is attempted')
    with tempfile.TemporaryDirectory(prefix='paper1-plos-docx-') as temp:
        completed = subprocess.run([str(SOFFICE), f'-env:UserInstallation=file://{temp}', '--headless',
                                    '--convert-to', 'pdf', '--outdir', str(path.parent), str(path)],
                                   check=True, capture_output=True, text=True, timeout=180)
    destination = path.with_suffix('.pdf')
    if not destination.is_file() or not destination.stat().st_size:
        raise RuntimeError(f'No PDF produced for {path.name}: {completed.stdout[:500]}')
    return destination


def normalized(text: str) -> str:
    text = unicodedata.normalize('NFKC', text).replace('\u200b', '').replace('\u00ad', '')
    return re.sub(r'\s+', '', text)


def content_tokens(text: str) -> Counter:
    text = unicodedata.normalize('NFKC', text).replace('\u200b', '').replace('\u00ad', '')
    # Chinese administrative documents can wrap within an unspaced sentence.
    # Compare every Han character while retaining word-level checks for Latin text.
    text = re.sub(r'([\u3400-\u4dbf\u4e00-\u9fff])', r' \1 ', text)
    return Counter(re.findall(r'[^\W_]+', text, flags=re.UNICODE))


def document_checks(source: Path, docx_path: Path, pdf_path: Path, *, review: bool = False, cover: bool = False) -> dict:
    doc = Document(docx_path)
    import zipfile
    with zipfile.ZipFile(docx_path) as archive:
        xml = archive.read('word/document.xml').decode()
        media = [name for name in archive.namelist() if name.startswith('word/media/')]
    with fitz.open(pdf_path) as pdf:
        empty, outside, page_text = [], [], []
        for number, page in enumerate(pdf, 1):
            text = page.get_text()
            page_text.append(text)
            if not text.strip() and not page.get_images():
                empty.append(number)
            for block in page.get_text('dict')['blocks']:
                if block['type'] != 0:
                    continue
                for line in block['lines']:
                    x0, y0, x1, y1 = line['bbox']
                    if x0 < 0 or y0 < 0 or x1 > page.rect.width + .5 or y1 > page.rect.height + .5:
                        outside.append({'page': number, 'bbox': [round(v, 2) for v in [x0,y0,x1,y1]]})
        text = '\n'.join(page_text)
        pages = len(pdf)
    # Check every non-code authored paragraph/cell as a whole in the rendered DOCX.
    # PDF text is also checked after ignoring inserted line/page-number tokens.
    authored = source.read_text(encoding='utf-8')
    doc_text = '\n'.join([p.text for p in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
    # PDF conversion must retain the editable document's lexical content. Added
    # page/line numbers and duplicated review captions may add tokens, never remove them.
    missing_pdf_tokens = content_tokens(doc_text) - content_tokens(text)
    missing = []
    for line in authored.splitlines():
        line = line.strip()
        if not line or line.startswith(('![', '```', '<!--')) or line in ['---', '***']:
            continue
        if line.startswith('|'):
            rows = parse_table([line]) if not re.match(r'^\|?\s*:?-{3,}', line) else []
            candidates = rows[0] if rows else []
        else:
            candidates = [re.sub(r'^#{1,6}\s+|^[-*]\s+', '', line)]
        for candidate in candidates:
            candidate = plain_inline(candidate)
            if normalized(candidate) and normalized(candidate) not in normalized(doc_text):
                missing.append(candidate[:160])
    forbidden = ['RESEARCH WORK VERSION', 'INDEPENDENT CONFIRMATION PENDING', 'Scientific draft for author review']
    report = {'pages': pages, 'empty_pages': empty, 'text_outside_page': outside,
              'authored_lines_missing_from_docx': missing, 'text_characters_pdf': len(text),
              'content_tokens_missing_from_pdf': dict(missing_pdf_tokens),
              'tables': len(doc.tables), 'image_objects': len(media),
              'continuous_line_numbering': 'w:lnNumType' in xml and 'w:restart="continuous"' in xml,
              'main_body_double_spaced': doc.styles['Normal'].paragraph_format.line_spacing == 2,
              'author_core_metadata': doc.core_properties.author,
              'author_name_metadata_verified': doc.core_properties.author == AUTHOR_DISPLAY_NAME,
              'author_display_name_in_text': AUTHOR_DISPLAY_NAME in doc_text,
              'forbidden_stale_export_labels': [word for word in forbidden if word in text],
              'source_sha256': sha(source), 'docx_sha256': sha(docx_path), 'pdf_sha256': sha(pdf_path)}
    if empty or outside or missing or missing_pdf_tokens or report['forbidden_stale_export_labels']:
        raise ValueError(f'Export verification failed: {json.dumps(report)}')
    if not cover and not report['continuous_line_numbering']:
        raise ValueError('Continuous line numbering missing')
    if not report['author_name_metadata_verified']:
        raise ValueError('Author name metadata differs from confirmed given/family-name order')
    if not review and not cover and media:
        raise ValueError('Upload-ready source documents must not embed image objects')
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--only', choices=['main', 'review', 'supplement', 'cover', 'ethics', 'response', 'fields'])
    args = parser.parse_args()
    jobs = [
        ('main', ROOT/'manuscript_en.md', ROOT/'main_manuscript.docx', False, False),
        ('review', ROOT/'manuscript_en.md', ROOT/'review_manuscript.docx', True, False),
        ('supplement', ROOT/'supplementary_en.md', ROOT/'S1_Appendix.docx', False, False),
        ('cover', ROOT/'submission'/'Cover_Letter.md', ROOT/'submission'/'Cover_Letter.docx', False, True),
        ('ethics', ROOT/'submission'/'Ethics_Author_Statement.md', ROOT/'submission'/'Ethics_Author_Statement.docx', False, True),
        ('response', ROOT/'Review_Response_zh.md', ROOT/'Review_Response_zh.docx', False, True),
        ('fields', ROOT/'submission'/'Submission_Fields.md', ROOT/'submission'/'Submission_Fields.docx', False, True),
    ]
    receipt_path = ROOT/'evidence'/'build_receipt.json'
    prior = json.loads(receipt_path.read_text()) if args.only and receipt_path.exists() else {}
    inputs, outputs, checks = {}, {}, prior.get('document_checks', {})
    for name, source, destination, review, cover in jobs:
        if args.only and args.only != name:
            continue
        if not source.is_file():
            raise FileNotFoundError(f'Awaiting authored source {source.relative_to(ROOT)}')
        figures = render(source, destination, append_figures=review, cover=cover)
        pdf = pdf_export(destination)
        checks[name] = document_checks(source, destination, pdf, review=review, cover=cover)
        for path in [source] + figures:
            inputs[str(path.relative_to(ROOT))] = sha(path)
        for path in [destination, pdf]:
            outputs[str(path.relative_to(ROOT))] = sha(path)
    receipt = {'schema_version': 1, 'status': 'EDITORIAL_EXPORTS_BUILT_AND_PROGRAMMATICALLY_CHECKED',
               'built_at_utc': datetime.now(timezone.utc).isoformat(),
               'scope': 'Authored text and existing figure rendering only; no EEG or model execution or new tests of scientific hypotheses',
               'model_fits': 0, 'prediction_inference_runs': 0,
               'generator_sha256': sha(Path(__file__).resolve()),
               'fonts': {'body': FONT, 'east_asian': 'Noto Serif CJK SC',
                         'Times_New_Roman_installed_font_substitute': 'Liberation Serif',
                         'Times_New_Roman_fontconfig_resolution': subprocess.run(
                             ['fc-match', '-f', '%{family}\\n', 'Times New Roman'],
                             check=True, capture_output=True, text=True).stdout.strip(),
                         'font_note': 'Times New Roman is not installed in this rendering environment; the actual installed Times-compatible Liberation Serif is declared, not relabelled.'},
               'author_name_check': {'given_name': AUTHOR_GIVEN_NAME, 'family_name': AUTHOR_FAMILY_NAME,
                                     'display_name': AUTHOR_DISPLAY_NAME,
                                     'basis': 'Author-confirmed romanized name; every DOCX core-author value checked after export',
                                     'all_exported_metadata_match': all(v['author_name_metadata_verified'] for v in checks.values())},
               'render_versions': {name: importlib.metadata.version(name) for name in ['python-docx', 'PyMuPDF']},
               'input_sha256': {**prior.get('input_sha256', {}), **inputs},
               'output_sha256': {**prior.get('output_sha256', {}), **outputs}, 'document_checks': checks,
               'visual_review': 'Separate rendered-page inspection required; automated bounding-box checks do not prove typographic perfection'}
    receipt_path.parent.mkdir(parents=True, exist_ok=True)
    receipt_path.write_text(json.dumps(receipt, indent=2) + '\n')
    print(json.dumps({'status': receipt['status'], 'document_checks': checks}, indent=2))


if __name__ == '__main__':
    main()
