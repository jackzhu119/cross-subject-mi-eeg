"""Check manuscript consistency and saved evidence; no EEG or model execution."""
from pathlib import Path
import hashlib
import importlib.metadata
import json
import re
import subprocess

import fitz
import pandas as pd
from docx import Document

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
def load(name): return json.loads((OUT / name).read_text())
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
data = load('manuscript_content.json')
q15 = load('evidence/q15_numbers.json')
internal = load('evidence/internal_review.json')
review = load('evidence/final_scientific_review.json')
checks = {}
def check(name, condition):
    checks[name] = bool(condition)
    if not condition: raise AssertionError(name)

check('manuscript_preparation_no_fit_or_checkpoint_inference', data['new_fits'] == data['new_checkpoint_inference'] == 0)
check('completed_q15_independent_arithmetic', q15['passed'])
check('original_source_fits15_new_source0_target0', q15['original_source_fits'] == 15 and q15['new_source_fits'] == q15['target_fits'] == 0)
check('earlier_saved_prediction_review', internal['arm_count_recomputed'] == 38 and internal['prediction_rows_rescored_sum'] == 809434)
check('final_scientific_review', review['status'] == 'scientific_and_numerical_review_passed' and not review['checks_failed'])
refs = load('references.json')['references']
check('22_verified_unique_references', len(refs) == len({r['key'] for r in refs}) == 22 and all(r['verified'] for r in refs))
check('22_manuscript_references_cited', len(data['references']) == len(set(data['reference_keys'])) == 22)
check('author_correspondence', data['authors'][0]['name'] == 'Ziyuan Zhu' and data['authors'][0]['email'] == 'zzy2630816871@gmail.com' and data['authors'][0]['corresponding'])
check('published_affiliation_usage', data['authors'][0]['affiliation'] == load('evidence/affiliation_verification.json')['recommended_affiliation_en'])
tables = [b for b in data['blocks'] if b['type'] == 'table']
figures = [b for b in data['blocks'] if b['type'] == 'figure']
check('six_figures_eight_tables', len(figures) == 6 and len(tables) == 8)
check('table_columns_complete', all(len(row) == len(b['headers']) for b in tables for row in b['rows']))
for b in figures:
    for ext in ['png', 'pdf', 'svg']: check(b['name'] + '_' + ext, (OUT / 'figures' / (b['name'] + '.' + ext)).stat().st_size > 1000)

doc = Document(OUT / 'manuscript_en.docx')
check('DOCX_tables_figures_complete', len(doc.tables) == 8 and len(doc.inline_shapes) == 6)
doc_text = '\n'.join(p.text for p in doc.paragraphs)
markdown = (OUT / 'manuscript_en.md').read_text()
tex = (OUT / 'manuscript.tex').read_text()
pdf = fitz.open(OUT / 'manuscript_en.pdf')
pdf_text = '\n'.join(p.get_text(clip=fitz.Rect(0, 0, p.rect.width, p.rect.height - 40)) for p in pdf)
def normalized(s): return re.sub(r'\s+', '', s).replace('\u00ad', '')
for b in data['blocks']:
    if b['type'] == 'paragraph':
        check('DOCX_paragraph_' + hashlib.sha256(b['text'].encode()).hexdigest()[:12], b['text'] in doc_text)
        check('Markdown_paragraph_' + hashlib.sha256(b['text'].encode()).hexdigest()[:12], b['text'] in markdown)
        check('PDF_paragraph_' + hashlib.sha256(b['text'].encode()).hexdigest()[:12], normalized(b['text']) in normalized(pdf_text))
check('PDF_all_pages_nonempty', all(len(p.get_text().strip()) > 30 for p in pdf))
check('PDF_author_email', data['authors'][0]['name'] in pdf_text and data['authors'][0]['email'] in pdf_text)
check('LaTeX_standalone_source', '\\begin{document}' in tex and '\\end{document}' in tex and '\\input{' not in tex and '\\includegraphics{' not in tex and '\\bibliography{' not in tex)
check('LaTeX_six_embedded_figures', tex.count('\\begin{figure}') == 6 and tex.count('\\begin{tikzpicture}') == 6)
check('no_obsolete_author_placeholders', 'Author details pending' not in markdown and ', ;' not in markdown)
check('JNE_abstract_under300', len(re.findall(r'\b[\w’-]+\b', data['blocks'][1]['text'])) <= 300)

people = pd.read_csv(OUT / 'tables/q15_external_subjects.csv')
summary = pd.read_csv(OUT / 'tables/q15_model_summary.csv')
seeds = pd.read_csv(OUT / 'tables/q15_seed_metrics.csv')
inventory = pd.read_csv(OUT / 'tables/completed_experiment_inventory.csv')
check('Q15_106_person_rows', len(people) == 106 and people.groupby('dataset').size().to_dict() == {'Cho2017': 52, 'Lee2019_MI': 54})
check('Q15_six_model_rows742_seed_rows', len(summary) == 6 and len(seeds) == 742)
check('complete_inventory104_rows', len(inventory) == 104)
check('unique_trials_correct', summary.groupby('dataset').n_trials.first().to_dict() == {'Cho2017': 10520, 'Lee2019_MI': 10800})
check('model_summary_contains_primary_FWER', abs(summary.iloc[0].primary_holm_p - 9.99950002499875e-5) < 1e-12 and abs(summary.iloc[3].primary_holm_p - .6036698165091745) < 1e-12)

required = ['supplementary_methods.md', 'supplementary_inventory.md', 'reproducibility_readme.md', 'journal_strategy_zh.md', 'cover_letter_draft.md', 'author_information_template.md', 'submission_checklist_zh.md', '中文说明.md']
for name in required: check(name, (OUT / name).stat().st_size > 300)

# Bind the exact reviewed sources and delivered exports, without leaking secrets.
input_hashes = {}
for name in ['manuscript_content.json', 'build_content.py', 'export_manuscript.py', 'references.json', 'tables/q15_model_summary.csv', 'evidence/q15_numbers.json', 'evidence/internal_review.json', 'evidence/final_scientific_review.json', 'manuscript_en.md', 'manuscript_en.docx', 'manuscript_en.pdf', 'manuscript.tex']:
    input_hashes[name] = sha(OUT / name)
versions = {name: importlib.metadata.version(name) for name in ['numpy', 'pandas', 'scipy', 'matplotlib', 'python-docx', 'reportlab', 'PyMuPDF']}
report = {'status': 'passed', 'checks': checks, 'failed_checks': [], 'pdf_pages': len(pdf),
          'visual_review': 'Title/abstract, results figure page, discussion and Q15 figures inspected; PDF text is complete.',
          'pdf_export_engine': 'ReportLab', 'native_latex_compilation': load('evidence/native_latex_check.json'),
          'input_sha256': input_hashes, 'runtime_distributions': versions,
          'new_fits': 0, 'new_checkpoint_inference': 0, 'raw_EEG_loaded': False,
          'original_scientific_files_modified': False, 'author_declarations_pending': True,
          'journal_submission_performed': False}
(OUT / 'evidence/delivery_validation.json').write_text(json.dumps(report, indent=2, ensure_ascii=False) + '\n')
source = {'paper_base_commit': '7af1a137e2676a018e1e880ab076de6cae4ce30b',
          'previous_reviewed_paper_commit': 'ac75a339c8db2861ff8e7d072e50690c79c602c4',
          'q15_scientific_code_revision': '271af288a2f3863430ab80e3145c2dee9bd5571d',
          'q15_verified_results_commit': 'bc48b257eb44f412ad069f50d0f1a72a33c3c520',
          'q15_results_branch': 'q15/run-20261005T050511Z-migration-from-r2-b82ad79b',
          'paper_branch': 'paper/completed-q15-20261005',
          'source_fit_count': 15, 'target_fit_count': 0, 'new_model_fits': 0, 'new_checkpoint_inference': 0,
          'input_evidence_sha256': input_hashes, 'raw_data_redistributed': False}
(OUT / 'evidence/source_snapshot.json').write_text(json.dumps(source, indent=2) + '\n')
print(json.dumps({'status': report['status'], 'checks': len(checks), 'pages': len(pdf), 'new_fits': 0}))
